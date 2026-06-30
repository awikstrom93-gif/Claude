"""
r2k_view_concentration.py  --  step8's "Weight Concentration" tab as a pure projection of the panel
(weights only; both indices).  step8's Return Breadth / Return Concentration tabs are returns-based
and remain in r2k_step8_concentration.py.

RUN:  python r2k_view_concentration.py
"""
from r2k_universe import get_panel, by_index_year, diff_sheet, print_sheet, BASE

CONC = BASE / "R2000G_Concentration.xlsx"
TOP_NS = [1, 5, 10, 25, 50]
HDR = ["Year", "R2KG #", "R2KG Top10%", "R2KG Top25%", "R2KG Max name%", "R2KG HHI", "R2KG EffN",
       "600G #", "600G Top10%", "600G Top25%", "600G Max name%", "600G HHI", "600G EffN"]


def weight_conc(members):
    """Mirror step8.weight_conc exactly, from panel member rows."""
    ws = sorted((r["weight"] for r in members), reverse=True)
    tw = sum(ws) or 1e-9
    shares = [w / tw for w in ws]
    return {"n": len(ws), "top": {k: round(sum(ws[:k]) / tw * 100, 2) for k in TOP_NS},
            "max": round(ws[0] / tw * 100, 2) if ws else None,
            "hhi": round(sum((s * 100) ** 2 for s in shares), 1),
            "effn": round(1 / sum(s * s for s in shares), 0) if shares else None}


def weight_concentration_rows(panel):
    grp = by_index_year(panel)
    years = sorted(y for (ix, y) in grp if ix == "R2KG")
    out = []
    for y in years:
        a = weight_conc(grp[("R2KG", y)])
        smem = grp.get(("SP600G", y))
        b = weight_conc(smem) if smem else None
        row = [y, a["n"], a["top"][10], a["top"][25], a["max"], a["hhi"], a["effn"]]
        row += ([b["n"], b["top"][10], b["top"][25], b["max"], b["hhi"], b["effn"]] if b else [None] * 6)
        out.append(row)
    return out


def write_sheet(wb, panel):
    from openpyxl.styles import Font
    ws = wb.create_sheet("Weight Concentration")
    ws.cell(row=1, column=1, value="Weight concentration -- how top-heavy each benchmark is").font = Font(bold=True, size=12)
    for c, h in enumerate(HDR, 1):
        ws.cell(row=3, column=c, value=h).font = Font(bold=True)
    for i, row in enumerate(weight_concentration_rows(panel), start=4):
        for c, v in enumerate(row, 1):
            ws.cell(row=i, column=c, value=v)
    return ws


def main():
    panel = get_panel(index=None)
    rows = weight_concentration_rows(panel)
    print("\n  Weight Concentration (panel-derived):")
    print_sheet(HDR, rows, ncols=7)
    diff_sheet("Weight Concentration", HDR, rows, src=CONC)


if __name__ == "__main__":
    main()
