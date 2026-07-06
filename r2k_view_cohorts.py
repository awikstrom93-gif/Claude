"""
r2k_view_cohorts.py  --  Profitability Cohorts sheet as a pure projection of the canonical panel.
Reproduces step3's "Profitability Cohorts" exactly (counts + index-weight of unprofitable / never-
profitable / fallen-angel / multi-year-profitable cohorts).  Only intended change vs the current
workbook is the base-first identity + canonical universe already validated on Quality Trends.

RUN:  python r2k_view_cohorts.py        # builds the sheet from the panel + diffs vs current
"""
from r2k_universe import get_panel, diff_sheet, print_sheet, year_snapshot

SHEET = "Profitability Cohorts"
TITLE_TEXT = "Profitability cohorts (count and index weight)"
HDR = ["Snapshot", "Covered", "Unprof NI %cnt", "Unprof NI %wt", "Unprof OI %cnt", "Unprof OI %wt",
       "Never (cnt)", "  confirmed", "  ltd-hist", "Fallen (cnt)", "Prof 2y %wt", "Prof 3y %wt",
       "% with Revenue"]


def cohorts_rows(panel):
    years = sorted({int(r["year"]) for r in panel})
    snap_of = year_snapshot(panel)                       # actual snapshot date per year (SNAP_MONTH-driven)
    out = []
    for yr in years:
        cov = [r for r in panel if int(r["year"]) == yr and r["covered"]]
        n = len(cov)
        wtot = sum(r["weight"] for r in cov) or 1
        cls = [r for r in cov if r["prof_ni"] is not None]
        un = [r for r in cov if r["prof_ni"] is False]
        upc = 100 * len(un) / len(cls) if cls else None
        upw = 100 * sum(r["weight"] for r in un) / (sum(r["weight"] for r in cls) or 1) if cls else None
        clo = [r for r in cov if r["prof_oi"] is not None]
        uno = [r for r in cov if r["prof_oi"] is False]
        uoc = 100 * len(uno) / len(clo) if clo else None
        uow = 100 * sum(r["weight"] for r in uno) / (sum(r["weight"] for r in clo) or 1) if clo else None
        never = [r for r in un if r["cohort"] == "never_profitable"]
        nconf = [r for r in never if r["never_basis"] == "confirmed"]
        nlim = [r for r in never if r["never_basis"] == "limited_history"]
        fallen = [r for r in un if r["cohort"] == "fallen"]
        w2 = 100 * sum(r["weight"] for r in cov if r["p2"]) / wtot
        w3 = 100 * sum(r["weight"] for r in cov if r["p3"]) / wtot
        prev = 100 * sum(r["weight"] for r in cov if r["has_rev"]) / wtot
        out.append([snap_of.get(yr, f"{yr}-06-30"), n,
                    round(upc, 1) if upc is not None else None, round(upw, 1) if upw is not None else None,
                    round(uoc, 1) if uoc is not None else None, round(uow, 1) if uow is not None else None,
                    len(never), len(nconf), len(nlim), len(fallen),
                    round(w2, 1), round(w3, 1), round(prev, 1)])
    return out


def write_sheet(wb, panel):
    from openpyxl.styles import Font
    ws = wb.create_sheet(SHEET)
    ws.cell(row=1, column=1, value=TITLE_TEXT).font = Font(bold=True, size=12)
    for c, h in enumerate(HDR, 1):
        ws.cell(row=3, column=c, value=h).font = Font(bold=True)
    for i, row in enumerate(cohorts_rows(panel), start=4):
        for c, v in enumerate(row, 1):
            ws.cell(row=i, column=c, value=v)
    return ws


def main():
    panel = get_panel()
    rows = cohorts_rows(panel)
    print(f"\n  {SHEET} (panel-derived):")
    print_sheet(HDR, rows, ncols=6)
    diff_sheet(SHEET, HDR, rows)


if __name__ == "__main__":
    main()
