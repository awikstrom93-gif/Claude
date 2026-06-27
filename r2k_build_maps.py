"""
============================================================
r2k_build_maps.py  --  build the holdings->CIK maps AND the combined universe CIK list directly
from the holdings workbooks, so the pipeline is fully self-contained on the DERA foundation
(no dependency on the old r2k_step2_asfiled.py for these).
============================================================
Closes the two workflow gaps:
  1. security_cik_map.json / temporal_cik_map.json  -- ticker->CIK resolution used by steps 3/5/6.
  2. universe_ciks.csv  -- every CIK across R2000G + S&P 600 Growth holdings, so the DERA index
     (r2k_dera_index.py) covers BOTH indices and step 6's comparison works.

Resolution per holding (best -> fallback):
  CIK column (authoritative; the holdings carry it) -> CUSIP via cusip2cik.json -> Ticker via
  ticker2cik.json. Whatever can't be resolved is reported, never guessed.

INPUTS  (R2KG_BASE)
  *[Hh]olding*.xlsx          all holdings workbooks (R2000G + S&P 600 Growth), snapshot tabs
                             like 4.30.2026; each row has CIK / Name / Ticker / CUSIP / weight.
  cusip2cik.json, ticker2cik.json   optional fallbacks (from r2k_morningstar_parse.py)
OUTPUT
  security_cik_map.json      {ticker: cik}            (union across all snapshots)
  temporal_cik_map.json      {"YYYY-04-30": {ticker: cik}}   (point-in-time, matches step 3's key)
  universe_ciks.csv          one CIK per line (cik column) -> feed r2k_dera_index.py

RUN:  python r2k_build_maps.py
SELFTEST: python r2k_build_maps.py --selftest
============================================================
"""
from pathlib import Path
import os, re, csv, sys, json
from collections import defaultdict

import openpyxl

BASE = Path(os.environ.get("R2KG_BASE", "."))
SEC_MAP = BASE / "security_cik_map.json"
TEMP_MAP = BASE / "temporal_cik_map.json"
UNIVERSE = BASE / "universe_ciks.csv"
CUSIP2CIK = BASE / "cusip2cik.json"
TICKER2CIK = BASE / "ticker2cik.json"


def snapshot_year(sn):
    m = re.search(r"(\d+)\.(\d+)\.(\d+)", str(sn))
    if not m:
        return None
    g3 = m.group(3)
    return int("20" + g3) if len(g3) == 2 else int(g3)


def clean_cik(v):
    if v is None:
        return ""
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return str(int(s)) if s.isdigit() else ""


def find_holdings():
    pats = ["*[Rr]ussell*[Gg]rowth*[Hh]olding*.xlsx", "*600*[Gg]rowth*[Hh]olding*.xlsx",
            "*[Hh]olding*.xlsx"]
    files = []
    for p in pats:
        files += list(BASE.glob(p))
    # dedup, keep order
    seen, out = set(), []
    for f in files:
        if f.resolve() not in seen:
            seen.add(f.resolve()); out.append(f)
    return out


def load_json(p):
    try:
        return json.loads(p.read_text()) if p.exists() else {}
    except Exception:
        return {}


def build(holdings_files, cusip2cik, ticker2cik):
    sec_map = {}                              # ticker -> cik
    temporal = defaultdict(dict)             # "YYYY-04-30" -> {ticker: cik}
    universe = set()
    src = defaultdict(int); unresolved = []
    for f in holdings_files:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        for sn in wb.sheetnames:
            yr = snapshot_year(sn)
            if yr is None:
                continue
            ws = wb[sn]; hdr = None; ix = {}
            for r in ws.iter_rows(values_only=True):
                if hdr is None:
                    hdr = [str(c).strip() if c else "" for c in r]
                    ix = {h: i for i, h in enumerate(hdr)}
                    if "Ticker" not in ix:
                        hdr = None
                    continue
                def g(col):
                    i = ix.get(col)
                    return r[i] if (i is not None and i < len(r)) else None
                tk = g("Ticker")
                if not tk:
                    continue
                tk = str(tk).strip()
                cik = clean_cik(g("CIK"))
                if cik:
                    src["CIK column"] += 1
                else:
                    cusip = str(g("CUSIP") or "").strip()
                    if cusip and cusip in cusip2cik:
                        cik = clean_cik(cusip2cik[cusip]); src["CUSIP fallback"] += 1
                    elif tk in ticker2cik:
                        cik = clean_cik(ticker2cik[tk]); src["ticker fallback"] += 1
                if not cik:
                    unresolved.append((f.name, yr, tk)); continue
                sec_map[tk] = cik
                temporal[f"{yr:04d}-04-30"][tk] = cik
                universe.add(cik)
        wb.close()
    return sec_map, dict(temporal), sorted(universe), dict(src), unresolved


def main():
    holdings = find_holdings()
    if not holdings:
        raise SystemExit("!! no *Holding*.xlsx found in R2KG_BASE.")
    print(f"  holdings workbooks: {[h.name for h in holdings]}")
    cusip2cik = load_json(CUSIP2CIK); ticker2cik = load_json(TICKER2CIK)
    sec_map, temporal, universe, src, unresolved = build(holdings, cusip2cik, ticker2cik)

    SEC_MAP.write_text(json.dumps(sec_map, indent=0))
    TEMP_MAP.write_text(json.dumps(temporal, indent=0))
    with open(UNIVERSE, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(["cik"])
        for c in universe:
            w.writerow([c])

    print(f"  resolution: " + ", ".join(f"{k} {v}" for k, v in src.items()))
    print(f"  -> {SEC_MAP.name}: {len(sec_map)} tickers")
    print(f"  -> {TEMP_MAP.name}: {len(temporal)} snapshots")
    print(f"  -> {UNIVERSE.name}: {len(universe)} unique CIKs (R2000G + S&P 600 Growth)")
    if unresolved:
        u = sorted({(t,) for _, _, t in unresolved})
        print(f"  !! {len(unresolved)} holding-rows unresolved ({len(u)} distinct tickers) -- "
              f"no CIK/CUSIP/ticker match. e.g. {[t[0] for t in u[:8]]}")
    print(f"\n  Next: point the DERA index at the full universe so step 6 (R2000G vs 600G) works:")
    print(f"     set R2KG_CIK_FILE={UNIVERSE.name}  (or r2k_dera_index.py auto-detects it)")


def selftest():
    import tempfile
    d = Path(tempfile.mkdtemp()); wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "4.30.2026"
    ws.append(["CIK", "Name", "Ticker", "ISIN", "CUSIP", "Portfolio Weighting %"])
    ws.append([1664703, "Bloom Energy", "BE", "US0937121079", "093712107", 3.65])
    ws.append([None, "NoCik Co", "NCK", "US123", "12345678X", 0.5])      # resolve via cusip fallback
    ws2 = wb.create_sheet("4.30.2025")
    ws2.append(["CIK", "Name", "Ticker", "CUSIP", "Portfolio Weighting %"])
    ws2.append([36029, "First Financial", "FFIN", "320218108", 0.9])
    hp = d / "Russell_2000_Growth_Holdings.xlsx"; wb.save(hp)
    global BASE, SEC_MAP, TEMP_MAP, UNIVERSE, CUSIP2CIK, TICKER2CIK
    BASE = d; SEC_MAP = d / "security_cik_map.json"; TEMP_MAP = d / "temporal_cik_map.json"
    UNIVERSE = d / "universe_ciks.csv"
    CUSIP2CIK = d / "cusip2cik.json"; CUSIP2CIK.write_text(json.dumps({"12345678X": "999999"}))
    TICKER2CIK = d / "t.json"
    main()
    sec = json.loads(SEC_MAP.read_text()); temp = json.loads(TEMP_MAP.read_text())
    uni = [r["cik"] for r in csv.DictReader(open(UNIVERSE))]
    ok = (sec.get("BE") == "1664703" and sec.get("NCK") == "999999" and sec.get("FFIN") == "36029"
          and temp.get("2026-04-30", {}).get("BE") == "1664703"
          and temp.get("2025-04-30", {}).get("FFIN") == "36029"
          and set(uni) == {"1664703", "999999", "36029"})
    print(f"\n  SELFTEST (CIK col + CUSIP fallback + temporal key + universe): {'PASS' if ok else 'FAIL'}")


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        selftest()
    else:
        main()
