"""
============================================================
r2k_refresh_charts_data.py  --  push fresh data from the step-7 output INTO your charted
workbook, updating only the cell values and leaving every chart byte-for-byte untouched.
============================================================
Re-running step 7 rebuilds the workbook and round-trips charts through openpyxl, which
drops some chart formatting. This avoids that entirely: it opens your charted file as a
zip, rewrites only the cell VALUES inside each worksheet's XML to match the fresh step-7
output, and copies every other part (charts, drawings, styles, themes) verbatim. The
charts never pass through a chart library, so they are preserved exactly.

HEADER-AWARE MATCHING (why this exists)
    The charts in your backup reference cells by POSITION -- e.g. 'Qual R2000G'!$M$4:$M$15.
    If an analytics step reorders or inserts columns (e.g. step 6 inserting the
    dollar-aggregate margin columns), a naive position-for-position copy would write the
    NEW column's data into the OLD column's cells, so every chart would silently plot the
    wrong metric under the right legend.

    Instead, for every data table this tool matches columns by their HEADER LABEL and rows
    by position (rows are stable across runs -- same years / constituents). For each column
    in your charts file it finds the column in the fresh data with the same header and pulls
    THAT column's header + values into the charts-file position. So the metric a chart
    expects always lands where the chart looks, no matter how the fresh output reordered its
    columns. Tables that didn't move are unaffected (label match is identity). Cells outside
    any detected table (titles, prose, notes) fall back to same-position matching.

    A few columns were RENAMED between layouts (the margin columns moved from a weight-
    weighted-average basis to the validated dollar-aggregate basis). ALIASES below map the
    old chart header to the fresh column that replaces it, so existing margin charts keep
    working and now show the corrected basis. Set MARGIN_BASIS=wavg to prefer the
    weight-weighted columns instead (only where the fresh output still carries them).
    Anything that still can't be matched is REPORTED, never silently mis-filled.

INPUTS  (R2KG_BASE)
    R2000G_SmallCapGrowth_Benchmark_Review.xlsx   the fresh step-7 output (DATA source)
                                                   override with R2KG_DATA_FILE
    R2000G_charts_backup.xlsx                      your charted workbook (CHARTS to keep)
                                                   override with R2KG_CHARTS_FILE
OUTPUT
    R2000G_Benchmark_Review_charted.xlsx           fresh data + your exact charts
                                                   override with R2KG_CHARTS_OUT

RUN: python r2k_refresh_charts_data.py
============================================================
"""
from pathlib import Path
from xml.sax.saxutils import escape, unescape
import os, re, zipfile

import openpyxl
from openpyxl.utils import get_column_letter

BASE = Path(os.environ.get("R2KG_BASE", "."))
DATA_FILE = BASE / os.environ.get("R2KG_DATA_FILE", "R2000G_SmallCapGrowth_Benchmark_Review.xlsx")
CHARTS_FILE = BASE / os.environ.get("R2KG_CHARTS_FILE", "R2000G_charts_backup.xlsx")
OUT = BASE / os.environ.get("R2KG_CHARTS_OUT", "R2000G_Benchmark_Review_charted.xlsx")
# which margin basis the renamed margin charts should follow: "agg" (validated, default) or "wavg"
MARGIN_BASIS = os.environ.get("MARGIN_BASIS", "agg").strip().lower()
# "header" (default, robust to column reordering) or "position" (legacy same-cell copy)
REFRESH_MODE = os.environ.get("R2KG_REFRESH_MODE", "header").strip().lower()


def build_cell(ref, style, val):
    st = f' s="{style}"' if style else ""
    if isinstance(val, bool):
        return f'<c r="{ref}"{st} t="b"><v>{1 if val else 0}</v></c>'
    if isinstance(val, (int, float)):
        v = ("%d" % val) if (isinstance(val, int) or float(val).is_integer()) else repr(val)
        return f'<c r="{ref}"{st}><v>{v}</v></c>'
    return f'<c r="{ref}"{st} t="inlineStr"><is><t xml:space="preserve">{escape(str(val))}</t></is></c>'


CELL_RE = re.compile(r'<c r="([A-Z]+\d+)"([^>]*?)(/>|>.*?</c>)', re.DOTALL)
STYLE_RE = re.compile(r'\bs="(\d+)"')


def update_worksheet_xml(xml_text, fresh):
    """fresh: {A1ref: value}. Update matching VALUE cells in place; leave FORMULA cells alone
    (they recompute from the refreshed inputs -- and overwriting them orphans calcChain.xml,
    which corrupts the file). Returns (new_text, n_updated, n_skipped_formula, missing_refs)."""
    seen = set(); skipped_formula = [0]
    def repl(m):
        ref = m.group(1)
        if ref not in fresh:
            return m.group(0)
        body = m.group(3)
        if "<f>" in body or "<f " in body or "<f/>" in body:   # formula cell -- never touch
            skipped_formula[0] += 1; seen.add(ref); return m.group(0)
        sm = STYLE_RE.search(m.group(2))                       # preserve the cell's style
        seen.add(ref)
        return build_cell(ref, sm.group(1) if sm else None, fresh[ref])
    new_text = CELL_RE.sub(repl, xml_text)
    missing = [r for r in fresh if r not in seen]
    return new_text, len(seen) - skipped_formula[0], skipped_formula[0], missing


# ---------------------------------------------------------------------------
# header-aware column matching
# ---------------------------------------------------------------------------
def norm(s):
    return re.sub(r"\s+", " ", str(s)).strip()


def alias_candidates(label):
    """Ordered fresh-header candidates for an old charts-file header that was renamed.
    The margin columns moved from weight-weighted average to dollar-aggregate; older
    layouts also used a bare 'OpMgn'/'GrossMgn'/'NetMgn' (which was the wavg figure)."""
    lbl = norm(label)
    cands = [lbl]
    prefer = ["$agg", "wavg"] if MARGIN_BASIS != "wavg" else ["wavg", "$agg"]

    # bare margin name (old single column) -> qualified variants
    m = re.fullmatch(r"(Gross|Op|Net)Mgn", lbl)
    if m:
        for q in prefer:
            cands.append(f"{m.group(0)} {q}")

    # "<metric> (wavg) <suffix>" <-> "<metric> ($agg) <suffix>"  (Comparison tab style)
    if "(wavg)" in lbl or "($agg)" in lbl:
        for q in prefer:
            cands.append(lbl.replace("(wavg)", f"({q})").replace("($agg)", f"({q})"))

    # "<metric> wavg"/"<metric> $agg" trailing-qualifier style
    m2 = re.fullmatch(r"(.*?)\s+(wavg|\$agg)", lbl)
    if m2:
        for q in prefer:
            cands.append(f"{m2.group(1)} {q}")

    out = []
    for c in cands:
        if c not in out:
            out.append(c)
    return out


def is_header_row(grid, r, maxc):
    """Heuristic: a header row has several short string labels and is followed by a row of
    mostly numbers. Distinguishes table headers from titles, prose and footnotes."""
    row = grid.get(r, {})
    strs = [v for v in row.values() if isinstance(v, str) and v.strip()]
    short = [v for v in strs if len(v) <= 60]
    if len(short) < 2:
        return False
    nxt = grid.get(r + 1, {})
    nums = [v for v in nxt.values() if isinstance(v, (int, float)) and not isinstance(v, bool)]
    return len(nums) >= 2


def grid_of(ws):
    g = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                g.setdefault(c.row, {})[c.column] = c.value
    return g


def remap_sheet(bak_grid, fresh_grid, report_unmatched, report_alias, sheet):
    """Return {A1ref: value} to write into the charts(bak) worksheet, with each table's
    columns matched fresh<-bak by header label. Rows are matched by position. Cells outside
    a detected table fall back to same-position fresh values."""
    maxr = max([*bak_grid, *fresh_grid, 1])
    maxc = max([max(r) for r in [*bak_grid.values(), *fresh_grid.values()] if r] + [1])

    # locate header rows (shared row indices -- layout is row-stable across runs)
    header_rows = [r for r in range(1, maxr + 1) if is_header_row(bak_grid, r, maxc)]
    # the table a header governs runs until the next header row (or end)
    regions = []
    for i, hr in enumerate(header_rows):
        end = header_rows[i + 1] - 1 if i + 1 < len(header_rows) else maxr
        regions.append((hr, end))

    covered = set()                       # (row) indices governed by some table region
    out = {}
    for hr, end in regions:
        bak_hdr = {c: norm(v) for c, v in bak_grid.get(hr, {}).items() if isinstance(v, str) and v.strip()}
        fresh_hdr_by_label = {}
        for c, v in fresh_grid.get(hr, {}).items():
            if isinstance(v, str) and v.strip():
                fresh_hdr_by_label.setdefault(norm(v), c)
        if not bak_hdr or not fresh_hdr_by_label:
            continue
        # column map: bak col -> fresh col, by header label (with aliases)
        col_map = {}
        for bc, lbl in bak_hdr.items():
            target = None
            for cand in alias_candidates(lbl):
                if cand in fresh_hdr_by_label:
                    target = fresh_hdr_by_label[cand]
                    if cand != lbl:
                        report_alias.append((sheet, lbl, cand))
                    break
            if target is None:
                report_unmatched.append((sheet, hr, lbl))
            else:
                col_map[bc] = target
        # write header label + every data row for each matched column
        for r in range(hr, end + 1):
            covered.add(r)
            for bc, fc in col_map.items():
                val = fresh_grid.get(r, {}).get(fc)
                if val is not None:
                    out[f"{get_column_letter(bc)}{r}"] = val

    # cells outside any table region -> same-position fresh value (titles, notes, single cells)
    for r, cols in fresh_grid.items():
        if r in covered:
            continue
        for c, v in cols.items():
            out[f"{get_column_letter(c)}{r}"] = v
    return out


def main():
    if not DATA_FILE.exists():
        raise SystemExit(f"!! {DATA_FILE.name} not found (run step 7 first).")
    if not CHARTS_FILE.exists():
        raise SystemExit(f"!! {CHARTS_FILE.name} not found "
                         f"(run r2k_snapshot_charts.py, or set R2KG_CHARTS_FILE).")

    # 1. fresh values from the step-7 output, plus the charts file's own grids (for headers)
    src = openpyxl.load_workbook(DATA_FILE, data_only=True)
    fresh_grids = {sn: grid_of(src[sn]) for sn in src.sheetnames}
    src.close()
    bak = openpyxl.load_workbook(CHARTS_FILE, data_only=True)
    bak_grids = {sn: grid_of(bak[sn]) for sn in bak.sheetnames}
    bak.close()

    report_unmatched, report_alias = [], []
    remapped = {}
    for sn, fg in fresh_grids.items():
        if sn not in bak_grids:
            continue
        if REFRESH_MODE == "position":
            remapped[sn] = {f"{get_column_letter(c)}{r}": v
                            for r, cols in fg.items() for c, v in cols.items()}
        else:
            remapped[sn] = remap_sheet(bak_grids[sn], fg, report_unmatched, report_alias, sn)

    # 2. map sheet name -> worksheet xml path in the charts file
    with zipfile.ZipFile(CHARTS_FILE) as z:
        wbxml = z.read("xl/workbook.xml").decode("utf-8")
        relsxml = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid_target = {}
    for rel in re.findall(r"<Relationship\b[^>]*/>", relsxml):
        i = re.search(r'Id="([^"]+)"', rel); t = re.search(r'Target="([^"]+)"', rel)
        if i and t: rid_target[i.group(1)] = t.group(1)
    name_to_path = {}
    for s in re.findall(r"<sheet\b[^>]*/>", wbxml):
        nm = re.search(r'name="([^"]+)"', s); rid = re.search(r'r:id="([^"]+)"', s)
        if nm and rid and rid.group(1) in rid_target:
            tgt = rid_target[rid.group(1)]
            path = tgt[1:] if tgt.startswith("/") else "xl/" + tgt
            name_to_path[unescape(nm.group(1))] = path.replace("xl/xl/", "xl/")

    # 3. update each matching worksheet's XML in place
    updated_parts = {}; report = []
    for sn, path in name_to_path.items():
        if sn not in remapped: continue
        with zipfile.ZipFile(CHARTS_FILE) as z:
            xml = z.read(path).decode("utf-8")
        new_xml, n_set, n_formula, missing = update_worksheet_xml(xml, remapped[sn])
        updated_parts[path] = new_xml.encode("utf-8")
        report.append((sn, n_set, n_formula, len(missing)))

    # strip the (optional) shared-string count attributes -- after converting some shared-string
    # cells to inline strings the totals no longer match, which Excel flags. Removing them is valid
    # OOXML and Excel recomputes on open.
    with zipfile.ZipFile(CHARTS_FILE) as z:
        if "xl/sharedStrings.xml" in z.namelist():
            ss = z.read("xl/sharedStrings.xml").decode("utf-8")
            ss = re.sub(r'\s+count="\d+"', "", ss, count=1)
            ss = re.sub(r'\s+uniqueCount="\d+"', "", ss, count=1)
            updated_parts["xl/sharedStrings.xml"] = ss.encode("utf-8")
    # force a full recalc on open so the chart-driving helper formulas pick up the refreshed data
    if "<calcPr" in wbxml and "fullCalcOnLoad" not in wbxml:
        updated_parts["xl/workbook.xml"] = re.sub(
            r"<calcPr\b", '<calcPr fullCalcOnLoad="1"', wbxml, count=1).encode("utf-8")

    in_charts = set(name_to_path); in_data = set(fresh_grids)
    only_data = sorted(in_data - in_charts)
    only_charts = sorted(in_charts - in_data)

    # 4. re-zip: replace updated worksheet parts, copy everything else (charts!) verbatim
    if OUT.exists(): OUT.unlink()
    with zipfile.ZipFile(CHARTS_FILE) as zin, zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = updated_parts.get(item.filename) or zin.read(item.filename)
            zout.writestr(item, data)

    print(f"  data source : {DATA_FILE.name}")
    print(f"  charts file : {CHARTS_FILE.name}")
    print(f"  match mode  : {REFRESH_MODE} (margins: {MARGIN_BASIS})")
    print(f"  -> {OUT.name}  (charts copied verbatim; columns matched by header)\n")
    print(f"  {'tab':<26}{'values updated':>15}{'formulas kept':>15}{'not-found':>11}")
    for sn, n, nf, miss in report:
        flag = "  <- extra fresh rows" if miss else ""
        print(f"  {sn[:26]:<26}{n:>15}{nf:>15}{miss:>11}{flag}")
    if report_alias:
        print("\n  renamed columns matched via alias (old chart header <- fresh column):")
        seen_al = set()
        for sn, old, new in report_alias:
            k = (sn, old, new)
            if k in seen_al: continue
            seen_al.add(k)
            print(f"    {sn[:22]:<22}  {old!r:<26} <- {new!r}")
    if report_unmatched:
        print("\n  *** UNMATCHED chart columns (no fresh column with this header -- left as-is,")
        print("      so any chart series on these still shows the PRIOR run's values):")
        seen_un = set()
        for sn, hr, lbl in report_unmatched:
            k = (sn, lbl)
            if k in seen_un: continue
            seen_un.add(k)
            print(f"    {sn[:22]:<22}  row {hr:<3}  {lbl!r}")
        print("      -> the fresh output renamed/removed this metric; re-point or rebuild that")
        print("         chart's series, or set MARGIN_BASIS to the basis your output still carries.")
    if only_data:
        print(f"\n  tabs in the fresh data but NOT in your charts file (won't be added): {only_data}")
    if only_charts:
        print(f"  tabs in your charts file but not in the fresh data (left as-is): {only_charts}")


if __name__ == "__main__":
    main()
