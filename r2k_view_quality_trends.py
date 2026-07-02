"""
r2k_view_quality_trends.py  --  FIRST ported view.  Reproduces step3's "Index Quality Trends" sheet
as a PURE PROJECTION of the canonical panel (r2k_universe.build_panel / r2k_panel.csv).  No holdings
or fundamentals are loaded here -- the panel is the only input.  This is the template every other
tab will follow.

The math (winsorized weight-averages, medians, dollar-aggregates) is the SAME helpers step3 uses, so
the only intended difference vs the current workbook tab is the identity fix (base-first CIK): 2016
Total NI 6.5 -> 13.2 and the handful of temporal-bug names.  main() prints a cell-by-cell diff against
the existing Russell2000Growth_Analytics.xlsx so you can confirm nothing else moved.

RUN:  python r2k_view_quality_trends.py        # builds the sheet from the panel + diffs vs current
"""
from r2k_universe import get_panel, diff_sheet, print_sheet
from r2k_step3_analytics import aggregate, dollar_agg, _p, _x

SHEET = "Index Quality Trends"
TITLE_TEXT = "Russell 2000 Growth -- Index Quality Trends (weight-weighted / median / dollar-aggregate)"
# NOTE on the margin/return aggregations (see NOTE_TEXT below): for a RATIO like operating margin,
# "$agg" = sum(numerator)/sum(denominator) is the revenue-weighted (blended, index-as-one-company)
# margin -- the economically standard figure. "med" is the typical constituent. "wavg" here is the
# INDEX-WEIGHT-weighted average of each name's (winsorized) per-name ratio; for margins that is tilted
# deeply negative because small-revenue heavy-loss names carry index weight with ratios floored at
# -200% -- so read OpMgn/GrossMgn "wavg (idx-wt)" as a loss-TILT signal, NOT the index's margin. Use
# $agg (blended) or med for the margin itself. Columns are unchanged in ORDER (charts/guard depend on
# position); only the wavg headers are clarified.
HDR = ["Snapshot", "Total Rev $B", "Total NI $B", "% with Revenue", "% Unprofitable (NI) wt",
       "% Unprofitable (OI) wt", "GrossMgn wavg (idx-wt)", "GrossMgn med", "OpMgn wavg (idx-wt)", "OpMgn med",
       "OpMgn $agg (blended)", "NetMgn med", "ROE wavg", "ROE med", "ROE $agg", "ROA med", "ROIC wavg",
       "ROIC med", "ROIC $agg", "GP/Assets med", "Accruals med", "CashConv med", "AssetTurn med",
       "Rev YoY wavg", "Rev 3yCAGR med", "FCF mgn med", "RuleOf40 med", "D/E wavg", "D/Cap wavg",
       "D/Cap $agg"]
NOTE_TEXT = ("Margins: '$agg (blended)' = sum(op income)/sum(revenue) is the index-as-one-company margin "
             "(the standard figure). 'med' = typical constituent. 'wavg (idx-wt)' = index-weight-weighted "
             "average of per-name margins, winsorized to +/-200%/500%; for margins it runs deeply NEGATIVE "
             "because small-revenue loss-makers carry index weight -- read it as a loss-tilt signal, not the "
             "index's operating margin. For the index's margin use OpMgn $agg (blended) or OpMgn med.")


def quality_trends_rows(panel):
    """The projection: panel -> one row per snapshot, columns == HDR. Mirrors step3 exactly."""
    years = sorted({int(r["year"]) for r in panel})
    out = []
    for yr in years:
        cov = [r for r in panel if int(r["year"]) == yr and r["covered"]]
        def col(k): return [(r[k], r["weight"]) for r in cov]
        ag = lambda k: aggregate(col(k), winsor=True)
        tot_rev = sum(r["revenue"] for r in cov if r["revenue"]) / 1e9
        tot_ni = sum(r["net_income"] for r in cov if r["net_income"] is not None) / 1e9
        wtot = sum(r["weight"] for r in cov) or 1
        pct_rev = 100 * sum(r["weight"] for r in cov if r["has_rev"]) / wtot
        ni_cls = [r for r in cov if r["prof_ni"] is not None]
        up_ni = (100 * sum(r["weight"] for r in ni_cls if r["prof_ni"] is False) /
                 (sum(r["weight"] for r in ni_cls) or 1)) if ni_cls else None
        oi_cls = [r for r in cov if r["prof_oi"] is not None]
        up_oi = (100 * sum(r["weight"] for r in oi_cls if r["prof_oi"] is False) /
                 (sum(r["weight"] for r in oi_cls) or 1)) if oi_cls else None
        opm_da = dollar_agg([(r["operating_income"], r["revenue"]) for r in cov])
        roe_da = dollar_agg([(r["net_income"], r["equity"]) for r in cov])
        roic_da = dollar_agg([(r["_nopat"], r["_ic"]) for r in cov])
        dcap_da = dollar_agg([(r["debt"], (r["debt"] + r["equity"])
                              if (r["debt"] is not None and r["equity"] is not None) else None) for r in cov])
        out.append([f"{yr}-04-30", round(tot_rev, 1), round(tot_ni, 1), round(pct_rev, 1),
                    round(up_ni, 1) if up_ni is not None else None,
                    round(up_oi, 1) if up_oi is not None else None,
                    _p(ag("gross_margin")["wavg"]), _p(ag("gross_margin")["median"]),
                    _p(ag("op_margin")["wavg"]), _p(ag("op_margin")["median"]), _p(opm_da),
                    _p(ag("net_margin")["median"]),
                    _p(ag("roe")["wavg"]), _p(ag("roe")["median"]), _p(roe_da), _p(ag("roa")["median"]),
                    _p(ag("roic")["wavg"]), _p(ag("roic")["median"]), _p(roic_da),
                    _p(ag("gp_to_assets")["median"]), _p(ag("accruals")["median"]),
                    _x(ag("cash_conversion")["median"]), _x(ag("asset_turnover")["median"]),
                    _p(ag("rev_yoy")["wavg"]), _p(ag("rev_cagr3")["median"]), _p(ag("fcf_margin")["median"]),
                    _p(ag("rule_of_40")["median"]),
                    _x(ag("d_to_equity")["wavg"]), _p(ag("d_to_capital")["wavg"]), _p(dcap_da)])
    return out


def write_sheet(wb, panel):
    """Add the Index Quality Trends sheet to an openpyxl workbook. Used by the future r2k_report.py."""
    from openpyxl.styles import Font
    ws = wb.create_sheet(SHEET)
    ws.cell(row=1, column=1, value=TITLE_TEXT).font = Font(bold=True, size=12)
    ws.cell(row=2, column=1, value=NOTE_TEXT).font = Font(size=9, italic=True, color="555555")
    for c, h in enumerate(HDR, 1):
        ws.cell(row=3, column=c, value=h).font = Font(bold=True)
    for i, row in enumerate(quality_trends_rows(panel), start=4):
        for c, v in enumerate(row, 1):
            ws.cell(row=i, column=c, value=v)
    return ws


def main():
    panel = get_panel()
    rows = quality_trends_rows(panel)
    print(f"\n  {SHEET} (panel-derived):")
    print_sheet(HDR, rows)
    diff_sheet(SHEET, HDR, rows)


if __name__ == "__main__":
    main()
