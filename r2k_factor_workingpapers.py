"""
r2k_factor_workingpapers.py -- emit ONLY the granular factor working files, with a lean dependency
footprint (no panel, no fundamentals_dera.csv needed -- just the factor inputs).

Reuses the SAME building blocks as the pipeline (r2k_audit_pack.factor_audit + factor_reliability),
so the output reconciles to the published Factor Summary / Attribution / Reliability by construction.

INPUTS required in the working directory (auto-globbed):
    *Monthly*Performance*.xlsx        constituent monthly TOTAL returns (Morningstar Direct export)
    *Market*Cap*.xlsx                 constituent month-end market cap (Morningstar Direct)
    fundamentals_dera_resolved.csv    as-filed, recovered fundamentals (value/quality/growth inputs)
    *Russell*Growth*Holdings*.xlsx    R2000G holdings: weight + GICS sector, point-in-time
    *600*Growth*Holdings*.xlsx        S&P600G holdings: weight + GICS sector, point-in-time

OUTPUTS:
    audit_factor_crosssection.csv     every regression INPUT   (index x month x name)
    audit_factor_coefficients.csv     every regression OUTPUT  (index x month)
    audit_factor_correlation.csv      factor-z correlation matrix + VIF, per index
    audit_factor_uni_vs_multi.csv     univariate vs multivariate slope + t, per index x factor
    R2000G_Factor_Granular.xlsx       sample-month LINEST reproduction + Factor Reliability tab

RUN:  python r2k_factor_workingpapers.py
"""
import openpyxl
import r2k_audit_pack as ap


def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ap.factor_audit(wb)          # -> crosssection + coefficients CSVs, and the sample-month tab
    ap.factor_reliability(wb)    # -> correlation + uni-vs-multi CSVs, and the Factor Reliability tab
    if not wb.sheetnames:
        print("  !! no factor tabs produced -- check that the factor input files are present (see the "
              "module docstring for the required globs).")
        return
    out = ap.BASE / "R2000G_Factor_Granular.xlsx"
    wb.save(out)
    print(f"  wrote {out.name}  ({len(wb.sheetnames)} tabs) + the four audit_factor_*.csv files")


if __name__ == "__main__":
    main()
