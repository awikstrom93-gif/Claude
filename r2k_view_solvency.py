"""
r2k_view_solvency.py  --  the solvency / interest-coverage TAIL of R2000G (the credit/rate-risk view
the IC asks about but the leverage tabs only proxy with D/E and D/Cap).  For each annual snapshot it
measures, by index weight: the share that CANNOT cover its interest (EBIT < interest), the share with
net-debt/EBITDA > 4x, the share with NEGATIVE EBITDA (and how much of that also carries net debt),
plus median interest coverage and median net-debt/EBITDA.  R2000G vs S&P 600 Growth.

Joins the panel (membership, weight, EBIT=operating_income, debt, cash) to fundamentals_dera.csv
(ebitda, interest_expense).  RUN:  python r2k_view_solvency.py
"""
import csv
from statistics import median

from r2k_universe import get_panel, by_index_year, BASE

FUND = BASE / "fundamentals_dera.csv"
OUT = BASE / "r2k_solvency.txt"
HDR = ["Year", "n", "%Can'tCoverInt", "%Cover<2x", "MedCover(x)", "%ND/EBITDA>4x",
       "MedND/EBITDA(x)", "%NegEBITDA", "%NegEBITDA+NetDebt", "600G %Can'tCover"]


def _f(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def _ck(c):
    s = str(c)
    return str(int(s)) if s.isdigit() else s


def _inputs():
    """{(cik, fy): {ebitda, interest}} from fundamentals_dera.csv."""
    d = {}
    if FUND.exists():
        for r in csv.DictReader(open(FUND, encoding="utf-8")):
            if r.get("fiscal_year", "").isdigit():
                d[(_ck(r["cik"]), r["fiscal_year"])] = {"ebitda": _f(r.get("ebitda")),
                                                         "interest": _f(r.get("interest_expense"))}
    return d


def _name_metrics(r, inp):
    """Per name: (coverage, net_debt/ebitda, ebitda, net_debt) or Nones where not computable."""
    k = (_ck(r["cik"]), str(r["fy0"]))
    e = inp.get(k, {})
    ebit = r.get("operating_income")
    interest = e.get("interest")
    ebitda = e.get("ebitda")
    debt, cash = r.get("debt"), r.get("cash")
    nd = (debt - (cash or 0)) if debt is not None else None
    cover = (ebit / interest) if (ebit is not None and interest and interest > 0) else None
    nde = (nd / ebitda) if (nd is not None and ebitda and ebitda > 0) else None
    return cover, nde, ebitda, nd


def _solvency_for(members, inp):
    cov = [r for r in members if r["covered"]]
    tw = sum(r["weight"] for r in cov) or 1e-9
    cant = c2 = nde4 = nege = nege_nd = 0.0
    covers, ndes = [], []
    for r in cov:
        cover, nde, ebitda, nd = _name_metrics(r, inp)
        w = r["weight"]
        if cover is not None:
            covers.append(cover)
            if cover < 1: cant += w
            if cover < 2: c2 += w
        if nde is not None:
            ndes.append(nde)
            if nde > 4: nde4 += w
        if ebitda is not None and ebitda < 0:
            nege += w
            if nd is not None and nd > 0:
                nege_nd += w
    return {"n": len(cov), "cant": 100 * cant / tw, "c2": 100 * c2 / tw,
            "medcov": median(covers) if covers else None, "nde4": 100 * nde4 / tw,
            "mednde": median(ndes) if ndes else None, "nege": 100 * nege / tw,
            "nege_nd": 100 * nege_nd / tw}


def solvency_rows(panel, inp):
    grp = by_index_year(panel)
    years = sorted(y for (ix, y) in grp if ix == "R2KG")
    out = []
    for y in years:
        a = _solvency_for(grp[("R2KG", y)], inp)
        b = _solvency_for(grp[("SP600G", y)], inp) if ("SP600G", y) in grp else None
        out.append([y, a["n"], round(a["cant"], 1), round(a["c2"], 1),
                    round(a["medcov"], 1) if a["medcov"] is not None else None, round(a["nde4"], 1),
                    round(a["mednde"], 1) if a["mednde"] is not None else None,
                    round(a["nege"], 1), round(a["nege_nd"], 1),
                    round(b["cant"], 1) if b else None])
    return out


def write_sheet(wb, panel=None):
    from openpyxl.styles import Font, PatternFill, Alignment
    panel = panel if panel is not None else get_panel(index=None)
    rows = solvency_rows(panel, _inputs())
    ws = wb.create_sheet("Solvency Tail")
    ws.cell(1, 1, "Solvency / interest-coverage tail -- the credit & rate-sensitivity risk of R2000G").font = Font(bold=True, size=12)
    ws.cell(2, 1, "By index weight: share that can't cover interest (EBIT<interest), high net-debt/EBITDA, "
                  "or negative EBITDA. EBIT=operating income; net debt=total debt-cash. 600G column for "
                  "contrast.").font = Font(size=9, italic=True, color="555555")
    fill = PatternFill("solid", fgColor="1F4E5F")
    for c, h in enumerate(HDR, 1):
        x = ws.cell(4, c, h); x.fill = fill; x.font = Font(bold=True, color="FFFFFF", size=10)
        x.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, row in enumerate(rows, start=5):
        for c, v in enumerate(row, 1):
            ws.cell(i, c, v)
    ws.freeze_panes = "B5"
    return "Solvency Tail"


def main():
    panel = get_panel(index=None)
    if not FUND.exists():
        print(f"  !! {FUND.name} not found -- run r2k_dera_classify.py first.")
        return
    rows = solvency_rows(panel, _inputs())
    L = ["SOLVENCY / INTEREST-COVERAGE TAIL  --  R2000G credit & rate risk (by index weight)", ""]
    L.append("  " + "".join(str(h)[:15].rjust(17) for h in HDR))
    L.append("  " + "-" * (17 * len(HDR)))
    for row in rows:
        L.append("  " + "".join(("" if v is None else str(v)).rjust(17) for v in row))
    L.append("")
    L.append("  READ: '%Can't cover interest' = index weight whose EBIT < interest expense (coverage <1x)")
    L.append("  -- the names most exposed to refinancing/rate risk. 'Neg EBITDA + net debt' = burning cash")
    L.append("  AND levered (the sharpest distress). 600G column shows R2000G's heavier low-quality tail.")
    OUT.write_text("\n".join(L), encoding="utf-8")
    print("\n".join(L))
    print(f"\n  -> {OUT.name}")


if __name__ == "__main__":
    main()
