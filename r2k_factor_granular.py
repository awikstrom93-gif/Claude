"""
r2k_factor_granular.py -- self-contained generator for the granular factor working files.

Depends only on r2k_factor_analysis.py (the factor engine that reads the inputs) + openpyxl.

INPUTS (auto-globbed from the working dir by r2k_factor_analysis):
    *Monthly*Performance*.xlsx        constituent monthly TOTAL returns (Morningstar Direct)
    *Market*Cap*.xlsx                 constituent month-end market cap (Morningstar Direct)
    fundamentals_dera_resolved.csv    as-filed, recovered fundamentals (value/quality/growth inputs)
    *Russell*Growth*Holdings*.xlsx    R2000G holdings: weight + GICS sector, point-in-time
    *600*Growth*Holdings*.xlsx        S&P600G holdings: weight + GICS sector, point-in-time

OUTPUTS:
    audit_factor_crosssection.csv     every regression INPUT   (index x month x name)
    audit_factor_coefficients.csv     every regression OUTPUT  (index x month)
    audit_factor_correlation.csv      factor-z correlation matrix + VIF, per index
    audit_factor_uni_vs_multi.csv     univariate vs multivariate slope + t, per index x factor
    R2000G_Factor_Granular.xlsx       sample-month LINEST reproduction + Factor Reliability tab

RUN:  python r2k_factor_granular.py
"""
import os
import csv
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = Path(os.environ.get("R2KG_BASE", "."))
OUT_XS = BASE / "audit_factor_crosssection.csv"
OUT_COEF = BASE / "audit_factor_coefficients.csv"
OUT_CORR = BASE / "audit_factor_correlation.csv"
OUT_UVM = BASE / "audit_factor_uni_vs_multi.csv"

TITLE = Font(bold=True, size=13, color="1F4E5F")
H = Font(bold=True, size=11, color="1F4E5F")
BODY = Font(size=10)
HDR = PatternFill("solid", fgColor="1F4E5F"); HF = Font(bold=True, color="FFFFFF", size=10)


def _hdr(ws, row, headers):
    for c, h in enumerate(headers, 1):
        x = ws.cell(row, c, h); x.fill = HDR; x.font = HF
        x.alignment = Alignment(horizontal="center", wrap_text=True)


# --------------------------------------------------------------------------- small stats helpers
def _pearson(xs, ys):
    """Pearson correlation of two equal-length lists; None if degenerate."""
    n = len(xs)
    if n < 3:
        return None
    mx = sum(xs) / n; my = sum(ys) / n
    sxx = sum((x - mx) ** 2 for x in xs); syy = sum((y - my) ** 2 for y in ys)
    sxy = sum((x - mx) * (y - my) for x, y in zip(xs, ys))
    d = (sxx * syy) ** 0.5
    return (sxy / d) if d > 1e-12 else None


def _inv(M):
    """Inverse of an n x n matrix via Gauss-Jordan with partial pivot; None if singular."""
    n = len(M)
    A = [list(M[i]) + [1.0 if i == j else 0.0 for j in range(n)] for i in range(n)]
    for c in range(n):
        piv = max(range(c, n), key=lambda r: abs(A[r][c]))
        if abs(A[piv][c]) < 1e-12:
            return None
        A[c], A[piv] = A[piv], A[c]
        pv = A[c][c]; A[c] = [v / pv for v in A[c]]
        for r in range(n):
            if r != c and A[r][c]:
                f = A[r][c]; A[r] = [A[r][k] - f * A[c][k] for k in range(2 * n)]
    return [row[n:] for row in A]


def _fm_stat(series):
    """Fama-MacBeth aggregate of a monthly slope series: (mean, t-stat, n). The t-stat divides the mean
    by the standard error of the month-to-month slopes -- so a slope that jumps around across months
    (noise / instability) earns a low t, which is the whole point of the reliability read."""
    s = [x for x in series if x is not None]
    n = len(s)
    if n < 2:
        return (None, None, n)
    m = sum(s) / n
    sd = (sum((x - m) ** 2 for x in s) / (n - 1)) ** 0.5
    t = (m / (sd / n ** 0.5)) if sd > 1e-12 else None
    return (m, t, n)


# --------------------------------------------------------------------------- factor cross-section + coefficients
def factor_audit(wb):
    """Emit the monthly factor cross-section and fitted coefficients, reproducing the pipeline's
    regression with the SAME building blocks; add a sample-month LINEST reproduction to the workbook."""
    try:
        import r2k_factor_analysis as fa
    except Exception as e:
        print(f"  (factor audit skipped -- cannot import r2k_factor_analysis: {e})")
        return
    try:
        months, ret, t2c, c2t = fa.load_returns()
        mcap = fa.load_mktcap(months)
        fund = fa.load_fundamentals()
    except FileNotFoundError as e:
        print(f"  (factor audit skipped -- input file not found: {e})")
        return
    FF = fa.FACTORS
    midx = {m: i for i, m in enumerate(months)}
    xs_rows = []                        # cross-section: one row per (index, month, name)
    coef_rows = []                      # coefficients: one row per (index, month)
    sample = None                       # (label, ym, header, matrix, pyslopes) for the in-sheet LINEST

    for pat, label in [("*Russell*Growth*Holdings*.xlsx", "R2000G"), ("*600*Growth*Holdings*.xlsx", "S&P600G")]:
        univ, sect = fa.load_holdings(pat)
        hmonths = sorted(univ)

        def snap(ym):
            prior = [h for h in hmonths if h <= ym]
            return prior[-1] if prior else None

        for ym in months:
            if int(ym[:4]) < fa.Y0:
                continue
            mi = midx.get(ym); s = snap(ym)
            if mi is None or mi + 1 >= len(months) or s is None:
                continue
            univ_ym, sect_ym = univ[s], sect[s]
            names = [c for c in univ_ym if (mi + 1) in ret.get(c, {})]
            if len(names) < fa.MIN_NAMES:
                continue
            w = {c: univ_ym[c] for c in names}; sec = {c: sect_ym.get(c, "?") for c in names}
            raw = fa.build_factor_scores(names, mi, ym, ret, mcap, fund)
            zsn = {ff: fa.zscore(raw[ff], sec) for ff in FF}      # sector-neutral z (efficacy)
            zr = {ff: fa.zscore(raw[ff]) for ff in FF}            # raw z (attribution)
            fwd = {c: ret[c][mi + 1] for c in names}
            for c in names:
                xs_rows.append(
                    [label, ym, c, c2t.get(c, ""), sec[c], round(w[c], 6)]
                    + [round(raw[ff][c], 6) if raw[ff].get(c) is not None else "" for ff in FF]
                    + [round(zr[ff][c], 4) if zr[ff].get(c) is not None else "" for ff in FF]
                    + [round(zsn[ff][c], 4) if zsn[ff].get(c) is not None else "" for ff in FF]
                    + [round(fwd[c], 6)])
            # --- reproduce the RAW Fama-MacBeth regression, exposing slope / exposure / contribution ---
            common = [c for c in names if all(c in zr[ff] for ff in FF)]
            if len(common) < fa.MIN_NAMES:
                continue
            tw = sum(w[c] for c in common) or 1.0
            ir = sum(w[c] * fwd[c] for c in common) / tw
            X = [[1.0] + [zr[ff][c] for ff in FF] for c in common]
            b = fa.ols(X, [fwd[c] for c in common])
            if not b:
                continue
            expo = {ff: sum(w[c] * zr[ff][c] for c in common) / tw for ff in FF}
            contrib = {ff: b[i + 1] * expo[ff] for i, ff in enumerate(FF)}
            market = b[0]; resid = ir - market - sum(contrib.values())
            # sector-adjusted contributions (styles controlled for GICS sector) + sector bucket
            dadj = fa._fm_decompose(common, w, zr, fwd, sect=sec)
            coef_rows.append(
                [label, ym, len(common), round(market, 8)]
                + [round(b[i + 1], 8) for i in range(len(FF))]           # raw slopes
                + [round(expo[ff], 8) for ff in FF]                       # index exposures
                + [round(contrib[ff], 8) for ff in FF]                    # raw contributions
                + [round(resid, 8), round(ir, 8)]
                + ([round(dadj["market"], 8)] + [round(dadj["style"][ff], 8) for ff in FF]
                   + [round(dadj["sect_bucket"], 8), round(dadj["resid"], 8)] if dadj else [""] * (len(FF) + 3)))
            if label == "R2000G":     # keep the most recent R2000G month as the in-sheet sample
                sample = (label, ym, common[:], [[zr[ff][c] for ff in FF] for c in common],
                          [fwd[c] for c in common], [b[i + 1] for i in range(len(FF))], b[0])

    # ---- CSVs ----
    xs_head = (["index", "month", "cik", "ticker", "gics_sector", "index_weight"]
               + [f"raw_{ff}" for ff in FF] + [f"z_{ff}" for ff in FF]
               + [f"zneutral_{ff}" for ff in FF] + ["fwd_return_next_month"])
    with open(OUT_XS, "w", newline="", encoding="utf-8") as f:
        wtr = csv.writer(f); wtr.writerow(xs_head); wtr.writerows(xs_rows)
    print(f"  wrote {OUT_XS.name}  ({len(xs_rows)} name-months)")
    coef_head = (["index", "month", "n_names", "market_intercept"]
                 + [f"slope_{ff}" for ff in FF] + [f"exposure_{ff}" for ff in FF]
                 + [f"contribution_{ff}" for ff in FF] + ["residual", "index_return"]
                 + ["adj_market"] + [f"adj_contribution_{ff}" for ff in FF] + ["adj_sectors_bucket", "adj_residual"])
    with open(OUT_COEF, "w", newline="", encoding="utf-8") as f:
        wtr = csv.writer(f); wtr.writerow(coef_head); wtr.writerows(coef_rows)
    print(f"  wrote {OUT_COEF.name}  ({len(coef_rows)} index-months)")

    # ---- workbook: method + one month reproduced in-sheet with LINEST ----
    ws = wb.create_sheet("Regression (sample month)")
    ws.cell(1, 1, "Factor attribution -- one month reproduced from the raw cross-section").font = TITLE
    ws.cell(2, 1, "Each month the constituents' next-month return is regressed on their raw factor "
                  "z-scores (a multivariate OLS / Fama-MacBeth cross-section). A factor's contribution = "
                  "its slope x the index's weighted exposure; market (intercept) + factor contributions + "
                  "residual = the index return. The full monthly inputs and outputs are in "
                  "audit_factor_crosssection.csv and audit_factor_coefficients.csv; below, one month is "
                  "reproduced in-sheet so LINEST matches the coefficients file.").font = BODY
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[2].height = 60
    if not sample:
        ws.cell(4, 1, "(no sample month available -- factor inputs were absent)").font = BODY
        return
    label, ym, cnames, Z, Y, pyslopes, pyint = sample
    ws.cell(4, 1, f"Sample: {label}  month {ym}  ({len(cnames)} names)").font = H
    hrow = 6
    _hdr(ws, hrow, ["cik"] + [f"z_{ff}" for ff in FF] + ["fwd_return"])
    for i, (c, zrow, y) in enumerate(zip(cnames, Z, Y), hrow + 1):
        ws.cell(i, 1, c)
        for j, zv in enumerate(zrow, 2):
            ws.cell(i, j, round(zv, 4))
        ws.cell(i, 2 + len(FF), round(y, 6))
    last = hrow + len(cnames)
    zc0, zc1 = get_column_letter(2), get_column_letter(1 + len(FF))          # z columns span
    yc = get_column_letter(2 + len(FF))                                       # fwd-return column
    # LINEST reproduces the OLS slopes (returned in REVERSE variable order, intercept last)
    lr = last + 2
    ws.cell(lr, 1, "LINEST(fwd_return, z_1..z_6)  ->  slopes in REVERSE order, then intercept:").font = H
    ws.cell(lr + 1, 1, f"=LINEST({yc}{hrow+1}:{yc}{last},{zc0}{hrow+1}:{zc1}{last},TRUE,FALSE)")
    ws.cell(lr + 3, 1, "Python slopes from audit_factor_coefficients.csv (same month, natural order):").font = H
    _hdr(ws, lr + 4, list(FF) + ["intercept"])
    for j, sv in enumerate(pyslopes, 1):
        ws.cell(lr + 5, j, round(sv, 6))
    ws.cell(lr + 5, len(FF) + 1, round(pyint, 6))
    ws.cell(lr + 6, 1, "Note: LINEST lists coefficients right-to-left (z_6 first ... z_1), intercept last; "
                       "reverse it to line up with the natural order above.").font = Font(italic=True, size=9)
    ws.column_dimensions["A"].width = 16


# --------------------------------------------------------------------------- collinearity / stability
def factor_reliability(wb):
    """A. correlation matrix of the raw z regressors; B. VIF per factor (diag of inverse corr matrix);
    C. univariate vs multivariate Fama-MacBeth slope per factor with FM t-stats. Per index."""
    try:
        import r2k_factor_analysis as fa
    except Exception as e:
        print(f"  (factor reliability skipped -- cannot import r2k_factor_analysis: {e})")
        return
    try:
        months, ret, t2c, c2t = fa.load_returns()
        mcap = fa.load_mktcap(months)
        fund = fa.load_fundamentals()
    except FileNotFoundError as e:
        print(f"  (factor reliability skipped -- input file not found: {e})")
        return
    FF = fa.FACTORS; nF = len(FF)
    midx = {m: i for i, m in enumerate(months)}
    per_index = {}   # label -> dict(corr, vif, uni{ff:[]}, multi{ff:[]})

    for pat, label in [("*Russell*Growth*Holdings*.xlsx", "R2000G"), ("*600*Growth*Holdings*.xlsx", "S&P600G")]:
        univ, sect = fa.load_holdings(pat)
        hmonths = sorted(univ)

        def snap(ym):
            prior = [h for h in hmonths if h <= ym]
            return prior[-1] if prior else None

        corr_sum = [[0.0] * nF for _ in range(nF)]; corr_n = [[0] * nF for _ in range(nF)]
        uni = {ff: [] for ff in FF}; multi = {ff: [] for ff in FF}
        for ym in months:
            if int(ym[:4]) < fa.Y0:
                continue
            mi = midx.get(ym); s = snap(ym)
            if mi is None or mi + 1 >= len(months) or s is None:
                continue
            univ_ym = univ[s]
            names = [c for c in univ_ym if (mi + 1) in ret.get(c, {})]
            if len(names) < fa.MIN_NAMES:
                continue
            w = {c: univ_ym[c] for c in names}
            raw = fa.build_factor_scores(names, mi, ym, ret, mcap, fund)
            zr = {ff: fa.zscore(raw[ff]) for ff in FF}                 # raw z = the attribution regressors
            common = [c for c in names if all(c in zr[ff] for ff in FF)]
            if len(common) < fa.MIN_NAMES:
                continue
            fwd = {c: ret[c][mi + 1] for c in common}
            cols = {ff: [zr[ff][c] for c in common] for ff in FF}
            # A. accumulate this month's pairwise correlations
            for a in range(nF):
                for b in range(a, nF):
                    r = _pearson(cols[FF[a]], cols[FF[b]])
                    if r is not None:
                        corr_sum[a][b] += r; corr_n[a][b] += 1
                        if a != b:
                            corr_sum[b][a] += r; corr_n[b][a] += 1
            # C. multivariate slopes (all six + intercept) and univariate slopes (one factor + intercept)
            y = [fwd[c] for c in common]
            bm = fa.ols([[1.0] + [zr[ff][c] for ff in FF] for c in common], y)
            if bm:
                for i, ff in enumerate(FF):
                    multi[ff].append(bm[i + 1])
            for ff in FF:
                bu = fa.ols([[1.0, zr[ff][c]] for c in common], y)
                if bu:
                    uni[ff].append(bu[1])
        # average correlation matrix -> VIF from its inverse diagonal
        corr = [[(corr_sum[a][b] / corr_n[a][b]) if corr_n[a][b] else (1.0 if a == b else 0.0)
                 for b in range(nF)] for a in range(nF)]
        inv = _inv(corr)
        vif = [inv[i][i] if inv else None for i in range(nF)]
        per_index[label] = dict(corr=corr, vif=vif, uni=uni, multi=multi)

    if not per_index:
        print("  (factor reliability skipped -- no index-months produced)")
        return

    # ---- CSVs ----
    with open(OUT_CORR, "w", newline="", encoding="utf-8") as f:
        wtr = csv.writer(f); wtr.writerow(["index", "factor"] + list(FF) + ["VIF"])
        for label, d in per_index.items():
            for i, ff in enumerate(FF):
                wtr.writerow([label, ff] + [round(d["corr"][i][j], 4) for j in range(nF)]
                             + [round(d["vif"][i], 3) if d["vif"][i] is not None else ""])
    print(f"  wrote {OUT_CORR.name}")

    def _ann(m):     # monthly slope (return per +1 SD of z) -> annualized %
        return round(m * 12 * 100, 2) if m is not None else None

    uvm_rows = []
    for label, d in per_index.items():
        for ff in FF:
            um, ut, un = _fm_stat(d["uni"][ff])
            mm, mt, mn = _fm_stat(d["multi"][ff])
            dchg = (mm - um) if (um is not None and mm is not None) else None
            pct = (100 * abs(dchg) / abs(um)) if (dchg is not None and um not in (None, 0)) else None
            rel = (mt is not None and abs(mt) >= 2.0)
            stable = (pct is not None and pct <= 25)
            verdict = ("reliable (significant, stable across specs)" if rel and stable else
                       "significant but collinearity-sensitive" if rel and not stable else
                       "not individually significant")
            uvm_rows.append([label, ff, un, _ann(um), round(ut, 2) if ut is not None else None,
                             _ann(mm), round(mt, 2) if mt is not None else None,
                             _ann(dchg), round(pct, 0) if pct is not None else None, verdict])
    with open(OUT_UVM, "w", newline="", encoding="utf-8") as f:
        wtr = csv.writer(f)
        wtr.writerow(["index", "factor", "n_months", "univariate_slope_ann_pct", "univariate_t",
                      "multivariate_slope_ann_pct", "multivariate_t", "delta_ann_pct", "abs_pct_change", "verdict"])
        wtr.writerows(uvm_rows)
    print(f"  wrote {OUT_UVM.name}")

    # ---- workbook tab ----
    ws = wb.create_sheet("Factor Reliability")
    ws.cell(1, 1, "Factor reliability -- collinearity (correlation, VIF) and coefficient stability "
                  "(univariate vs multivariate)").font = TITLE
    ws.cell(2, 1, "The attribution is ONE multivariate regression per month, so correlated factors share "
                  "explanatory power and errors in one regressor can move the others' coefficients. This tab "
                  "sizes that: (A) how correlated the regressors are, (B) VIF = variance inflation, (C) whether "
                  "each factor's slope survives moving from a one-factor to the full six-factor regression. A "
                  "factor significant AND stable across the two specs is not an artifact of the joint estimation.").font = BODY
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[2].height = 74
    r = 4
    for label in per_index:
        d = per_index[label]
        ws.cell(r, 1, f"{label} — A. Average cross-sectional correlation of the raw factor z-scores "
                      "(|r|>0.5 = materially collinear)").font = H
        r += 1
        _hdr(ws, r, [""] + list(FF))
        for i, ff in enumerate(FF):
            ws.cell(r + 1 + i, 1, ff).font = Font(bold=True, size=10)
            for j in range(nF):
                cell = ws.cell(r + 1 + i, 2 + j, round(d["corr"][i][j], 2))
                if i != j and abs(d["corr"][i][j]) > 0.5:
                    cell.font = Font(bold=True, color="B00000")
        r += nF + 2
        ws.cell(r, 1, f"{label} — B. Variance Inflation Factor (VIF < 5 low · 5-10 moderate · > 10 high)").font = H
        r += 1
        _hdr(ws, r, list(FF))
        for i in range(nF):
            v = d["vif"][i]
            cell = ws.cell(r + 1, 1 + i, round(v, 2) if v is not None else "n/a")
            if v is not None and v > 5:
                cell.font = Font(bold=True, color="B00000")
        r += 3
        ws.cell(r, 1, f"{label} — C. Univariate vs multivariate Fama-MacBeth slope (annualized %, per +1 SD) "
                      "with FM t-stats").font = H
        r += 1
        _hdr(ws, r, ["Factor", "Univariate slope %", "Univ t", "Multivariate slope %", "Multi t",
                     "Delta slope (pts)", "|Delta| / univ %", "Verdict"])
        r += 1
        for row in [x for x in uvm_rows if x[0] == label]:
            _, ff, un, us, ut, ms, mt, dl, pc, verdict = row
            for c, v in enumerate([ff, us, ut, ms, mt, dl, pc, verdict], 1):
                cell = ws.cell(r, c, v)
                if c == 1:
                    cell.font = Font(bold=True, size=10)
                if c == 8 and verdict.startswith("reliable"):
                    cell.font = Font(bold=True, color="1F6E1F")
            r += 1
        r += 2
    ws.column_dimensions["A"].width = 24
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 17


def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    factor_audit(wb)          # -> crosssection + coefficients CSVs, and the sample-month tab
    factor_reliability(wb)    # -> correlation + uni-vs-multi CSVs, and the Factor Reliability tab
    if not wb.sheetnames:
        print("  !! no factor tabs produced -- check that the factor input files are present.")
        return
    out = BASE / "R2000G_Factor_Granular.xlsx"
    wb.save(out)
    print(f"  wrote {out.name}  ({len(wb.sheetnames)} tabs) + the four audit_factor_*.csv files")


if __name__ == "__main__":
    main()
