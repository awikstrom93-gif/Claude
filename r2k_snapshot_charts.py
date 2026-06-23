"""
r2k_snapshot_charts.py  --  save a durable backup of the charts in the consolidated
workbook. Run this after you've added/edited charts (e.g. with the Claude Excel
plugin). r2k_step7_consolidate.py preserves charts from the live output on every
re-run; this backup is the fallback used if that output file is ever lost.

RUN: python r2k_snapshot_charts.py
"""
from pathlib import Path
import os, shutil
import openpyxl

BASE = Path(os.environ.get("R2KG_BASE", "."))
OUT = BASE / "R2000G_SmallCapGrowth_Benchmark_Review.xlsx"
BACKUP = BASE / os.environ.get("R2KG_CHART_BACKUP", "R2000G_charts_backup.xlsx")


def main():
    if not OUT.exists():
        raise SystemExit(f"!! {OUT.name} not found -- nothing to back up.")
    shutil.copy(OUT, BACKUP)
    n = sum(len(getattr(openpyxl.load_workbook(BACKUP)[s], "_charts", []))
            for s in openpyxl.load_workbook(BACKUP).sheetnames)
    print(f"  backed up {OUT.name} -> {BACKUP.name}  ({n} charts snapshotted)")
    print("  step 7 will restore from this if the live workbook is ever missing.")


if __name__ == "__main__":
    main()
