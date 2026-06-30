"""
r2k_view_valuation.py  --  what did the market PAY for R2000G's unprofitable tail?  The valuation
dimension the book was missing -- derived with NO new data by using index weight as a float-market-cap
proxy (Russell weights are float-cap weights).

Math: weight_i = mktcap_i / TotalCap, so a cohort's aggregate P/S = Sum(mktcap)/Sum(sales) =
TotalCap * Sum(weight)/Sum(sales). Expressed as a MULTIPLE of the index's own P/S, the unknown
TotalCap cancels exactly -- so 'unprofitable names trade at X.X times the index P/S' is exact and
comparable across years. Same for P/B (book). (P/E is undefined for the unprofitable tail.)

Caveats (in the tab): weight is FLOAT-cap (float adjustment adds mild noise, ~washes in aggregates);
multiples are RELATIVE TO THIS INDEX (not absolute levels / cross-index without an external anchor);
revenue & book are as-filed fiscal-year matched to the April snapshot.

Pure panel projection.  RUN:  python r2k_view_valuation.py
"""
from r2k_universe import get_panel, by_index_year, BASE

OUT = BASE / "r2k_valuation.txt"
HDR = ["Year", "%NoRev wt", "Prof P/S xIdx", "Unprof P/S xIdx", "Never P/S xIdx",
       "Unprof/Prof P/S", "Prof P/B xIdx", "Unprof P/B xIdx", "Unprof/Prof P/B"]


def _agg_ps(rows, vf):
    """Sum(weight)/Sum(value) over rows with value>0 -- the cohort's P/S (or P/B) in TotalCap units."""
    sw = sum(r["weight"] for r in rows if r.get(vf) and r[vf] > 0)
    sv = sum(r[vf] for r in rows if r.get(vf) and r[vf] > 0)
    return (sw / sv) if sv > 0 else None


def _mult(rows_c, rows_all, vf):
    """Cohort P/S (or P/B) as a MULTIPLE of the index's -- TotalCap cancels, so this is exact."""
    uc, ua = _agg_ps(rows_c, vf), _agg_ps(rows_all, vf)
    return (uc / ua) if (uc and ua) else None


def _ratio(rows_a, rows_b, vf):
    """Cohort A's P/S relative to cohort B's (e.g. unprofitable vs profitable). TotalCap cancels."""
    ua, ub = _agg_ps(rows_a, vf), _agg_ps(rows_b, vf)
    return (ua / ub) if (ua and ub) else None


def valuation_rows(panel):
    grp = by_index_year(panel)
    years = sorted(y for (ix, y) in grp if ix == "R2KG")
    out = []
    for y in years:
        allr = [r for r in grp[("R2KG", y)] if r["covered"]]
        prof = [r for r in allr if r["cohort"] == "profitable"]
        unp = [r for r in allr if r["prof_ni"] is False]
        nev = [r for r in allr if r["cohort"] == "never_profitable"]
        tw = sum(r["weight"] for r in allr) or 1e-9
        norev = 100 * sum(r["weight"] for r in allr if not r["has_rev"]) / tw
        out.append([y, round(norev, 1),
                    round(_mult(prof, allr, "revenue") or 0, 2) or None,
                    round(_mult(unp, allr, "revenue") or 0, 2) or None,
                    round(_mult(nev, allr, "revenue") or 0, 2) or None,
                    round(_ratio(unp, prof, "revenue") or 0, 2) or None,
                    round(_mult(prof, allr, "equity") or 0, 2) or None,
                    round(_mult(unp, allr, "equity") or 0, 2) or None,
                    round(_ratio(unp, prof, "equity") or 0, 2) or None])
    return out


def write_sheet(wb, panel=None):
    from openpyxl.styles import Font, PatternFill, Alignment
    panel = panel if panel is not None else get_panel(index="R2KG")
    rows = valuation_rows(panel)
    ws = wb.create_sheet("Valuation of the Tail")
    ws.cell(1, 1, "What did the market pay for the unprofitable tail?  Cohort valuation vs the index").font = Font(bold=True, size=12)
    ws.cell(2, 1, "Index weight used as a float-market-cap proxy (Russell = float-cap weighted). Each "
                  "cohort's P/S and P/B shown as a MULTIPLE of the index's (the cap constant cancels, so "
                  "this is exact and comparable across years). >1 = richer than the index. P/E omitted "
                  "(undefined for loss-makers). Relative to THIS index, not an absolute level.").font = Font(size=9, italic=True, color="555555")
    fill = PatternFill("solid", fgColor="1F4E5F")
    for c, h in enumerate(HDR, 1):
        x = ws.cell(4, c, h); x.fill = fill; x.font = Font(bold=True, color="FFFFFF", size=10)
        x.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, row in enumerate(rows, start=5):
        for c, v in enumerate(row, 1):
            ws.cell(i, c, v)
    ws.freeze_panes = "B5"
    return "Valuation of the Tail"


def main():
    panel = get_panel(index="R2KG")
    rows = valuation_rows(panel)
    L = ["VALUATION OF THE TAIL  --  cohort P/S & P/B as a MULTIPLE of the R2000G index (weight=float-cap proxy)", ""]
    L.append("  " + "".join(str(h)[:15].rjust(17) for h in HDR))
    L.append("  " + "-" * (17 * len(HDR)))
    for row in rows:
        L.append("  " + "".join(("" if v is None else str(v)).rjust(17) for v in row))
    L.append("")
    L.append("  READ: 'Unprof P/S xIdx' >1 means unprofitable names traded RICHER than the index per dollar")
    L.append("  of sales. 'Unprof/Prof' is the tail's premium over profitable names. Watch the 2020-2021")
    L.append("  spike (the growth bubble paid up for sales-without-earnings) and the 2022 de-rate. P/E is")
    L.append("  omitted -- the unprofitable tail has no earnings to price. Multiples are RELATIVE to this")
    L.append("  index (float-cap proxy), not absolute levels.")
    OUT.write_text("\n".join(L), encoding="utf-8")
    print("\n".join(L))
    print(f"\n  -> {OUT.name}")


if __name__ == "__main__":
    main()
