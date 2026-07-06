"""
r2k_view_reliability.py  --  the IC-facing "Data Reliability" tab, extracted from r2k_step7_consolidate
into a standalone, testable module (so step7 and the future r2k_report.py both call it instead of it
living inside the monolith).  One row per company-year (full historical ledger): reliability tier,
three-statement tie-out confidence, which identities break, plausibility flags, and the provenance of
any engineered values.  PREDECESSOR-entity years (reverse mergers) are excluded.

Inputs (graceful-degrade if absent): plausibility_flags.csv [r2k_plausibility], fundamentals_dera.csv
[r2k_dera_classify], plus index weights/tickers.  Logic is a faithful copy of the step7 tab.

RUN:  python r2k_view_reliability.py        # prints the headline + ledger summary for validation
"""
import csv as _csv
import json as _json
from collections import Counter

from r2k_universe import BASE

FLAGS = BASE / "plausibility_flags.csv"
FUND = BASE / "fundamentals_dera.csv"
CMAP = BASE / "security_cik_map.json"

HEADERS = ["Ticker", "CIK", "Index Wt % (that yr)", "Fiscal Year", "Core IS+BS", "Full (all 3)",
           "Tie-out conf", "Identity breaks", "Plausibility flags", "Provenance (engineered values)"]


def _short(s, n=180):
    return s if len(s) <= n else s[:n - 1] + "…"


def _ck(c):
    """Canonical int-string CIK so panel / flags / fundamentals keys always match."""
    s = str(c)
    return str(int(s)) if s.isdigit() else s


def _panel_weights():
    """From the panel: per-year R2000G weight per (cik, fiscal_year), current weight per cik, and
    a ticker per cik. Per-year weight is the index weight at the April snapshot that USED that 10-K
    (earliest such snapshot when the data was fresh); 0 means the name was not an R2000G constituent
    that year. Returns ({}, {}, {}) if the panel can't be built."""
    wt_by_cikfy, current_wt, tkr = {}, {}, {}
    try:
        from r2k_universe import get_panel
        panel = get_panel(index="R2KG")
    except Exception:
        return wt_by_cikfy, current_wt, tkr
    if not panel:
        return wt_by_cikfy, current_wt, tkr
    max_year = max(int(r["year"]) for r in panel)
    for r in sorted(panel, key=lambda x: int(x["year"])):   # ascending -> earliest snapshot wins
        c = _ck(r["cik"]) if r["cik"] else None
        if not c:
            continue
        tkr.setdefault(c, r.get("ticker") or "")
        if int(r["year"]) == max_year:
            current_wt[c] = r["weight"]
        if r["covered"] and r.get("fy0") is not None:
            wt_by_cikfy.setdefault((c, str(r["fy0"])), r["weight"])
    return wt_by_cikfy, current_wt, tkr


def reliability_data():
    """Returns (rows, meta) or (None, None) if the flags file is absent. Each row carries the name's
    ACTUAL R2000G index weight in that fiscal year (from the panel); the headline core/full figures
    are measured on the CURRENT snapshot weight. rows are sorted by that-year weight desc."""
    if not FLAGS.exists():
        return None, None
    flags = [r for r in _csv.DictReader(open(FLAGS, encoding="utf-8")) if r.get("fiscal_year", "").isdigit()]

    latest = {}                                   # latest-year flag row per cik (for the headline tier)
    for r in flags:
        c = _ck(r.get("cik"))
        if c not in latest or r["fiscal_year"] > latest[c]["fiscal_year"]:
            latest[c] = r

    prov = {}                                     # (cik, fy) -> fundamentals row (provenance/breaks/entity_flag)
    if FUND.exists():
        for r in _csv.DictReader(open(FUND, encoding="utf-8")):
            if r.get("fiscal_year", "").isdigit():
                prov[(_ck(r.get("cik")), r["fiscal_year"])] = r

    wt_by_cikfy, current_wt, tkr = _panel_weights()
    cik2tkr = dict(tkr)
    if CMAP.exists():                             # ticker fallback for non-constituent names in the ledger
        try:
            for tk, v in _json.load(open(CMAP)).items():
                c = v.get("cik") if isinstance(v, dict) else v
                cik2tkr.setdefault(_ck(c), tk)
        except Exception:
            pass

    rows, pred_excluded = [], 0
    for fr in flags:
        c, fy = _ck(fr.get("cik")), fr["fiscal_year"]
        pr = prov.get((c, fy), {})
        if pr.get("entity_flag") == "PREDECESSOR":
            pred_excluded += 1
            continue
        why = fr.get("watch_reason", "") if fr.get("tier") == "watch" else ""
        flg = ";".join(x for x in (fr.get("critical", ""), fr.get("watch", "")) if x)
        rows.append({
            "tkr": cik2tkr.get(c, ""), "cik": fr.get("cik"), "wt": wt_by_cikfy.get((c, fy), 0.0), "fy": fy,
            "tier": fr.get("tier", ""), "core": fr.get("core_reliable", ""),
            "conf": fr.get("confidence", "") or pr.get("confidence", ""),
            "breaks": pr.get("breaks", ""),
            "flags": (flg + (f" ({why})" if why and not flg else (f" — {why}" if why else ""))).strip(),
            "prov": _short(pr.get("provenance", "")),
        })
    rows.sort(key=lambda x: (-x["wt"], x["cik"], x["fy"]))

    tw = sum(current_wt.values()) or 0.0
    clean_w = sum(current_wt.get(c, 0.0) for c, fr in latest.items() if fr.get("tier") == "clean")
    core_w = sum(current_wt.get(c, 0.0) for c, fr in latest.items() if fr.get("tier") != "review")
    meta = {"tw": tw, "clean_w": clean_w, "core_w": core_w, "pred_excluded": pred_excluded,
            "tiers": Counter(r["tier"] for r in rows), "n_names": len({r["cik"] for r in rows})}
    return rows, meta


def write_sheet(wb):
    """Render the Data Reliability sheet into an openpyxl workbook. Returns the sheet name or None."""
    from openpyxl.styles import Font, PatternFill, Alignment
    rows, meta = reliability_data()
    if rows is None:
        return None
    TITLE = Font(bold=True, size=14, color="1F4E5F"); BODY = Font(size=10)
    HF = Font(bold=True, color="FFFFFF", size=10); HDR = PatternFill("solid", fgColor="1F4E5F")
    tw, clean_w, core_w = meta["tw"], meta["clean_w"], meta["core_w"]
    tc = meta["tiers"]

    ws = wb.create_sheet("Data Reliability")
    ws.cell(1, 1, "Data Reliability — three-statement tie-out, sanity, and provenance").font = TITLE
    sub = (f"{len(rows):,} company-years across {meta['n_names']:,} names (full historical ledger, "
           f"2015–present). "
           + (f"CORE reliability (income statement + balance sheet tie out AND are plausible — the basis "
              f"of the quality/growth/leverage analytics): {100*core_w/tw:.1f}% of current index weight.  "
              f"Full three-statement articulation (adds the cash-flow roll-forward): {100*clean_w/tw:.1f}%."
              if tw else "(index weights unavailable — name-count view.)"))
    if meta["pred_excluded"]:
        sub += f"  ({meta['pred_excluded']:,} predecessor-entity years excluded — reverse mergers.)"
    ws.cell(2, 1, sub).font = BODY
    ws.cell(3, 1, f"company-years:  clean {tc['clean']:,}   |   watch {tc['watch']:,}   |   review {tc['review']:,}"
                  "      clean = all three statements tie · watch = IS & BS tie; a cash-flow gap or a "
                  "growth-typical value · review = a real concern.   Index Wt % = the name's R2000G "
                  "weight in that fiscal year (blank = not an R2000G constituent then).").font = Font(size=9, italic=True, color="555555")
    # staleness tripwire: tiers come from plausibility (FLAGS); if they predate the last accounting-engine
    # run (FUND), they don't reflect the current classify -- surface it instead of shipping stale silently.
    if FUND.exists() and FLAGS.stat().st_mtime < FUND.stat().st_mtime - 1:
        print(f"  !! Data Reliability is STALE: {FLAGS.name} is older than {FUND.name} -- "
              f"re-run r2k_plausibility.py before r2k_report.py so the tiers reflect the latest engine.")
        ws.cell(4, 1, "⚠ STALE: these tiers were computed by an earlier plausibility run than the current "
                      "fundamentals — re-run r2k_plausibility.py, then r2k_report.py, to refresh them."
                ).font = Font(bold=True, color="C00000", size=10)
    for c, h in enumerate(HEADERS, 1):
        x = ws.cell(5, c, h); x.fill = HDR; x.font = HF
        x.alignment = Alignment(horizontal="center", wrap_text=True)
    tier_fill = {"clean": PatternFill("solid", fgColor="E2EFDA"), "watch": PatternFill("solid", fgColor="FFF2CC"),
                 "review": PatternFill("solid", fgColor="FCE4D6")}
    green = PatternFill("solid", fgColor="E2EFDA"); red = PatternFill("solid", fgColor="FCE4D6")
    for i, r in enumerate(rows, start=6):
        ws.cell(i, 1, r["tkr"]); ws.cell(i, 2, r["cik"])
        ws.cell(i, 3, round(r["wt"], 3) if r["wt"] else None); ws.cell(i, 4, r["fy"])
        core_ok = (r.get("core") or ("Y" if r["tier"] != "review" else "N")) == "Y"
        cc = ws.cell(i, 5, "reliable" if core_ok else "review"); cc.fill = green if core_ok else red
        cell = ws.cell(i, 6, r["tier"]); cell.fill = tier_fill.get(r["tier"], PatternFill())
        ws.cell(i, 7, r["conf"]); ws.cell(i, 8, r["breaks"]); ws.cell(i, 9, r["flags"]); ws.cell(i, 10, r["prov"])
        for col in range(1, 11):
            ws.cell(i, col).font = BODY
    for col, w in enumerate([10, 12, 11, 11, 11, 12, 11, 22, 30, 56], 1):
        ws.column_dimensions[chr(64 + col)].width = w
    ws.auto_filter.ref = f"A5:J{5 + len(rows)}"
    ws.freeze_panes = "A6"
    return "Data Reliability"


def main():
    rows, meta = reliability_data()
    if rows is None:
        print(f"  !! {FLAGS.name} not found -- run r2k_plausibility.py first.")
        return
    tw = meta["tw"]
    print("\n  DATA RELIABILITY (panel-era extract) -- compare to the current tab's header:")
    print(f"    company-years : {len(rows):,}   names: {meta['n_names']:,}   "
          f"predecessor-excluded: {meta['pred_excluded']:,}")
    if tw:
        print(f"    CORE reliable : {100*meta['core_w']/tw:.1f}% of current index weight")
        print(f"    FULL clean    : {100*meta['clean_w']/tw:.1f}% of current index weight")
    tc = meta["tiers"]
    print(f"    tiers         : clean {tc['clean']:,} | watch {tc['watch']:,} | review {tc['review']:,}")
    print("\n    top 8 by weight:")
    print(f"      {'tkr':<7}{'cik':<10}{'wt%':>7}{'fy':>6}  {'core':<9}{'full':<7}flags")
    for r in rows[:8]:
        core_ok = (r.get("core") or ("Y" if r["tier"] != "review" else "N")) == "Y"
        print(f"      {r['tkr']:<7}{r['cik']:<10}{(r['wt'] or 0):>7.3f}{r['fy']:>6}  "
              f"{'reliable' if core_ok else 'review':<9}{r['tier']:<7}{r['flags'][:40]}")


if __name__ == "__main__":
    main()
