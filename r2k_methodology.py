"""
============================================================
r2k_methodology.py  --  generate METHODOLOGY.md, a concise standalone reader's guide to the
benchmark-review workbook: WHAT each tab shows, HOW the data was built, and the conventions/caveats
behind every number -- so anyone can talk through the analysis without opening 42 tabs.
============================================================
The "how it was done" narrative (data foundation, pipeline, conventions, key adjustments) is stable
process knowledge and lives here as static prose. The tab-by-tab guide and the glossary are pulled
LIVE from the workbook's own `Reading Guide` and `Glossary` tabs, so they can never drift from what
the workbook actually contains -- and a coverage check warns if any tab is documented nowhere.

Deliberately FIGURE-LIGHT: results live in the memo (r2k_memo.py) and the workbook. This document is
about method, so re-running it after a data refresh changes only the tab list, never a headline
number that could fall out of sync.

RUN:      python r2k_methodology.py            (writes METHODOLOGY.md next to the workbook)
SELFTEST: python r2k_methodology.py --selftest (parse + coverage check, write nothing)
Override with R2KG_WORKBOOK / R2KG_BASE / R2KG_METHODOLOGY_OUT.
============================================================
"""
from pathlib import Path
import os, sys

def _esc(s):
    """Escape pipes so workbook text (e.g. 'R2000G | 600G | Diff') can't break a markdown table."""
    return str(s).replace("|", "\\|").strip()


BASE = Path(os.environ.get("R2KG_BASE", "."))
WB = BASE / os.environ.get("R2KG_WORKBOOK", "R2000G_SmallCapGrowth_Benchmark_Review.xlsx")
OUT = BASE / os.environ.get("R2KG_METHODOLOGY_OUT", "METHODOLOGY.md")

# Front-matter tabs the workbook carries but the Reading Guide doesn't describe. One-liners live here;
# the builder verifies each named tab is actually present and warns on any workbook tab covered neither
# here nor in the Reading Guide. (The panel-native exhibit tabs -- Quality Factor Spreads, Solvency
# Tail, Cohort Persistence, Valuation of the Tail -- are documented in the Reading Guide itself, so they
# are intentionally NOT repeated here.)
SUPPORT_TABS = {
    "Executive Summary": "The one-page version of the whole review: the question, the three-part answer, and the bottom line, each pointing to the tabs that prove it.",
    "Reading Guide": "What each analytical tab shows and how to read it (the source of the tab guide below).",
    "Glossary": "Plain-language definition of every metric and convention used in the workbook.",
    "Contents": "The tab index, grouped by section.",
    "Key Charts": "The exhibit charts (Growth of $1, % unprofitable by weight, the earnings-screen counterfactual) in one place.",
    "Data Reliability": "How trustworthy the reconstructed fundamentals are, by weight: the share of index weight whose three statements tie out cleanly vs. flagged for review, so a reader can weight the conclusions accordingly.",
}

# The methodology narrative -- static process knowledge. (Markdown; headings are '## ...'.)
NARRATIVE = r"""
## What this review answers

Active US Small-Cap Growth managers, benchmarked to the **Russell 2000 Growth (R2000G)**, lagged that
index over the recent manager window. This review tests one explanation: that the shortfall is
**structural** — a property of how the benchmark is built — rather than a loss of manager skill. It
does so by comparing R2000G against the **S&P SmallCap 600 Growth (S&P600G)**, whose one decisive
difference is an **earnings screen** (a company needs positive trailing GAAP earnings to enter the
S&P 600). A quality- or earnings-disciplined manager's portfolio tends to resemble the S&P600G, so the
gap between the two indices is a clean proxy for the cost — or benefit — of that discipline.

The argument runs in three moves, each with its own tabs: **(1)** R2000G carries a much larger tail of
unprofitable, often pre-revenue companies; **(2)** that tail *led* the benchmark over the exact window
managers were judged on, so avoiding it mechanically caused underperformance; **(3)** over the full
cycle the same discipline *won* — more return at lower risk. The memo (`R2000G_SCG_Benchmark_Review_Memo.md`)
states the current figures; this document explains how each was produced.

## How the fundamentals were built (the data foundation)

Every fundamental in this review is reconstructed from the SEC's own filings, not bought from a vendor
feed, so each number is traceable to an as-filed 10-K.

- **Source: SEC DERA / XBRL Financial Statement Data Sets.** Each quarter's structured facts from every
  10-K, 20-F and 40-F are indexed and extracted, then a classification engine reconstructs all three
  statements (income statement, balance sheet, cash flow) for each company-year.
- **As-filed, by original accession.** Values are taken as *originally* reported in each 10-K — no
  restatement or vintage blending. What the market saw at the time is what the analysis uses.
- **Point-in-time, no look-ahead.** A constituent's fiscal year at any snapshot is the **latest 10-K
  filed before that snapshot** — never financials that were not yet public. Index membership is
  historical (survivorship-free): a name is in a year only if it was actually in the index then.
- **Identity-checked reconstruction.** The engine doesn't just copy tags; it enforces accounting
  identities (the income statement foots to net income, the balance sheet foots, cash-flow D&A ties to
  the income statement) and only accepts a value when the statement ties out. Blanks left by
  non-standard XBRL tags are recovered **identity-first** (reconstruct from what must be true), then
  from validated as-filed tags — a closed loop that shrinks the gap at the source rather than plugging
  numbers.
- **Deliberate, documented adjustments** for cases where the raw filing would mislead an index
  aggregate. The clearest example: a **commodity/securities broker-dealer** (e.g. StoneX) reports
  "Revenues" grossed up by pass-through physical-commodity sales — tens of billions that are not
  comparable to an operating company's revenue and would swamp any index revenue or margin total. Such
  a filer is carried on a **net operating-revenue** basis (revenue net of the pass-through cost), the
  same convention banks and insurers already use, and only where a strict matched-book signature holds
  (tiny gross margin *and* tiny net margin), so genuine low-margin operating companies are untouched.
- **Consolidated registrant only.** Filers with public debt file Rule 3-10 *guarantor consolidating
  schedules* — a parent-only and a subsidiary column carrying the same XBRL tags as the consolidated
  company. Only the consolidated registrant's facts (no co-registrant, no segment dimension) are kept,
  so a parent holdco's tiny stand-alone figure can never overwrite the consolidated total.
- **Reliability is measured, not assumed.** The `Data Reliability` tab reports, by index weight, how
  much of the reconstruction ties out cleanly versus is flagged for review, so the reader can weight
  the conclusions. Headline results are dominated by the high-confidence core.

## Revenue capture — what "no revenue" means, and one bounded limitation

The top line is captured from the as-filed consolidated figure under a broad tag set — the standard
`Revenues` / `RevenueFromContractWithCustomer` concepts plus industry-specific operating lines
(homebuilding, hospital patient-service, marine, mining, fitness, franchisor) — and regardless of
whether the filer presents its income statement as a standalone statement, a **combined statement of
operations and comprehensive income**, or an **uncategorized presentation**. Two deliberate rules
shape what a *blank* top line means:

- **Financial-sector filers are not "no-revenue."** Banks, insurers, mortgage REITs, asset managers and
  BDCs report a net-interest / premium / fee top line, not a `Revenue` tag, so a blank revenue is
  *definitional*, not a sign of a pre-commercial company. These names are carried without a revenue
  figure and are **excluded from the "no-revenue" cohort** — otherwise a bank would masquerade as a
  clinical-stage biotech. The no-revenue weight therefore reads as *genuinely pre-commercial* names,
  which are overwhelmingly biotech (the driver of the 2026 step-up).
- **Bounded limitation — segment-only revenue totals.** A small number of filers tag their consolidated
  revenue total *only* with a business-segment dimension and never file an undimensioned company-level
  figure (e.g. **M.D.C. Holdings 2012–2017**, **Meritage Homes 2018+**, and a few others). Because the
  pipeline adopts only the as-filed consolidated value and does **not** reconstruct a total by summing
  segment members — which would risk double-counting nested sub-tiers or omitting inter-segment
  eliminations — these company-years are left blank rather than filled with an inferred number. The
  effect is bounded and immaterial: it is **under ~1 percentage point of index weight in any year**, is
  confined to historical years (mostly names no longer in the index), and does **not** touch the current
  2026 reading. This is a conscious accuracy-over-coverage choice: a blank is more honest than a
  reconstructed total that could be wrong.

## How the analysis was done (conventions)

- **Two aggregation lenses, stated explicitly.** Index-level figures use the **dollar-aggregate**
  convention (sum of numerators / sum of denominators — the index treated as one big company), which is
  the index-representative measure and is validated against FactSet. Per-name distribution views also
  show weight-weighted-average and median, because tiny-revenue loss-makers distort a simple average.
- **Each company counted once in dollar totals.** A name can sit in the index under two share classes
  (e.g. `CENTA`/`CENT`) or appear twice in a holdings file; both rows carry the same company's identical
  fundamentals. Dollar-level totals (index revenue, net income, the $-aggregate margins) **de-duplicate
  by company** so a dual-listed name's financials are not double-counted, while the weight-based quality
  percentages keep every row (the split weights correctly sum to the company's true index weight).
- **Profitability cohorts are point-in-time labels** from each name's as-filed net-income history:
  *Profitable* (net income > 0 in the latest filed year), *Fallen* (was profitable, now not),
  *Never-profitable* (no profitable year on record), *Unknown* (net income not reported). The
  *never-profitable* weight is the cleanest expression of the earnings-screen gap.
- **Attribution is Carino-linked** so that single-period cohort contributions sum *exactly* to the
  index's multi-period cumulative return — cohort contributions add up to the whole, with a small
  explicitly-labelled reconstruction residual (coverage + weight drift between snapshots).
- **The counterfactual is the cleanest test.** Rather than compare two different indices, it rebuilds
  R2000G's **own** constituents as a "profitable-only" (or ex-biotech) portfolio, reweighted monthly.
  The gap between the real index path and the screened path *is* the realized cost (or benefit) of the
  screen, holding the universe fixed. Its full-period result independently lands near the actual
  S&P600G return — two constructions agreeing that the earnings screen is the mechanism.
- **The manager window is data-driven, not chosen.** The `Perf Window Proof` tab finds when R2000G's
  cumulative excess over S&P600G troughed and began a persistent run, and shows a table of candidate
  windows so the conclusion doesn't hinge on the exact start month.
- **Ratios use average (opening + closing) denominators** (CFA convention); per-name ratios are
  winsorized before any averaging.
""".strip("\n")

NARRATIVE_ARC = r"""
## Talking through it in one minute

1. **Start with `Qual Comparison` (the "why").** One rule — the S&P 600 earnings screen — creates a
   persistent gap: R2000G runs many more points of unprofitable and never-profitable weight every year.
   `Bio Weight & Quality` shows biotech is the embodiment of that gap.
2. **Move to `Attr Contribution` + `Attr Counterfactual` (the "cost").** Decompose R2000G's return by
   quality cohort: over the manager window the unprofitable tail led. The counterfactual rebuilds
   R2000G's own names profitable-only — the sign flips vs. the full cycle, and that flip *is* the
   manager's shortfall.
3. **Close with `Perf Summary` + `Perf Window Proof` (the "context").** Over the full cycle the
   screened index delivered more return at lower risk; the recent window is the cost of the discipline
   during a low-quality rally, and the window dates are justified from the data, not picked.

**Bottom line:** the underperformance is a benchmark-construction effect, not lost skill — which is why
the S&P SmallCap 600 Growth is often the more representative yardstick for a quality-disciplined mandate.
""".strip("\n")


# --------------------------------------------------------------------------- workbook parsing
def _reading_guide(wb):
    """Parse the Reading Guide into [(section, [(tab, what, how), ...]), ...] preserving order.
    A section-header row has text in col A but no col-B description; a tab row has both."""
    ws = wb["Reading Guide"]
    sections, cur = [], None
    for r in ws.iter_rows(min_row=1, values_only=True):
        a = str(r[0]).strip() if r and r[0] is not None else ""
        b = str(r[1]).strip() if (r and len(r) > 1 and r[1] is not None) else ""
        c = str(r[2]).strip() if (r and len(r) > 2 and r[2] is not None) else ""
        if not a or a.lower().startswith("reading guide") or a == "Tab":
            continue
        if b:                                  # tab row (name + description)
            if cur is None:
                cur = ("", [])
                sections.append(cur)
            cur[1].append((a, b, c))
        else:                                  # section header (name only)
            cur = (a, [])
            sections.append(cur)
    return [(s, tabs) for s, tabs in sections if tabs]


def _glossary(wb):
    ws = wb["Glossary"]
    out = []
    for r in ws.iter_rows(min_row=1, values_only=True):
        a = str(r[0]).strip() if r and r[0] is not None else ""
        b = str(r[1]).strip() if (r and len(r) > 1 and r[1] is not None) else ""
        c = str(r[2]).strip() if (r and len(r) > 2 and r[2] is not None) else ""
        if not a or a.lower().startswith("glossary") or a == "Metric":
            continue
        if b:
            out.append((a, b, c))
    return out


def coverage(wb, sections):
    """Every workbook tab should be documented either in the Reading Guide or in SUPPORT_TABS.
    Returns (uncovered_tabs, stale_support) so the run can warn without failing."""
    guided = {t for _, tabs in sections for t, _, _ in tabs}
    present = set(wb.sheetnames)
    uncovered = [s for s in wb.sheetnames if s not in guided and s not in SUPPORT_TABS]
    stale = [s for s in SUPPORT_TABS if s not in present]
    return uncovered, stale


# --------------------------------------------------------------------------- build
def build(wb):
    sections = _reading_guide(wb)
    gloss = _glossary(wb)
    uncovered, stale = coverage(wb, sections)
    n_tabs = len(wb.sheetnames)

    P = []
    P.append("# Methodology & Reader's Guide")
    P.append("### Russell 2000 Growth vs. S&P SmallCap 600 Growth — benchmark review\n")
    P.append(f"*Companion to `{WB.name}` ({n_tabs} tabs) and `R2000G_SCG_Benchmark_Review_Memo.md`. "
             f"This guide is generated from the workbook's own Reading Guide and Glossary, so the tab "
             f"list and definitions below always match the workbook; the method narrative is stable "
             f"process documentation.*\n")
    P.append("---\n")
    P.append(NARRATIVE + "\n")
    P.append("---\n")
    P.append(NARRATIVE_ARC + "\n")
    P.append("---\n")

    # tab-by-tab guide, grouped by the Reading Guide's own sections
    P.append("## What each tab shows\n")
    for section, tabs in sections:
        if section:
            P.append(f"### {section}\n")
        P.append("| Tab | What it shows | How to read it |")
        P.append("|---|---|---|")
        for tab, what, how in tabs:
            P.append(f"| **{_esc(tab)}** | {_esc(what)} | {_esc(how)} |")
        P.append("")

    # supporting / front-matter / exhibit tabs
    support_present = [(t, d) for t, d in SUPPORT_TABS.items() if t in wb.sheetnames]
    if support_present:
        P.append("### Front matter & standalone exhibits\n")
        P.append("| Tab | What it shows |")
        P.append("|---|---|")
        for t, d in support_present:
            P.append(f"| **{_esc(t)}** | {_esc(d)} |")
        P.append("")

    if uncovered:
        P.append("> ⚠️ **Coverage note:** the following workbook tab(s) are documented neither in the "
                 "Reading Guide nor here, so this guide may be incomplete — add them to the Reading "
                 f"Guide or to `SUPPORT_TABS` in `r2k_methodology.py`: {', '.join(uncovered)}.\n")

    # glossary
    P.append("---\n")
    P.append("## Glossary\n")
    P.append("| Term | Definition | Notes |")
    P.append("|---|---|---|")
    for term, dfn, note in gloss:
        P.append(f"| **{_esc(term)}** | {_esc(dfn)} | {_esc(note)} |")
    P.append("")
    P.append("---\n")
    P.append("*Generated by `r2k_methodology.py` from the workbook. Re-run after any workbook change so "
             "the tab guide and glossary stay in sync.*")
    return "\n".join(P) + "\n", (uncovered, stale)


# --------------------------------------------------------------------------- entrypoints
def _load():
    import openpyxl
    if not WB.exists():
        raise SystemExit(f"!! workbook not found: {WB}  (run r2k_report.py first, or set R2KG_WORKBOOK)")
    return openpyxl.load_workbook(WB, read_only=True, data_only=True)


def main():
    wb = _load()
    doc, (uncovered, stale) = build(wb)
    OUT.write_text(doc, encoding="utf-8")
    sections = _reading_guide(wb)
    ntabs = sum(len(t) for _, t in sections)
    print(f"  -> {OUT.name}: {len(doc.splitlines())} lines  "
          f"({ntabs} analytical tabs from Reading Guide + {len(SUPPORT_TABS)} support/exhibit tabs, "
          f"{len(_glossary(wb))} glossary terms)")
    if uncovered:
        print(f"  !! {len(uncovered)} workbook tab(s) documented nowhere: {', '.join(uncovered)}")
    if stale:
        print(f"  !! SUPPORT_TABS names a tab not in the workbook (renamed?): {', '.join(stale)}")


def selftest():
    wb = _load()
    sections = _reading_guide(wb)
    gloss = _glossary(wb)
    uncovered, stale = coverage(wb, sections)
    doc, _ = build(wb)
    print(f"  sections: {len(sections)}  analytical tabs: {sum(len(t) for _,t in sections)}  "
          f"glossary terms: {len(gloss)}")
    print(f"  coverage: {'OK (every tab documented)' if not uncovered else 'GAPS -> '+', '.join(uncovered)}")
    if stale:
        print(f"  stale SUPPORT_TABS: {', '.join(stale)}")
    print(f"  build OK: {len(doc.splitlines())} lines")
    print("  SELFTEST:", "PASS" if not uncovered and not stale else "WARN")


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        selftest()
    else:
        main()
