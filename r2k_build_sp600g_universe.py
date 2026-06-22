"""
============================================================
r2k_build_sp600g_universe.py  --  build the extraction universe for the
S&P SmallCap 600 Growth holdings, so the as-filed engine can cover its
constituents (the ones never in R2000G aren't in the fundamentals CSV yet).
============================================================
Reads the S&P 600 Growth quarterly holdings workbook, collects every unique CIK
(union across all snapshots -- point-in-time / survivorship-free), and writes
    sp600g_cik_map.json     { ticker: {cik, ticker, name} }
which r2k_step2_asfiled.py picks up automatically (EXTRA_MAPS, additive: it never
overrides an existing R2000G entry). Then re-run r2k_step2_asfiled.py once to
extend edgar_annual_fundamentals_ASFILED.csv to the combined universe.

It also reports how many 600G CIKs are ALREADY in the fundamentals CSV vs new,
so you know the incremental extraction size before you run it.

RUN
    python r2k_build_sp600g_universe.py
============================================================
"""
from pathlib import Path
import os, json, csv, re
import openpyxl

BASE = Path(os.environ.get("R2KG_BASE", "."))
OUT = BASE / "sp600g_cik_map.json"
FUND = BASE / "edgar_annual_fundamentals_ASFILED.csv"


def find_holdings():
    pats = ["*[Ss][Pp]*600*[Gg]rowth*[Hh]olding*.xlsx", "*600*[Gg]rowth*[Hh]olding*.xlsx",
            "*[Ss]mall[Cc]ap*600*[Hh]olding*.xlsx"]
    for p in pats:
        c = list(BASE.glob(p))
        if c: return c[0]
    raise FileNotFoundError("S&P 600 Growth holdings workbook not found in R2KG_BASE")


def to_i(x):
    try: return int(float(str(x).replace(",", "")))
    except (TypeError, ValueError): return None


def main():
    path = find_holdings()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    uni = {}                      # cik(10) -> {ticker, name, n_snaps}
    n_sheets = 0
    for sn in wb.sheetnames:
        if not re.search(r"\d{1,2}[.\-/]\d{1,2}[.\-/]\d{2,4}", sn): continue
        ws = wb[sn]; hdr = None; idx = {}
        for r in ws.iter_rows(values_only=True):
            if hdr is None:
                hdr = [str(c).strip() if c else "" for c in r]; idx = {h: i for i, h in enumerate(hdr)}
                if "CIK" not in idx and "Ticker" not in idx: hdr = None
                continue
            def g(col):
                i = idx.get(col); return r[i] if (i is not None and i < len(r) and r[i] is not None) else None
            cik = to_i(g("CIK"))
            if cik is None: continue
            c10 = str(cik).zfill(10)
            rec = uni.setdefault(c10, {"ticker": str(g("Ticker") or "").strip(),
                                       "name": str(g("Name") or "").strip(), "n_snaps": 0})
            rec["n_snaps"] += 1
            if not rec["ticker"] and g("Ticker"): rec["ticker"] = str(g("Ticker")).strip()
        n_sheets += 1
    wb.close()

    # already-covered check
    have = set()
    if FUND.exists():
        for row in csv.DictReader(open(FUND, encoding="utf-8-sig")):
            cv = to_i(row.get("cik"))
            if cv is not None: have.add(str(cv).zfill(10))
    new = [c for c in uni if c not in have]

    out = {}
    for c10, rec in uni.items():
        key = rec["ticker"] or c10
        if key in out: key = f"{key}|{c10}"          # ticker collision across delisted/recycled names
        out[key] = {"cik": int(c10), "ticker": rec["ticker"], "name": rec["name"]}
    json.dump(out, open(OUT, "w"), indent=0)

    print(f"  holdings: {path.name}  ({n_sheets} snapshots)")
    print(f"  unique S&P 600 Growth CIKs (survivorship-free union): {len(uni)}")
    print(f"  already in fundamentals CSV: {len(uni) - len(new)}   NEW to extract: {len(new)}")
    print(f"  wrote {OUT.name}  ->  re-run r2k_step2_asfiled.py to extend the dataset to the combined universe.")


if __name__ == "__main__":
    main()
