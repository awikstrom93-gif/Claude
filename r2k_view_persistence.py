"""
r2k_view_persistence.py  --  is R2000G's unprofitable tail a PERMANENT drag or a churn of maturing
names?  Tracks every constituent's profitability cohort year-over-year and tabulates the transitions:
from {profitable, fallen, never-profitable} to the same set, "unknown", or "exited" (no longer an
R2000G constituent the next year -- acquired, delisted, or style-migrated).  The headline is the
NEVER-PROFITABLE tail's annual fate: what share GRADUATES to profitable, EXITS, or persists.

Pure projection of the panel (cohort + membership across years).  RUN:  python r2k_view_persistence.py
"""
from collections import defaultdict

from r2k_universe import get_panel, BASE

OUT = BASE / "r2k_persistence.txt"
COH = ["profitable", "fallen", "never_profitable"]
TO_STATES = ["profitable", "fallen", "never_profitable", "unknown", "exited"]
LABEL = {"profitable": "Profitable", "fallen": "Fallen", "never_profitable": "Never-prof",
         "unknown": "Unknown", "exited": "Exited index"}


def transitions(panel):
    """{from_cohort: {to_state: [count, weight]}} aggregated over all consecutive year-pairs."""
    by = defaultdict(dict)                       # year -> {cik: (cohort, weight)}
    for r in panel:
        if r["covered"]:
            by[int(r["year"])][r["cik"]] = (r["cohort"], r["weight"])
    years = sorted(by)
    trans = {c: defaultdict(lambda: [0, 0.0]) for c in COH}
    for i in range(len(years) - 1):
        y0, y1 = years[i], years[i + 1]
        for cik, (c0, w0) in by[y0].items():
            if c0 not in COH:
                continue
            if cik in by[y1]:
                c1 = by[y1][cik][0]
                to = c1 if c1 in COH else "unknown"
            else:
                to = "exited"
            trans[c0][to][0] += 1
            trans[c0][to][1] += w0
    return years, trans


def _pct_row(d, by_weight):
    tot = sum(v[1 if by_weight else 0] for v in d.values()) or 1e-9
    return {s: 100 * d[s][1 if by_weight else 0] / tot for s in TO_STATES}


def main():
    panel = get_panel(index="R2KG")
    years, trans = transitions(panel)
    L = [f"COHORT PERSISTENCE  --  R2000G profitability transitions, year over year ({years[0]}-{years[-1]})",
         "  rows = cohort in year T; columns = state in year T+1 (share of that cohort's NAMES)", ""]
    corner = "from / to"
    L.append("  " + f"{corner:<16}" + "".join(f"{LABEL[s]:>14}" for s in TO_STATES) + f"{'n/yr':>8}")
    L.append("  " + "-" * (16 + 14 * len(TO_STATES) + 8))
    for c in COH:
        p = _pct_row(trans[c], by_weight=False)
        navg = sum(v[0] for v in trans[c].values()) / max(1, len(years) - 1)
        L.append("  " + f"{LABEL[c]:<16}" + "".join(f"{p[s]:>14.1f}" for s in TO_STATES) + f"{navg:>8.0f}")
    L.append("")
    # the tail's fate, by NAME and by WEIGHT
    np_n = _pct_row(trans["never_profitable"], by_weight=False)
    np_w = _pct_row(trans["never_profitable"], by_weight=True)
    L.append("  NEVER-PROFITABLE tail -- annual fate:")
    L.append(f"      graduates to profitable : {np_n['profitable']:.1f}% of names  /  {np_w['profitable']:.1f}% of tail weight")
    L.append(f"      exits the index         : {np_n['exited']:.1f}% of names  /  {np_w['exited']:.1f}% of tail weight")
    L.append(f"      stays never-profitable  : {np_n['never_profitable']:.1f}% of names  /  {np_w['never_profitable']:.1f}% of tail weight")
    L.append("")
    L.append("  READ: a high graduation + exit rate means the unprofitable tail CHURNS (names mature into")
    L.append("  profitability or leave) rather than permanently dragging. A high 'stays never-profitable'")
    L.append("  share means a persistent structural tail. 'Exited' = acquired, delisted, or left R2000G.")
    OUT.write_text("\n".join(L), encoding="utf-8")
    print("\n".join(L))
    print(f"\n  -> {OUT.name}")


def write_sheet(wb, panel=None):
    from openpyxl.styles import Font, PatternFill, Alignment
    panel = panel if panel is not None else get_panel(index="R2KG")
    years, trans = transitions(panel)
    ws = wb.create_sheet("Cohort Persistence")
    ws.cell(1, 1, "Cohort persistence -- does the unprofitable tail graduate, exit, or persist?").font = Font(bold=True, size=12)
    ws.cell(2, 1, "Rows = profitability cohort in year T; columns = state in year T+1 (share of that "
                  "cohort's names). 'Exited' = no longer an R2000G constituent (acquired/delisted/migrated).").font = Font(size=9, italic=True, color="555555")
    head = ["From \\ To"] + [LABEL[s] for s in TO_STATES]
    fill = PatternFill("solid", fgColor="1F4E5F")
    for c, h in enumerate(head, 1):
        x = ws.cell(4, c, h); x.fill = fill; x.font = Font(bold=True, color="FFFFFF", size=10)
        x.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, coh in enumerate(COH, start=5):
        p = _pct_row(trans[coh], by_weight=False)
        ws.cell(i, 1, LABEL[coh]).font = Font(bold=True)
        for c, s in enumerate(TO_STATES, 2):
            ws.cell(i, c, round(p[s], 1))
    ws.freeze_panes = "B5"
    return "Cohort Persistence"


if __name__ == "__main__":
    main()
