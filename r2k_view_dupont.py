"""
r2k_view_dupont.py  --  DuPont (dollar-aggregate ROE decomposition) sheet as a pure projection of
the canonical panel.  Reproduces step3's "DuPont": index ROE = Net margin x Asset turnover x
Leverage, all dollar-aggregated over the covered constituents.

RUN:  python r2k_view_dupont.py
"""
from r2k_universe import get_panel, diff_sheet, print_sheet, year_snapshot
from r2k_step3_analytics import dollar_agg, _p, _x, dedup_cik

SHEET = "DuPont"
TITLE_TEXT = "DuPont: index ROE = Net margin x Asset turnover x Leverage (dollar-aggregate)"
HDR = ["Snapshot", "Net margin", "Asset turnover (x)", "Leverage (Assets/Equity, x)",
       "Implied ROE", "ROE $agg (check)"]


# the four inputs the DuPont identity telescopes through -- a name contributes to the decomposition
# only if it has ALL of them (see COMMON-UNIVERSE note below).
DUPONT_INPUTS = ("net_income", "revenue", "assets", "equity")


def dupont_rows(panel):
    years = sorted({int(r["year"]) for r in panel})
    snap_of = year_snapshot(panel)                       # actual snapshot date per year (SNAP_MONTH-driven)
    out = []
    for yr in years:
        cov = [r for r in panel if int(r["year"]) == yr and r["covered"]]
        # COMMON UNIVERSE -- the fix for the non-reconciliation. DuPont only telescopes
        #   ($NI/$Rev) x ($Rev/$Assets) x ($Assets/$Eq) = $NI/$Eq = ROE
        # when every factor is a dollar-aggregate over the SAME set of names. If each factor is allowed
        # its own present-set (as before), the $Rev in Net-margin's denominator no longer equals the
        # $Rev in Asset-turnover's numerator, the $Assets don't cancel, and the product drifts from the
        # ROE check (2026: implied 3.5% vs ROE 2.2%, because pre-revenue losers were in ROE but dropped
        # from Net-margin/Asset-turnover). Restrict all four sums to names with ALL inputs present, and
        # implied ROE == ROE $agg by construction. NOTE: this common universe can be slightly narrower
        # than the headline ROE $agg on Index Quality Trends (which keeps any name with NI+equity), so
        # the check column here is the ROE *of the DuPont universe*, not necessarily the headline ROE.
        common = dedup_cik([r for r in cov if all(r[k] is not None for k in DUPONT_INPUTS)])  # one row/company
        # Use AVERAGE assets/equity (opening+closing) so the DuPont ROE matches the per-name ROE and the
        # ROIC $agg convention; fall back to ending only for an old cached panel without _aeq/_ata. The
        # identity still telescopes: NM x (Rev/avgAssets) x (avgAssets/avgEquity) = NI/avgEquity = ROE.
        ata = lambda r: r["_ata"] if r.get("_ata") is not None else r["assets"]
        aeq = lambda r: r["_aeq"] if r.get("_aeq") is not None else r["equity"]
        nm = dollar_agg([(r["net_income"], r["revenue"]) for r in common])
        at = dollar_agg([(r["revenue"], ata(r)) for r in common])
        lev = dollar_agg([(ata(r), aeq(r)) for r in common])
        roe_da = dollar_agg([(r["net_income"], aeq(r)) for r in common])
        implied = (nm * at * lev) if (nm is not None and at is not None and lev is not None) else None
        out.append([snap_of.get(yr, f"{yr}-06-30"), _p(nm), _x(at), _x(lev), _p(implied), _p(roe_da)])
    return out


def _selftest():
    """The DuPont identity must close (implied ROE == ROE check) on the common universe, even when a
    name is missing one input; AND the ROE must be computed on AVERAGE equity/assets (_aeq/_ata),
    not ending -- so a regression to ending equity is caught."""
    # complete names carry _aeq/_ata DIFFERENT from ending equity/assets, so ending vs average is
    # distinguishable: SNI=15, S_aeq=75 -> ROE=15/75=20.0% (ending would give 15/90=16.7%).
    panel = [
        {"year": 2020, "covered": 1, "net_income": 10.0, "revenue": 100.0, "assets": 200.0, "equity": 50.0, "_ata": 180.0, "_aeq": 40.0},
        {"year": 2020, "covered": 1, "net_income": 5.0,  "revenue": 80.0,  "assets": 160.0, "equity": 40.0, "_ata": 150.0, "_aeq": 35.0},
        {"year": 2020, "covered": 1, "net_income": 3.0,  "revenue": 60.0,  "assets": 120.0, "equity": None},   # no equity -> dropped
        {"year": 2020, "covered": 1, "net_income": -8.0, "revenue": None,  "assets": 300.0, "equity": 90.0},   # pre-revenue -> dropped
    ]
    implied, check = dupont_rows(panel)[0][4], dupont_rows(panel)[0][5]
    ok = (implied is not None and check is not None and abs(implied - check) < 1e-9
          and abs(check - 20.0) < 0.05)   # 20.0 = average-based; 16.7 would mean ending equity (regression)
    print(f"  SELFTEST DuPont identity closes on AVERAGE equity "
          f"(implied={implied}, check={check}, expect 20.0): {'PASS' if ok else 'FAIL'}")
    assert ok, (implied, check)


def write_sheet(wb, panel):
    from openpyxl.styles import Font
    ws = wb.create_sheet(SHEET)
    ws.cell(row=1, column=1, value=TITLE_TEXT).font = Font(bold=True, size=12)
    for c, h in enumerate(HDR, 1):
        ws.cell(row=3, column=c, value=h).font = Font(bold=True)
    last = 3
    for i, row in enumerate(dupont_rows(panel), start=4):
        for c, v in enumerate(row, 1):
            ws.cell(row=i, column=c, value=v)
        last = i
    ws.cell(row=last + 2, column=1, value=(
        "Net margin, Asset turnover, Leverage and ROE are all dollar-aggregates on the COMMON universe "
        "(names with revenue, net income, average assets AND average equity all present), so the three "
        "factors multiply EXACTLY to the ROE check. ROE/turnover/leverage use AVERAGE (opening+closing) "
        "assets & equity -- the same convention as the per-name ROE and ROIC. Because this common universe "
        "is slightly narrower than the headline 'ROE $agg' on Index Quality Trends (which keeps any name "
        "with net income + equity, incl. those missing revenue/assets), the ROE here can differ from that "
        "headline by a few tenths of a point; it is the ROE OF THE DuPont universe, by construction."
    )).font = Font(size=9, italic=True, color="555555")
    return ws


def main():
    import sys
    if "--selftest" in sys.argv:
        _selftest(); return
    _selftest()
    panel = get_panel()
    rows = dupont_rows(panel)
    print(f"\n  {SHEET} (panel-derived):")
    print_sheet(HDR, rows, ncols=6)
    diff_sheet(SHEET, HDR, rows)


if __name__ == "__main__":
    main()
