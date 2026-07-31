#!/usr/bin/env python3
"""
build_attribution_json.py
=========================

Phase 2 of the Battle Book Agent data pipeline.

Converts Morningstar Direct attribution exports into per-manager JSON that
answers one question for the commentary agent:

    "Why did this manager outperform or underperform?"

This is a COMMENTARY tool, not an attribution audit tool. The battle-book-facing
JSON deliberately contains no residual, attribution-gap, coverage, expense-ratio
or benchmark-reconstruction fields. Those numbers are still computed, but they
live only in debug_attribution_report.json for internal QA.

Usage
-----
    python build_attribution_json.py --quarter "2026 Q2"
    python build_attribution_json.py --quarter latest
    python build_attribution_json.py --quarter "2026 Q2" --no-patch
    python build_attribution_json.py --quarter "2026 Q2" \
        --input-root "D:\\test\\in" --output-root "D:\\test\\out"

Prerequisite
------------
Phase 1 must have run for the same quarter. This script reads the asset-class
JSON files to resolve each manager, confirm the asset class and benchmark, and
then patches attribution availability pointers back into them.

Inputs
------
    <input-root>/<quarter>/Attribution/*.xlsx

Outputs
-------
    <output-root>/<quarter>/attribution/<manager_slug>.json
    <output-root>/<quarter>/debug_attribution_report.json
    <output-root>/<quarter>/{large,mid,small}_growth.json   (patched in place)

Workbook layout
---------------
Every attribution export carries two sheets, `Template` (an export manifest) and
`Attribution` (the grid). The grid is:

    rows 1-5    metadata (Name / Portfolio / Benchmark / Currency / Exported)
    row  8      period band labels, e.g. "4-1-2026 - 6-30-2026"
    row  9      metric group labels (Weights % / Return % / ...)
    row  10     leaf headers (Portfolio / Benchmark / +/- / Allocation % / ...)
    row  11+    data: column A opens a section, column B/C are securities

Columns A-C are identity; the remainder is seven 13-column period bands. Nothing
here is addressed by fixed cell reference: bands are matched on the date range
parsed from their label, and leaf columns are mapped from the row 9/10 header
text, so a workbook with a different band count or column order still parses.

Structure the code so additional attribution views can be added later.
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError:  # pragma: no cover - guidance for a fresh machine
    sys.stderr.write(
        "ERROR: openpyxl is not installed.\n"
        "       Run:  pip install -r requirements.txt\n"
    )
    raise


# =============================================================================
# SECTION 1 - CONSTANTS / DISCOVERED LAYOUT ASSUMPTIONS
# -----------------------------------------------------------------------------
# Everything derived from inspecting real Morningstar attribution exports. If a
# template changes, this should be the only section that needs editing.
# =============================================================================

SCRIPT_VERSION = "1.0.0"

# --- Roots (mirror build_quarterly_json.py) -----------------------------------
DEFAULT_INPUT_ROOT = (
    r"C:\Users\alex.wikstrom\OneDrive - NFP Corp\Research\Battle Book Agent"
    r"\Quarterly Battle Book Inputs"
)
DEFAULT_OUTPUT_ROOT = (
    r"C:\Users\alex.wikstrom\OneDrive - NFP Corp\Research\Battle Book Agent"
    r"\Quarterly Battle Book Data"
)

ATTRIBUTION_INPUT_SUBFOLDER = "Attribution"
ATTRIBUTION_OUTPUT_SUBFOLDER = "attribution"
PACK_OUTPUT_SUBFOLDER = "packs"
DEBUG_REPORT_FILENAME = "debug_attribution_report.json"

# --- Battle book packs --------------------------------------------------------
# A pack is one self-contained file per manager per quarter: everything needed to
# write that battle book, in a single fetch. It exists so runtime retrieval is a
# deterministic file read rather than a search across large asset-class files.
#
# A pack is only written when an attribution workbook was supplied, so the
# pack's existence IS the attribution-available gate: no pack means no battle
# book. Nothing else has to stay in sync.
#
# Contents are trimmed deliberately, because the whole pack is returned into the
# agent's context in one response:
#   * `summaries` is carried instead of the full `market_data` row lists - the
#     top and bottom tens are what the market backdrop actually cites.
#   * Extra market periods are carried as summaries only, never full rows.
#   * `security_attribution` is not carried; the securities that get discussed
#     are the top movers, so those records are enriched with weights and
#     contribution instead.
PACK_EXTRA_MARKET_PERIODS: Tuple[str, ...] = ("ytd",)
# Fields lifted from security_attribution onto each top mover.
PACK_MOVER_ENRICHMENT: Tuple[str, ...] = (
    "portfolio_weight",
    "benchmark_weight",
    "contribution_portfolio",
    "contribution_benchmark",
    "contribution_active",
    "selection_effect_bps",
)
# Warn when a pack grows past this; the whole file is returned in one response.
PACK_SIZE_WARN_KB = 60

# The manifest maps every lookup key a manager can be asked for - fund name,
# ticker, strategy name - onto its pack path.
#
# It exists so nothing downstream has to rebuild the slug. Slugs strip
# characters that normalisation keeps (hyphens in "Mid-Cap", the ampersand in
# "Segall Bryant & Hamill", the mark in "Ultra(R)"), so a caller reconstructing a
# filename from a manager name gets it wrong for a sixth of the universe.
# Normalisation is reproducible anywhere with lowercase, strip periods and
# collapse spaces; the filename is not. So callers normalise, and the manifest
# resolves.
PACK_MANIFEST_FILENAME = "_manifest.json"

# Morningstar fund names carry a share-class suffix that nobody says out loud:
# a request for "Fidelity Blue Chip Growth" means "Fidelity Blue Chip Growth K".
# Stripping these produces an extra manifest key, but only when the stripped
# name resolves to exactly one manager in the whole universe. Anything that
# would become ambiguous is rejected rather than guessed, so the failure mode is
# the agent asking which fund was meant - never quietly serving the wrong one.
SHARE_CLASS_SUFFIX_RE = re.compile(
    r"\s+(?:class\s+)?(?:a|c|i|is|k|k5|k6|m|n|r|r3|r4|r5|r6|s|y|z|i2|i3|"
    r"inst|instl|institutional|inv|investor|adm|admiral|adv|advisor|premier|"
    r"retail|svc|service|shares?)$",
    re.IGNORECASE,
)
SHARE_CLASS_STRIP_ROUNDS = 3

# --- Quarter folder naming ----------------------------------------------------
QUARTER_FOLDER_PATTERNS = (
    re.compile(r"^\s*(?P<year>\d{4})\s*[-_ ]?\s*[Qq](?P<quarter>[1-4])\s*$"),
    re.compile(r"^\s*[Qq](?P<quarter>[1-4])\s*[-_ ]?\s*(?P<year>\d{4})\s*$"),
)
QUARTER_FOLDER_TEMPLATE = "{year} Q{quarter}"

# --- Phase 1 asset-class files -----------------------------------------------
ASSET_CLASS_FILES: Tuple[Tuple[str, str], ...] = (
    ("large_growth", "large_growth.json"),
    ("mid_growth", "mid_growth.json"),
    ("small_growth", "small_growth.json"),
)

# --- Worksheet names ----------------------------------------------------------
TEMPLATE_SHEET_ALIASES = ("template",)
ATTRIBUTION_SHEET_ALIASES = ("attribution",)

# --- Template / metadata keys -------------------------------------------------
METADATA_KEYS = {
    "name": "name",
    "portfolio": "portfolio",
    "benchmark": "benchmark",
    "currency": "currency",
    "date exported": "exported_at",
}
CLASSIFICATION_HINT_RE = re.compile(r"^\s*\d+\.\s*(?P<value>.+?)\s*$")
ATTRIBUTION_MODEL_RE = re.compile(r"three[- ]factor", re.IGNORECASE)

# --- Grid geometry ------------------------------------------------------------
# Expected positions; all of them are validated by content, never trusted blind.
EXPECTED_BAND_LABEL_ROW = 8
EXPECTED_METRIC_GROUP_ROW = 9
EXPECTED_LEAF_HEADER_ROW = 10
HEADER_SEARCH_MAX_ROW = 30
IDENTITY_COLUMN_COUNT = 3  # A=GICS Sector, B=Name, C=Ticker
EXPECTED_BAND_WIDTH = 13

# Band labels look like "4-1-2026 - 6-30-2026".
BAND_LABEL_DATE_RANGE_RE = re.compile(
    r"^\s*(?P<start>\d{1,2}-\d{1,2}-\d{4})\s*-\s*(?P<end>\d{1,2}-\d{1,2}-\d{4})\s*$"
)
BAND_LABEL_DATE_FORMATS = ("%m-%d-%Y", "%m/%d/%Y")

# --- Metric mapping: (metric group, leaf header) -> logical field -------------
METRIC_GROUP_ALIASES = {
    "weights %": "weights",
    "weight %": "weights",
    "return %": "returns",
    "contribution to return %": "contribution",
    "attribution effect": "effects",
}
LEAF_FIELD_MAP: Dict[Tuple[str, str], str] = {
    ("weights", "portfolio"): "portfolio_weight",
    ("weights", "benchmark"): "benchmark_weight",
    ("weights", "+/-"): "active_weight",
    ("returns", "portfolio"): "portfolio_return",
    ("returns", "benchmark"): "benchmark_return",
    ("returns", "+/-"): "return_differential",
    ("contribution", "portfolio"): "contribution_portfolio",
    ("contribution", "benchmark"): "contribution_benchmark",
    ("contribution", "+/-"): "contribution_active",
    ("effects", "allocation %"): "allocation_effect",
    ("effects", "selection %"): "selection_effect",
    ("effects", "interaction %"): "interaction_effect",
    ("effects", "active ret%"): "total_effect",
    ("effects", "active ret %"): "total_effect",
    ("effects", "active return %"): "total_effect",
}
# Fields required before a band is considered usable.
REQUIRED_BAND_FIELDS = ("portfolio_weight", "total_effect")

# --- Row taxonomy -------------------------------------------------------------
GICS_SECTORS: Tuple[str, ...] = (
    "Communication Services",
    "Consumer Discretionary",
    "Consumer Staples",
    "Energy",
    "Financials",
    "Health Care",
    "Industrials",
    "Information Technology",
    "Materials",
    "Real Estate",
    "Utilities",
)
GICS_SECTOR_LOOKUP = {s.lower(): s for s in GICS_SECTORS}

# Attributed non-GICS buckets: they carry effects and belong in reconciliation,
# but they are not GICS sectors so they stay out of sector_attribution.
ATTRIBUTED_NON_SECTOR_ROWS = ("cash", "unclassified")
# Cash is not a GICS sector, but a material cash allocation effect is a genuine
# reason a manager lagged, so it is surfaced separately from sector_attribution.
CASH_ROW_LABEL = "cash"
# Buckets with no attribution effects at all - their securities are never stored.
UNATTRIBUTED_SECTION_ROWS = ("bond", "missing performance", "other")
# Control/total rows.
ROW_ATTRIBUTION_TOTAL = "attribution total"
ROW_TOTAL = "total"
ROW_REPORTED_TOTAL = "reported total"
ROW_EXPENSE_RATIO = "expense ratio"
ROW_RESIDUAL_PREFIX = "residual"
CONTROL_ROWS = (
    ROW_ATTRIBUTION_TOTAL,
    ROW_TOTAL,
    ROW_REPORTED_TOTAL,
    ROW_EXPENSE_RATIO,
)

# --- Logical periods ----------------------------------------------------------
# Resolved by matching each band's parsed date range against the calendar
# quarter, exactly as build_quarterly_json.py resolves its period blocks.
LOGICAL_PERIODS: Tuple[Tuple[str, int], ...] = (
    ("selected_quarter", 0),
    ("two_quarters_ago", 1),
    ("three_quarters_ago", 2),
    ("four_quarters_ago", 3),
)
QUARTER_TREND_SEQUENCE = tuple(key for key, _ in LOGICAL_PERIODS)  # newest first

# --- Materiality / sizing -----------------------------------------------------
# A security is stored when |total_effect| reaches this many percentage points.
# 0.05 pp == 5 bps.
MATERIALITY_THRESHOLD_PP = 0.05
TOP_N_SECURITIES = 10
TOP_N_SECTORS = 5
CONCENTRATION_TOP_N = 5
# Five names out of several hundred producing this share of one side's effect is
# concentrated by any reasonable reading. Precomputed so the agent never has to
# form an impression of breadth from the numbers.
CONCENTRATION_SHARE_THRESHOLD = 0.40

# --- Classification thresholds ------------------------------------------------
# Active weight (pp) within this band counts as neutral rather than over/under.
NEUTRAL_ACTIVE_WEIGHT_PP = 0.10
# A security is "held" when its portfolio weight exceeds this.
HELD_WEIGHT_EPSILON = 0.0
# Driver dominance: largest |effect| vs second largest.
DRIVER_HIGH_CONFIDENCE_RATIO = 2.00
DRIVER_MIN_DOMINANCE_RATIO = 1.25

# --- Output shaping -----------------------------------------------------------
ROUND_PCT_DECIMALS = 4
ROUND_BPS_DECIMALS = 1
ROUND_SHARE_DECIMALS = 4

# --- Reconciliation tolerances (debug only) -----------------------------------
RECONCILE_TOLERANCE_PP = 0.05

LOG = logging.getLogger("build_attribution_json")


# =============================================================================
# SECTION 2 - SHARED UTILITIES (same conventions as build_quarterly_json.py)
# =============================================================================


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ").replace("\r", " ")
    return re.sub(r"\s+", " ", text).strip().lower()


def normalize_lookup_name(value: Any) -> str:
    """
    Normalise a manager name for lookup keys.

    Identical rules to build_quarterly_json.py: lowercase, trim, collapse
    duplicate spaces, remove periods. Must stay in sync - this is the join key.
    """
    if value is None:
        return ""
    text = str(value).replace("\n", " ").replace("\r", " ")
    text = text.replace(".", "")
    return re.sub(r"\s+", " ", text).strip().lower()


def manager_slug(manager_name: str) -> str:
    """Filesystem-safe slug derived from the normalised manager name."""
    key = normalize_lookup_name(manager_name)
    slug = re.sub(r"[^a-z0-9]+", "_", key).strip("_")
    return slug or "unnamed_manager"


def to_float(value: Any) -> Optional[float]:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        result = float(value)
        return None if result != result else result
    text = str(value).strip()
    if not text or text in {"-", "--", "n/a", "N/A", "NA", "nan"}:
        return None
    text = text.replace(",", "").replace("%", "")
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    try:
        result = float(text)
    except ValueError:
        return None
    return -result if negative else result


def round_pct(value: Optional[float]) -> Optional[float]:
    return None if value is None else round(value, ROUND_PCT_DECIMALS)


def to_bps(percent_value: Optional[float]) -> Optional[float]:
    """
    Basis points from an already-rounded percentage.

    Battle books quote basis points while the workbook is in percent. Deriving
    bps from the rounded percent keeps `bps == pct * 100` exactly true, which is
    a validation gate.
    """
    if percent_value is None:
        return None
    return round(percent_value * 100, ROUND_BPS_DECIMALS)


def round_share(value: Optional[float]) -> Optional[float]:
    return None if value is None else round(value, ROUND_SHARE_DECIMALS)


def cell_text(worksheet: Worksheet, row: int, column: int) -> Optional[str]:
    value = worksheet.cell(row=row, column=column).value
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def iso_or_empty(day: Optional[date]) -> str:
    return day.isoformat() if day else ""


def write_json(path: Path, payload: Any, compact: bool = False) -> None:
    """Write JSON. `compact` drops indentation for files fetched into an agent's
    context, where whitespace is pure cost; everything else stays readable."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        if compact:
            json.dump(payload, handle, separators=(",", ":"), ensure_ascii=False)
        else:
            json.dump(payload, handle, indent=2, ensure_ascii=False)
        handle.write("\n")
    LOG.info("Wrote %s (%.1f KB)", path, path.stat().st_size / 1024)


# =============================================================================
# SECTION 3 - QUARTER RESOLUTION
# =============================================================================


@dataclass(frozen=True)
class Quarter:
    year: int
    quarter: int

    @property
    def label(self) -> str:
        return QUARTER_FOLDER_TEMPLATE.format(year=self.year, quarter=self.quarter)

    @property
    def start_date(self) -> date:
        return date(self.year, 3 * (self.quarter - 1) + 1, 1)

    @property
    def end_date(self) -> date:
        month = 3 * self.quarter
        return date(self.year, month, {3: 31, 6: 30, 9: 30, 12: 31}[month])

    def contains(self, day: date) -> bool:
        return self.start_date <= day <= self.end_date

    def shifted(self, quarters_back: int) -> "Quarter":
        absolute = self.year * 4 + (self.quarter - 1) - quarters_back
        return Quarter(absolute // 4, absolute % 4 + 1)


def parse_quarter_token(token: str) -> Optional[Quarter]:
    for pattern in QUARTER_FOLDER_PATTERNS:
        match = pattern.match(token)
        if match:
            return Quarter(int(match.group("year")), int(match.group("quarter")))
    return None


def discover_quarter_folders(input_root: Path) -> List[Tuple[Quarter, Path]]:
    if not input_root.is_dir():
        raise FileNotFoundError(
            f"Input root folder does not exist: {input_root}\n"
            f"       Check --input-root or the DEFAULT_INPUT_ROOT constant."
        )
    found: List[Tuple[Quarter, Path]] = []
    for child in sorted(input_root.iterdir()):
        if child.is_dir():
            quarter = parse_quarter_token(child.name)
            if quarter is not None:
                found.append((quarter, child))
    found.sort(key=lambda item: (item[0].year, item[0].quarter))
    return found


def resolve_quarter(input_root: Path, requested: str) -> Tuple[Quarter, Path]:
    available = discover_quarter_folders(input_root)
    if not available:
        raise FileNotFoundError(
            f"No quarter folders found under: {input_root}\n"
            f"       Expected folders named like '2026 Q2'."
        )
    if requested.strip().lower() == "latest":
        quarter, folder = available[-1]
        LOG.info("--quarter latest resolved to '%s'", quarter.label)
        return quarter, folder

    quarter = parse_quarter_token(requested)
    if quarter is None:
        raise ValueError(
            f"Could not understand --quarter {requested!r}.\n"
            f'       Use a value like "2026 Q2" or the keyword "latest".'
        )
    for candidate, folder in available:
        if candidate == quarter:
            return quarter, folder
    known = ", ".join(q.label for q, _ in available)
    raise FileNotFoundError(
        f"Quarter folder '{quarter.label}' not found under: {input_root}\n"
        f"       Available quarters: {known}"
    )


# =============================================================================
# SECTION 4 - PHASE 1 INDEX (manager linkage)
# =============================================================================


@dataclass
class ManagerRef:
    """A manager located inside a Phase 1 asset-class JSON."""

    lookup_key: str
    manager: str
    asset_class: str
    asset_class_key: str
    asset_class_file: str
    index: int
    benchmark: str
    excess_return: Optional[float]
    match_type: str = "manager_name"


@dataclass
class Phase1Index:
    """Everything Phase 2 needs from the Phase 1 outputs."""

    documents: Dict[str, Dict[str, Any]] = field(default_factory=dict)
    paths: Dict[str, Path] = field(default_factory=dict)
    by_manager_name: Dict[str, List[ManagerRef]] = field(default_factory=dict)
    by_alias: Dict[str, List[ManagerRef]] = field(default_factory=dict)

    def resolve(self, lookup_key: str) -> Tuple[List[ManagerRef], str]:
        """Primary manager-name matches win; aliases are the fallback."""
        primary = self.by_manager_name.get(lookup_key, [])
        if primary:
            return primary, "manager_name"
        return self.by_alias.get(lookup_key, []), "alias"


def load_phase1_index(output_folder: Path) -> Phase1Index:
    """Load the asset-class JSONs written by build_quarterly_json.py."""
    index = Phase1Index()
    missing: List[str] = []

    for asset_class_key, filename in ASSET_CLASS_FILES:
        path = output_folder / filename
        if not path.is_file():
            missing.append(filename)
            continue
        with path.open("r", encoding="utf-8") as handle:
            document = json.load(handle)
        index.documents[asset_class_key] = document
        index.paths[asset_class_key] = path

        managers = document.get("managers", [])
        for position, manager in enumerate(managers):
            key = normalize_lookup_name(manager.get("manager"))
            if not key:
                continue
            index.by_manager_name.setdefault(key, []).append(
                ManagerRef(
                    lookup_key=key,
                    manager=manager.get("manager", ""),
                    asset_class=document.get("asset_class", ""),
                    asset_class_key=asset_class_key,
                    asset_class_file=filename,
                    index=position,
                    benchmark=manager.get("benchmark", ""),
                    excess_return=manager.get("excess_return_cumulative"),
                )
            )

        # Phase 1 also indexes tickers and strategy names. Keep them as a
        # clearly-labelled fallback so a slightly different portfolio label
        # still resolves, without letting an alias outrank a real name.
        for alias, position in document.get("manager_lookup", {}).items():
            if not isinstance(position, int) or not 0 <= position < len(managers):
                continue
            manager = managers[position]
            if normalize_lookup_name(manager.get("manager")) == alias:
                continue
            index.by_alias.setdefault(alias, []).append(
                ManagerRef(
                    lookup_key=alias,
                    manager=manager.get("manager", ""),
                    asset_class=document.get("asset_class", ""),
                    asset_class_key=asset_class_key,
                    asset_class_file=filename,
                    index=position,
                    benchmark=manager.get("benchmark", ""),
                    excess_return=manager.get("excess_return_cumulative"),
                    match_type="alias",
                )
            )

    if missing:
        raise FileNotFoundError(
            f"Phase 1 output missing from {output_folder}: {', '.join(missing)}\n"
            f"       Run build_quarterly_json.py for this quarter first."
        )
    return index


# =============================================================================
# SECTION 5 - WORKBOOK LAYOUT DETECTION
# =============================================================================


@dataclass
class PeriodBand:
    """One 13-column attribution period band."""

    label: str
    first_column: int
    last_column: int
    start_date: Optional[date]
    end_date: Optional[date]
    field_columns: Dict[str, int] = field(default_factory=dict)

    @property
    def column_range(self) -> str:
        return (
            f"{get_column_letter(self.first_column)}:"
            f"{get_column_letter(self.last_column)}"
        )

    @property
    def width(self) -> int:
        return self.last_column - self.first_column + 1

    def describe(self) -> Dict[str, Any]:
        return {
            "label": self.label,
            "columns": self.column_range,
            "width": self.width,
            "start_date": iso_or_empty(self.start_date),
            "end_date": iso_or_empty(self.end_date),
            "fields_mapped": sorted(self.field_columns),
        }


def parse_band_dates(label: str) -> Tuple[Optional[date], Optional[date]]:
    """Parse '4-1-2026 - 6-30-2026' into start and end dates."""
    match = BAND_LABEL_DATE_RANGE_RE.match(label or "")
    if not match:
        return None, None

    def parse_one(token: str) -> Optional[date]:
        for fmt in BAND_LABEL_DATE_FORMATS:
            try:
                return datetime.strptime(token, fmt).date()
            except ValueError:
                continue
        return None

    return parse_one(match.group("start")), parse_one(match.group("end"))


def find_header_rows(worksheet: Worksheet) -> Tuple[int, int, int]:
    """
    Locate (band label row, metric group row, leaf header row).

    Anchored on the leaf header row, which is the row containing the
    'Allocation %' header. The two rows above it are the group and band rows.
    """
    target_leaves = {"allocation %", "selection %", "interaction %"}
    limit = min(worksheet.max_row, HEADER_SEARCH_MAX_ROW)
    for row in range(1, limit + 1):
        seen = set()
        for column in range(1, min(worksheet.max_column, 200) + 1):
            header = normalize_header(worksheet.cell(row=row, column=column).value)
            if header in target_leaves:
                seen.add(header)
                if len(seen) == len(target_leaves):
                    break
        if len(seen) == len(target_leaves):
            leaf_row = row
            return max(1, leaf_row - 2), max(1, leaf_row - 1), leaf_row

    raise ValueError(
        "Could not locate the attribution header rows (no row containing "
        "'Allocation %', 'Selection %' and 'Interaction %'). "
        "The export template may have changed."
    )


def detect_period_bands(
    worksheet: Worksheet, band_row: int, group_row: int, leaf_row: int
) -> List[PeriodBand]:
    """
    Discover every period band and map its leaf columns to logical fields.

    Band starts are the populated cells of the band label row beyond the
    identity columns. Within a band, the metric group is carried forward from
    the merged group row and combined with the leaf header to identify a field.
    """
    max_column = worksheet.max_column
    starts: List[int] = []
    labels: Dict[int, str] = {}

    for column in range(IDENTITY_COLUMN_COUNT + 1, max_column + 1):
        text = cell_text(worksheet, band_row, column)
        if text:
            starts.append(column)
            labels[column] = text

    if not starts:
        raise ValueError(
            f"No period bands found on header row {band_row}. "
            f"The export template may have changed."
        )

    bands: List[PeriodBand] = []
    for position, first_column in enumerate(starts):
        last_column = (
            starts[position + 1] - 1 if position + 1 < len(starts) else max_column
        )
        label = labels[first_column]
        start_date, end_date = parse_band_dates(label)

        field_columns: Dict[str, int] = {}
        current_group: Optional[str] = None
        for column in range(first_column, last_column + 1):
            group_text = normalize_header(worksheet.cell(row=group_row, column=column).value)
            if group_text:
                current_group = METRIC_GROUP_ALIASES.get(group_text, group_text)
            leaf_text = normalize_header(worksheet.cell(row=leaf_row, column=column).value)
            if not leaf_text or current_group is None:
                continue
            logical = LEAF_FIELD_MAP.get((current_group, leaf_text))
            if logical and logical not in field_columns:
                field_columns[logical] = column

        bands.append(
            PeriodBand(
                label=label,
                first_column=first_column,
                last_column=last_column,
                start_date=start_date,
                end_date=end_date,
                field_columns=field_columns,
            )
        )
    return bands


def select_band_for_quarter(
    bands: Sequence[PeriodBand], quarter: Quarter
) -> Optional[PeriodBand]:
    """
    Find the band covering a calendar quarter.

    Date match first (the label carries an explicit range); falls back to a
    quarter-shaped label such as 'Q2 2026' or '2026 Q2' if a future export
    stops using date ranges. Position is never used.
    """
    usable = [
        band
        for band in bands
        if all(f in band.field_columns for f in REQUIRED_BAND_FIELDS)
    ]
    for band in usable:
        if band.start_date == quarter.start_date and band.end_date == quarter.end_date:
            return band

    wanted = {
        f"q{quarter.quarter} {quarter.year}",
        f"{quarter.year} q{quarter.quarter}",
        quarter.label.lower(),
    }
    for band in usable:
        if normalize_header(band.label) in wanted:
            return band
    return None


# =============================================================================
# SECTION 6 - ROW CLASSIFICATION AND EXTRACTION
# =============================================================================

ROW_KIND_GICS_SECTOR = "gics_sector"
ROW_KIND_ATTRIBUTED_BUCKET = "attributed_bucket"      # Cash, Unclassified
ROW_KIND_UNATTRIBUTED_BUCKET = "unattributed_bucket"  # Bond, Other, Missing Perf
ROW_KIND_CONTROL = "control"                          # Attribution Total, Total...
ROW_KIND_RESIDUAL = "residual"
ROW_KIND_SECURITY = "security"
ROW_KIND_BLANK = "blank"


def classify_section_label(label: str) -> str:
    lowered = normalize_header(label)
    if lowered in GICS_SECTOR_LOOKUP:
        return ROW_KIND_GICS_SECTOR
    if lowered in ATTRIBUTED_NON_SECTOR_ROWS:
        return ROW_KIND_ATTRIBUTED_BUCKET
    if lowered in UNATTRIBUTED_SECTION_ROWS:
        return ROW_KIND_UNATTRIBUTED_BUCKET
    if lowered.startswith(ROW_RESIDUAL_PREFIX):
        return ROW_KIND_RESIDUAL
    if lowered in CONTROL_ROWS:
        return ROW_KIND_CONTROL
    return ROW_KIND_ATTRIBUTED_BUCKET  # unknown but effect-bearing; keep visible


def read_band_values(
    worksheet: Worksheet, row: int, band: PeriodBand
) -> Dict[str, Optional[float]]:
    """Read every mapped field of one band for one row."""
    return {
        name: to_float(worksheet.cell(row=row, column=column).value)
        for name, column in band.field_columns.items()
    }


@dataclass
class RawRow:
    """One parsed grid row, before it becomes an output record."""

    row: int
    kind: str
    section_label: Optional[str]
    section_kind: Optional[str]
    security_name: Optional[str]
    ticker: Optional[str]
    values: Dict[str, Optional[float]]


def scan_rows(
    worksheet: Worksheet, first_data_row: int, band: PeriodBand
) -> Tuple[List[RawRow], Dict[str, int]]:
    """Walk the grid once, classifying every row and reading the chosen band."""
    rows: List[RawRow] = []
    counts: Dict[str, int] = {}
    current_section: Optional[str] = None
    current_section_kind: Optional[str] = None

    for row in range(first_data_row, worksheet.max_row + 1):
        section_label = cell_text(worksheet, row, 1)
        security_name = cell_text(worksheet, row, 2)
        ticker = cell_text(worksheet, row, 3)

        if section_label:
            current_section = section_label
            current_section_kind = classify_section_label(section_label)
            kind = current_section_kind
        elif security_name:
            kind = ROW_KIND_SECURITY
        else:
            counts[ROW_KIND_BLANK] = counts.get(ROW_KIND_BLANK, 0) + 1
            continue

        counts[kind] = counts.get(kind, 0) + 1
        rows.append(
            RawRow(
                row=row,
                kind=kind,
                section_label=section_label or current_section,
                section_kind=current_section_kind,
                security_name=security_name,
                ticker=ticker,
                values=read_band_values(worksheet, row, band),
            )
        )
    return rows, counts


# =============================================================================
# SECTION 7 - DERIVED CLASSIFICATION (driver, position, reason)
# =============================================================================


def resolve_driver(
    allocation: Optional[float],
    selection: Optional[float],
    interaction: Optional[float],
) -> Tuple[str, str]:
    """
    Decide which effect drove a result, by absolute magnitude.

    Returns (driver, confidence). When no single effect clearly dominates the
    driver is 'mixed', so the commentary agent never claims a driver the numbers
    do not support.
    """
    candidates = [
        ("allocation", abs(allocation) if allocation is not None else None),
        ("selection", abs(selection) if selection is not None else None),
        ("interaction", abs(interaction) if interaction is not None else None),
    ]
    present = [(name, value) for name, value in candidates if value is not None]
    if not present:
        return "mixed", "insufficient_data"

    present.sort(key=lambda item: item[1], reverse=True)
    top_name, top_value = present[0]
    if top_value == 0:
        return "mixed", "low"

    runner_up = present[1][1] if len(present) > 1 else 0.0
    ratio = float("inf") if runner_up == 0 else top_value / runner_up

    if ratio >= DRIVER_HIGH_CONFIDENCE_RATIO:
        return top_name, "high"
    if ratio >= DRIVER_MIN_DOMINANCE_RATIO:
        return top_name, "medium"
    return "mixed", "low"


def resolve_position(
    portfolio_weight: Optional[float], active_weight: Optional[float]
) -> str:
    """overweight / underweight / neutral / not_held."""
    if portfolio_weight is None or portfolio_weight <= HELD_WEIGHT_EPSILON:
        return "not_held"
    if active_weight is None:
        return "neutral"
    if abs(active_weight) < NEUTRAL_ACTIVE_WEIGHT_PP:
        return "neutral"
    return "overweight" if active_weight > 0 else "underweight"


def effect_direction(effect: Optional[float]) -> Optional[str]:
    """
    Whether a sector's allocation or selection helped or hurt, in words.

    The sign is right there in the _bps field, but three drafts have credited a
    position for a gain its own effect contradicts - "an overweight position and
    strong selection combined to add 52 bps" when allocation cost 16, and "on
    the strength of both an overweight stance and strong stock picking" when the
    overweight cost 35. Instructing the agent to read the sign has not held, so
    the reading is done here.
    """
    if effect is None:
        return None
    if abs(effect) < MATERIALITY_THRESHOLD_PP:
        return "negligible"
    return "helped" if effect > 0 else "hurt"


def resolve_driver_label(driver: str, confidence: str) -> str:
    """
    What drove the result, as a phrase the agent can only copy.

    Where the leading factor is not clearly dominant, the honest statement is
    that both contributed - so say that, rather than naming one and hedging it
    with a confidence qualifier the commentary is not allowed to mention.
    """
    if driver == "mixed" or confidence == "low":
        return "both allocation and selection"
    return driver


def resolve_position_label(
    held: bool,
    benchmark_weight: Optional[float],
    active_weight: Optional[float],
) -> str:
    """
    The phrase the commentary agent should use for this security, precomputed.

    `reason` below encodes ownership and relative performance together, which
    makes it read like a position label when it is not: a name carrying
    "overweight_outperformer" with a zero benchmark weight is an out-of-benchmark
    holding, not an overweight. Instructing an agent to prefer benchmark_weight
    over reason proved unreliable across three wordings, so the judgement is made
    here instead and the agent is given a phrase it can only copy.
    """
    if not held:
        return "not held"
    if benchmark_weight is None or benchmark_weight <= HELD_WEIGHT_EPSILON:
        return "out-of-benchmark holding"
    if active_weight is None or abs(active_weight) < NEUTRAL_ACTIVE_WEIGHT_PP:
        return "in line with the benchmark"
    return "overweight" if active_weight > 0 else "underweight"


def resolve_security_reason(
    held: bool,
    active_weight: Optional[float],
    security_return: Optional[float],
    benchmark_total_return: Optional[float],
    total_effect: Optional[float],
) -> str:
    """
    Turn a security row into a phrase the commentary agent can use safely.

    The held / not-held split is the point of this enum: describing a benchmark
    constituent the manager never owned as a "holding" is the most damaging
    error attribution commentary can make.
    """
    outperformed: Optional[bool] = None
    if security_return is not None and benchmark_total_return is not None:
        outperformed = security_return > benchmark_total_return

    if not held:
        if outperformed is None:
            return "mixed"
        return "not_held_outperformer" if outperformed else "not_held_underperformer"

    if active_weight is None or abs(active_weight) < NEUTRAL_ACTIVE_WEIGHT_PP:
        if total_effect is None:
            return "mixed"
        return "held_positive_selection" if total_effect > 0 else "held_negative_selection"

    if outperformed is None:
        return "mixed"
    if active_weight > 0:
        return "overweight_outperformer" if outperformed else "overweight_underperformer"
    return "underweight_outperformer" if outperformed else "underweight_underperformer"


# =============================================================================
# SECTION 8 - BATTLE-BOOK SECTION BUILDERS
# -----------------------------------------------------------------------------
# Everything in this section feeds the manager-facing JSON. It answers "why did
# this manager outperform or underperform" and deliberately carries no
# reconciliation, residual, coverage or expense language - those belong to
# SECTION 9 and the debug report only.
# =============================================================================


def build_attribution_summary(total_values: Dict[str, Optional[float]]) -> Dict[str, Any]:
    allocation = round_pct(total_values.get("allocation_effect"))
    selection = round_pct(total_values.get("selection_effect"))
    interaction = round_pct(total_values.get("interaction_effect"))
    total_active = round_pct(total_values.get("total_effect"))
    driver, confidence = resolve_driver(allocation, selection, interaction)
    return {
        "allocation_effect": allocation,
        "allocation_effect_bps": to_bps(allocation),
        "selection_effect": selection,
        "selection_effect_bps": to_bps(selection),
        "interaction_effect": interaction,
        "interaction_effect_bps": to_bps(interaction),
        "total_active_return": total_active,
        "total_active_return_bps": to_bps(total_active),
        "primary_driver": driver,
        # driver_label collapses primary_driver and the confidence behind it
        # into the phrase to use. driver_confidence itself is deliberately not
        # emitted: with the value visible, four of five drafts wrote something
        # like "selection at high confidence" despite an explicit ban in two
        # places. Removing the value removes the temptation.
        "driver_label": resolve_driver_label(driver, confidence),
    }


def build_sector_attribution(sector_rows: Sequence[RawRow]) -> List[Dict[str, Any]]:
    """One record per GICS sector, ranked by absolute total effect."""
    records: List[Dict[str, Any]] = []
    for raw in sector_rows:
        values = raw.values
        allocation = round_pct(values.get("allocation_effect"))
        selection = round_pct(values.get("selection_effect"))
        interaction = round_pct(values.get("interaction_effect"))
        total = round_pct(values.get("total_effect"))
        portfolio_weight = round_pct(values.get("portfolio_weight"))
        active_weight = round_pct(values.get("active_weight"))
        driver, _ = resolve_driver(allocation, selection, interaction)

        records.append(
            {
                "sector": GICS_SECTOR_LOOKUP.get(
                    normalize_header(raw.section_label), raw.section_label
                ),
                "portfolio_weight": portfolio_weight,
                "benchmark_weight": round_pct(values.get("benchmark_weight")),
                "active_weight": active_weight,
                "portfolio_return": round_pct(values.get("portfolio_return")),
                "benchmark_return": round_pct(values.get("benchmark_return")),
                "return_differential": round_pct(values.get("return_differential")),
                "allocation_effect": allocation,
                "allocation_effect_bps": to_bps(allocation),
                "allocation_direction": effect_direction(allocation),
                "selection_effect": selection,
                "selection_effect_bps": to_bps(selection),
                "selection_direction": effect_direction(selection),
                "interaction_effect": interaction,
                "interaction_effect_bps": to_bps(interaction),
                "total_effect": total,
                "total_effect_bps": to_bps(total),
                "position": resolve_position(portfolio_weight, active_weight),
                "driver": driver,
                "effect_rank": 0,
            }
        )

    ranked = sorted(
        records,
        key=lambda record: abs(record["total_effect"] or 0.0),
        reverse=True,
    )
    for position, record in enumerate(ranked, start=1):
        record["effect_rank"] = position
    return records


def build_cash_attribution(bucket_rows: Sequence[RawRow]) -> Dict[str, Any]:
    """
    Cash allocation effect, reported separately from sector_attribution.

    Cash is not a GICS sector so it never joins the sector table, but holding
    cash in a rising market is a real reason a manager lagged. `material` flags
    whether the effect clears the same threshold used for securities, so the
    commentary agent can ignore trivial cash drag without judging it.
    """
    cash_row = next(
        (
            row
            for row in bucket_rows
            if normalize_header(row.section_label) == CASH_ROW_LABEL
        ),
        None,
    )
    if cash_row is None:
        return {
            "available": False,
            "allocation_effect": None,
            "allocation_effect_bps": None,
            "material": False,
        }

    allocation = round_pct(cash_row.values.get("allocation_effect"))
    return {
        "available": True,
        "allocation_effect": allocation,
        "allocation_effect_bps": to_bps(allocation),
        "material": allocation is not None
        and abs(allocation) >= MATERIALITY_THRESHOLD_PP,
    }


def build_benchmark_sector_context(
    sector_rows: Sequence[RawRow], total_values: Dict[str, Optional[float]]
) -> Dict[str, Any]:
    """
    The benchmark's own GICS sector composition, for the market backdrop.

    The attribution workbook's benchmark columns describe the index, not the
    manager, so they supply the sector detail the market backdrop needs and that
    the sector/industry/factor export cannot: all eleven GICS sectors (the market
    export carries nine Russell-scheme sectors, without Communication Services or
    Real Estate), plus each sector's index weight and its contribution to the
    index return.
    """
    benchmark_return = round_pct(total_values.get("benchmark_return"))
    sectors: List[Dict[str, Any]] = []

    for raw in sector_rows:
        values = raw.values
        contribution = round_pct(values.get("contribution_benchmark"))
        sectors.append(
            {
                "sector": GICS_SECTOR_LOOKUP.get(
                    normalize_header(raw.section_label), raw.section_label
                ),
                "benchmark_weight": round_pct(values.get("benchmark_weight")),
                "benchmark_return": round_pct(values.get("benchmark_return")),
                # No _bps twin on purpose. House style quotes index sector
                # contributions in percentage points, but the general "use the
                # _bps fields" convention kept pulling the agent to basis points
                # whenever one was available. Removing it settles the conflict.
                "contribution_to_benchmark_return": contribution,
            }
        )

    ranked = sorted(
        (s for s in sectors if s["benchmark_return"] is not None),
        key=lambda s: s["benchmark_return"],
        reverse=True,
    )
    contributors = sorted(
        (s for s in sectors if s["contribution_to_benchmark_return"] is not None),
        key=lambda s: s["contribution_to_benchmark_return"],
        reverse=True,
    )

    return {
        "benchmark_return": benchmark_return,
        "sectors_positive": sum(
            1 for s in sectors if (s["benchmark_return"] or 0) > 0
        ),
        "sectors_negative": sum(
            1 for s in sectors if (s["benchmark_return"] or 0) < 0
        ),
        "sector_count": len(sectors),
        "best_sector": ranked[0]["sector"] if ranked else None,
        "worst_sector": ranked[-1]["sector"] if ranked else None,
        "largest_contributor": contributors[0]["sector"] if contributors else None,
        "largest_detractor": contributors[-1]["sector"] if contributors else None,
        "sectors": sectors,
    }


def build_security_records(
    security_rows: Sequence[RawRow], benchmark_total_return: Optional[float]
) -> List[Dict[str, Any]]:
    """Every security inside a GICS sector, before materiality filtering."""
    records: List[Dict[str, Any]] = []
    for raw in security_rows:
        values = raw.values
        portfolio_weight = round_pct(values.get("portfolio_weight"))
        benchmark_weight = round_pct(values.get("benchmark_weight"))
        active_weight = round_pct(values.get("active_weight"))
        portfolio_return = round_pct(values.get("portfolio_return"))
        benchmark_return = round_pct(values.get("benchmark_return"))
        selection = round_pct(values.get("selection_effect"))
        # Securities carry selection only; Morningstar repeats it as Active Ret%.
        total = round_pct(values.get("total_effect"))
        if total is None:
            total = selection
        if selection is None:
            selection = total

        held = portfolio_weight is not None and portfolio_weight > HELD_WEIGHT_EPSILON
        security_return = portfolio_return if held and portfolio_return is not None else benchmark_return

        records.append(
            {
                "security": raw.security_name,
                "ticker": raw.ticker or "",
                "sector": GICS_SECTOR_LOOKUP.get(
                    normalize_header(raw.section_label), raw.section_label
                ),
                "held": held,
                "portfolio_weight": portfolio_weight,
                "benchmark_weight": benchmark_weight,
                "active_weight": active_weight,
                "portfolio_return": portfolio_return,
                "benchmark_return": benchmark_return,
                "contribution_portfolio": round_pct(values.get("contribution_portfolio")),
                "contribution_benchmark": round_pct(values.get("contribution_benchmark")),
                "contribution_active": round_pct(values.get("contribution_active")),
                "selection_effect": selection,
                "selection_effect_bps": to_bps(selection),
                "total_effect": total,
                "total_effect_bps": to_bps(total),
                "position_label": resolve_position_label(
                    held, benchmark_weight, active_weight
                ),
                "reason": resolve_security_reason(
                    held, active_weight, security_return, benchmark_total_return, total
                ),
                "source_row": raw.row,
            }
        )
    return records


def filter_material_securities(
    securities: Sequence[Dict[str, Any]], keep_names: Iterable[Tuple[str, str]]
) -> List[Dict[str, Any]]:
    """
    Keep securities whose effect reaches the materiality threshold.

    Names appearing in the top contributor / detractor lists are kept regardless,
    so those lists can never reference a security absent from the file.
    """
    protected = set(keep_names)
    kept: List[Dict[str, Any]] = []
    for record in securities:
        effect = record.get("total_effect")
        identity = (record.get("security") or "", record.get("ticker") or "")
        material = effect is not None and abs(effect) >= MATERIALITY_THRESHOLD_PP
        if material or identity in protected:
            kept.append(record)
    kept.sort(key=lambda record: abs(record.get("total_effect") or 0.0), reverse=True)
    return kept


def build_top_movers(
    securities: Sequence[Dict[str, Any]], contributors: bool
) -> List[Dict[str, Any]]:
    """Top N contributors (positive effect) or detractors (negative effect)."""
    pool = [
        record
        for record in securities
        if record.get("total_effect") is not None
        and (record["total_effect"] > 0 if contributors else record["total_effect"] < 0)
    ]
    pool.sort(key=lambda record: record["total_effect"], reverse=contributors)

    movers: List[Dict[str, Any]] = []
    for position, record in enumerate(pool[:TOP_N_SECURITIES], start=1):
        movers.append(
            {
                "rank": position,
                "security": record["security"],
                "ticker": record["ticker"],
                "sector": record["sector"],
                "held": record["held"],
                "total_effect": record["total_effect"],
                "total_effect_bps": record["total_effect_bps"],
                "active_weight": record["active_weight"],
                "portfolio_return": record["portfolio_return"],
                "benchmark_return": record["benchmark_return"],
                # position_label replaces reason in the agent-facing lists on
                # purpose: carrying both put a correct phrase next to one that
                # contradicts it, and the contradicting one kept winning.
                "position_label": record["position_label"],
            }
        )
    return movers


def build_sector_rankings(sectors: Sequence[Dict[str, Any]]) -> Dict[str, Any]:
    """Pre-sorted sector leaders and laggards, in basis points."""

    def by(field_name: str, best: bool) -> List[Dict[str, Any]]:
        pool = [s for s in sectors if s.get(field_name) is not None]
        pool.sort(key=lambda s: s[field_name], reverse=best)
        return pool

    def brief(record: Optional[Dict[str, Any]], field_name: str) -> Optional[Dict[str, Any]]:
        if record is None:
            return None
        return {
            "sector": record["sector"],
            field_name: record[field_name],
            "position": record["position"],
        }

    contributing = [
        {
            "sector": s["sector"],
            "total_effect_bps": s["total_effect_bps"],
            "driver": s["driver"],
            "position": s["position"],
        }
        for s in by("total_effect_bps", True)[:TOP_N_SECTORS]
        if (s["total_effect_bps"] or 0) > 0
    ]
    detracting = [
        {
            "sector": s["sector"],
            "total_effect_bps": s["total_effect_bps"],
            "driver": s["driver"],
            "position": s["position"],
        }
        for s in by("total_effect_bps", False)[:TOP_N_SECTORS]
        if (s["total_effect_bps"] or 0) < 0
    ]

    selection_ranked = by("selection_effect_bps", True)
    allocation_ranked = by("allocation_effect_bps", True)

    return {
        "top_contributing_sectors": contributing,
        "top_detracting_sectors": detracting,
        "strongest_sector_selection": brief(
            selection_ranked[0] if selection_ranked else None, "selection_effect_bps"
        ),
        "weakest_sector_selection": brief(
            selection_ranked[-1] if selection_ranked else None, "selection_effect_bps"
        ),
        "strongest_sector_allocation": brief(
            allocation_ranked[0] if allocation_ranked else None, "allocation_effect_bps"
        ),
        "weakest_sector_allocation": brief(
            allocation_ranked[-1] if allocation_ranked else None, "allocation_effect_bps"
        ),
    }


def build_period_record(
    band: Optional[PeriodBand], values: Optional[Dict[str, Optional[float]]]
) -> Dict[str, Any]:
    """One entry of attribution_periods."""
    if band is None or values is None:
        return {
            "label": "",
            "period_start_date": "",
            "period_end_date": "",
            "allocation_effect": None,
            "allocation_effect_bps": None,
            "selection_effect": None,
            "selection_effect_bps": None,
            "interaction_effect": None,
            "interaction_effect_bps": None,
            "total_active_return": None,
            "total_active_return_bps": None,
            "available": False,
        }
    allocation = round_pct(values.get("allocation_effect"))
    selection = round_pct(values.get("selection_effect"))
    interaction = round_pct(values.get("interaction_effect"))
    total = round_pct(values.get("total_effect"))
    return {
        "label": band.label,
        "period_start_date": iso_or_empty(band.start_date),
        "period_end_date": iso_or_empty(band.end_date),
        "allocation_effect": allocation,
        "allocation_effect_bps": to_bps(allocation),
        "selection_effect": selection,
        "selection_effect_bps": to_bps(selection),
        "interaction_effect": interaction,
        "interaction_effect_bps": to_bps(interaction),
        "total_active_return": total,
        "total_active_return_bps": to_bps(total),
        "available": total is not None,
    }


def build_attribution_trends(periods: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
    """
    Streaks and driver persistence across the quarter chain.

    Streaks walk newest to oldest and stop at the first quarter that breaks the
    run or has no data, matching the convention used by performance_trends in
    Phase 1.
    """

    def effect_of(period_key: str, field_name: str) -> Optional[float]:
        return periods.get(period_key, {}).get(field_name)

    def streak(field_name: str, negative: bool) -> int:
        count = 0
        for period_key in QUARTER_TREND_SEQUENCE:
            value = effect_of(period_key, field_name)
            if value is None:
                break
            if (value < 0) if negative else (value > 0):
                count += 1
            else:
                break
        return count

    drivers: Dict[str, str] = {}
    for period_key in QUARTER_TREND_SEQUENCE:
        record = periods.get(period_key, {})
        if not record.get("available"):
            continue
        driver, _ = resolve_driver(
            record.get("allocation_effect"),
            record.get("selection_effect"),
            record.get("interaction_effect"),
        )
        drivers[period_key] = driver

    totals = {"allocation": 0.0, "selection": 0.0, "interaction": 0.0}
    for period_key in QUARTER_TREND_SEQUENCE:
        record = periods.get(period_key, {})
        if not record.get("available"):
            continue
        for name, field_name in (
            ("allocation", "allocation_effect"),
            ("selection", "selection_effect"),
            ("interaction", "interaction_effect"),
        ):
            value = record.get(field_name)
            if value is not None:
                totals[name] += abs(value)

    if drivers:
        trailing_driver, _ = resolve_driver(
            totals["allocation"], totals["selection"], totals["interaction"]
        )
    else:
        trailing_driver = "mixed"

    distinct = {d for d in drivers.values() if d != "mixed"}
    if len(drivers) < 2:
        stability = "insufficient_data"
    elif len(distinct) <= 1:
        stability = "consistent"
    else:
        stability = "variable"

    return {
        "consecutive_quarters_selection_negative": streak("selection_effect", True),
        "consecutive_quarters_selection_positive": streak("selection_effect", False),
        "consecutive_quarters_allocation_negative": streak("allocation_effect", True),
        "consecutive_quarters_allocation_positive": streak("allocation_effect", False),
        "dominant_driver_selected_quarter": drivers.get("selected_quarter", "mixed"),
        "dominant_driver_trailing_4_quarters": trailing_driver,
        "driver_stability": stability,
        "quarters_available": len(drivers),
    }


def build_concentration(
    securities: Sequence[Dict[str, Any]], sectors: Sequence[Dict[str, Any]]
) -> Dict[str, Any]:
    """
    Concentrated or broad-based?

    Computed over the FULL security universe inside GICS sectors, before
    materiality filtering, so the denominators are honest.
    """
    effects = [
        record["total_effect"]
        for record in securities
        if record.get("total_effect") is not None
    ]
    positives = sorted((e for e in effects if e > 0), reverse=True)
    negatives = sorted((e for e in effects if e < 0))

    def share(part: float, whole: float) -> Optional[float]:
        return round_share(part / whole) if whole else None

    sector_effects = [
        s["total_effect"] for s in sectors if s.get("total_effect") is not None
    ]

    contributors_share = share(sum(positives[:CONCENTRATION_TOP_N]), sum(positives))
    detractors_share = share(
        abs(sum(negatives[:CONCENTRATION_TOP_N])), abs(sum(negatives))
    )

    def character(value: Optional[float]) -> Optional[str]:
        """Concentrated or broad-based, decided here rather than by the agent."""
        if value is None:
            return None
        return (
            "concentrated"
            if value >= CONCENTRATION_SHARE_THRESHOLD
            else "broad_based"
        )

    # Only the verdicts ship. The underlying shares and security counts are
    # precisely what OUTPUT forbids quoting - "do not quote shares of total
    # effect or counts of securities" - and leaving them visible produced
    # "the top five detractors accounting for roughly 31% of negative effect"
    # twice in one draft. The share is what the character label is FOR; the
    # agent never needs both. Sector counts stay: those are legitimate
    # commentary material and have been used correctly.
    return {
        "contributors_character": character(contributors_share),
        "detractors_character": character(detractors_share),
        "number_of_sectors_positive": sum(1 for e in sector_effects if e > 0),
        "number_of_sectors_negative": sum(1 for e in sector_effects if e < 0),
    }


# =============================================================================
# SECTION 9 - RECONCILIATION (DEBUG / INTERNAL QA ONLY)
# -----------------------------------------------------------------------------
# None of this reaches the battle-book JSON. It exists so a build can be trusted
# without the commentary agent ever being tempted to write about residuals,
# coverage or expense ratios.
# =============================================================================


def build_reconciliation(
    sectors: Sequence[Dict[str, Any]],
    securities: Sequence[Dict[str, Any]],
    bucket_rows: Sequence[RawRow],
    total_values: Dict[str, Optional[float]],
    control_values: Dict[str, Dict[str, Optional[float]]],
    manager_ref: Optional[ManagerRef],
) -> Dict[str, Any]:
    sector_sum = sum(s["total_effect"] or 0.0 for s in sectors)
    bucket_sum = sum(row.values.get("total_effect") or 0.0 for row in bucket_rows)
    security_sum = sum(s["total_effect"] or 0.0 for s in securities)

    attribution_total = total_values.get("total_effect")
    selection_total = total_values.get("selection_effect")
    coverage = total_values.get("portfolio_weight")

    reported = control_values.get(ROW_REPORTED_TOTAL, {})
    expense = control_values.get(ROW_EXPENSE_RATIO, {})
    residual = control_values.get("residual", {})

    reported_portfolio = reported.get("portfolio_return")
    reported_benchmark = reported.get("benchmark_return")
    reported_excess = (
        reported_portfolio - reported_benchmark
        if reported_portfolio is not None and reported_benchmark is not None
        else None
    )

    def close(left: Optional[float], right: Optional[float]) -> Optional[bool]:
        if left is None or right is None:
            return None
        return abs(left - right) <= RECONCILE_TOLERANCE_PP

    return {
        "sector_total_effect_sum_gics_only": round_pct(sector_sum),
        "attributed_bucket_effect_sum": round_pct(bucket_sum),
        "sector_plus_bucket_sum": round_pct(sector_sum + bucket_sum),
        "attribution_total_active_return": round_pct(attribution_total),
        "sector_sum_matches_attribution_total": close(
            sector_sum + bucket_sum, attribution_total
        ),
        "security_total_effect_sum": round_pct(security_sum),
        "sector_level_selection_total": round_pct(selection_total),
        "security_sum_matches_selection_total": close(security_sum, selection_total),
        "attribution_coverage_percent": round_pct(coverage),
        "reported_portfolio_return": round_pct(reported_portfolio),
        "reported_benchmark_return": round_pct(reported_benchmark),
        "reported_excess_return": round_pct(reported_excess),
        "phase1_excess_return": manager_ref.excess_return if manager_ref else None,
        "attribution_vs_reported_gap": round_pct(
            reported_excess - attribution_total
            if reported_excess is not None and attribution_total is not None
            else None
        ),
        "expense_ratio": round_pct(expense.get("portfolio_return")),
        "residual_portfolio": round_pct(residual.get("portfolio_return")),
        "residual_benchmark": round_pct(residual.get("benchmark_return")),
        "attribution_benchmark_return": round_pct(total_values.get("benchmark_return")),
        "note": (
            "Internal QA only. None of these values appear in the battle-book "
            "attribution JSON by design."
        ),
    }


# =============================================================================
# SECTION 10 - WORKBOOK PARSING
# =============================================================================


def find_sheet(workbook, aliases: Sequence[str]) -> Optional[Worksheet]:
    for name in workbook.sheetnames:
        if normalize_header(name) in aliases:
            return workbook[name]
    return None


def parse_template_metadata(worksheet: Optional[Worksheet]) -> Dict[str, Any]:
    """Read the Template sheet manifest: portfolio, benchmark, model, dates."""
    metadata: Dict[str, Any] = {
        "name": "",
        "portfolio": "",
        "benchmark": "",
        "currency": "",
        "exported_at": "",
        "classification": "",
        "model": "",
    }
    if worksheet is None:
        return metadata

    for row in range(1, min(worksheet.max_row, 60) + 1):
        for column in range(1, min(worksheet.max_column, 6) + 1):
            text = cell_text(worksheet, row, column)
            if not text:
                continue
            if ":" in text:
                label, _, value = text.partition(":")
                key = METADATA_KEYS.get(normalize_header(label))
                if key and not metadata[key]:
                    metadata[key] = value.strip()
            match = CLASSIFICATION_HINT_RE.match(text)
            if match and not metadata["classification"]:
                candidate = match.group("value")
                if "sector" in candidate.lower() or "industry" in candidate.lower():
                    metadata["classification"] = candidate
            if not metadata["model"] and ATTRIBUTION_MODEL_RE.search(text):
                metadata["model"] = "three_factor_brinson"
    return metadata


def parse_attribution_workbook(
    path: Path, quarter: Quarter
) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    """
    Parse one attribution workbook into its raw building blocks.

    Returns (parsed, debug). The caller links the manager and assembles the
    output document; this function only reads the workbook.
    """
    LOG.info("Parsing attribution workbook: %s", path.name)
    warnings: List[str] = []

    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet_names = list(workbook.sheetnames)
        template_sheet = find_sheet(workbook, TEMPLATE_SHEET_ALIASES)
        attribution_sheet = find_sheet(workbook, ATTRIBUTION_SHEET_ALIASES)
        if attribution_sheet is None:
            raise ValueError(
                f"No 'Attribution' worksheet found in {path.name}; "
                f"sheets present: {sheet_names}"
            )
        if template_sheet is None:
            warnings.append(
                f"{path.name}: no 'Template' worksheet; metadata will be read "
                f"from the Attribution sheet header instead."
            )

        metadata = parse_template_metadata(template_sheet)
        # Rows 1-5 of the Attribution sheet repeat the same manifest.
        if not metadata["portfolio"]:
            fallback = parse_template_metadata(attribution_sheet)
            for key, value in fallback.items():
                if value and not metadata[key]:
                    metadata[key] = value

        band_row, group_row, leaf_row = find_header_rows(attribution_sheet)
        first_data_row = leaf_row + 1
        bands = detect_period_bands(attribution_sheet, band_row, group_row, leaf_row)

        selected_band = select_band_for_quarter(bands, quarter)
        if selected_band is None:
            available = "; ".join(
                f"{b.label} [{iso_or_empty(b.start_date)} -> {iso_or_empty(b.end_date)}]"
                for b in bands
            )
            raise ValueError(
                f"{path.name}: no attribution period band matches {quarter.label} "
                f"({quarter.start_date} - {quarter.end_date}).\n"
                f"       Bands present: {available}"
            )

        missing_fields = [
            name
            for name in ("allocation_effect", "selection_effect", "interaction_effect")
            if name not in selected_band.field_columns
        ]
        if missing_fields:
            warnings.append(
                f"{path.name}: band {selected_band.label!r} is missing columns for "
                f"{missing_fields}; those effects will be null."
            )
        if selected_band.width != EXPECTED_BAND_WIDTH:
            warnings.append(
                f"{path.name}: band {selected_band.label!r} is {selected_band.width} "
                f"columns wide, expected {EXPECTED_BAND_WIDTH}. Columns were mapped "
                f"by header text, so this is informational."
            )

        rows, row_counts = scan_rows(attribution_sheet, first_data_row, selected_band)

        # Split the scan into the pieces the builders need.
        sector_rows = [r for r in rows if r.kind == ROW_KIND_GICS_SECTOR]
        bucket_rows = [r for r in rows if r.kind == ROW_KIND_ATTRIBUTED_BUCKET]
        security_rows_all = [r for r in rows if r.kind == ROW_KIND_SECURITY]
        security_rows_in_sectors = [
            r for r in security_rows_all if r.section_kind == ROW_KIND_GICS_SECTOR
        ]

        control_values: Dict[str, Dict[str, Optional[float]]] = {}
        for raw in rows:
            if raw.kind in (ROW_KIND_CONTROL, ROW_KIND_RESIDUAL):
                label = normalize_header(raw.section_label)
                key = "residual" if raw.kind == ROW_KIND_RESIDUAL else label
                control_values.setdefault(key, raw.values)

        total_values = control_values.get(ROW_ATTRIBUTION_TOTAL)
        if total_values is None:
            raise ValueError(
                f"{path.name}: no 'Attribution Total' row found. "
                f"The export template may have changed."
            )

        found_sectors = {
            GICS_SECTOR_LOOKUP.get(normalize_header(r.section_label)) for r in sector_rows
        }
        missing_sectors = [s for s in GICS_SECTORS if s not in found_sectors]
        if missing_sectors:
            warnings.append(
                f"{path.name}: GICS sectors absent from the workbook: {missing_sectors}."
            )

        # Prior quarters, resolved by date the same way as the selected quarter.
        period_bands: Dict[str, Optional[PeriodBand]] = {}
        period_values: Dict[str, Optional[Dict[str, Optional[float]]]] = {}
        period_debug: Dict[str, Any] = {}
        total_row_number = next(
            (
                r.row
                for r in rows
                if normalize_header(r.section_label) == ROW_ATTRIBUTION_TOTAL
                and r.kind == ROW_KIND_CONTROL
            ),
            None,
        )
        for period_key, offset in LOGICAL_PERIODS:
            target = quarter.shifted(offset)
            band = (
                selected_band if offset == 0 else select_band_for_quarter(bands, target)
            )
            period_bands[period_key] = band
            if band is not None and total_row_number is not None:
                period_values[period_key] = read_band_values(
                    attribution_sheet, total_row_number, band
                )
            else:
                period_values[period_key] = None
                if band is None:
                    warnings.append(
                        f"{path.name}: no attribution band for {target.label} "
                        f"('{period_key}'); that period will be null."
                    )
            period_debug[period_key] = {
                "target_quarter": target.label,
                "workbook_band": band.label if band else None,
                "columns": band.column_range if band else None,
                "resolved_by": "matched_by_date" if band else "not_found",
            }

        parsed = {
            "metadata": metadata,
            "selected_band": selected_band,
            "sector_rows": sector_rows,
            "bucket_rows": bucket_rows,
            "security_rows_in_sectors": security_rows_in_sectors,
            "security_rows_all": security_rows_all,
            "total_values": total_values,
            "control_values": control_values,
            "period_bands": period_bands,
            "period_values": period_values,
            "warnings": warnings,
        }

        debug = {
            "file": path.name,
            "absolute_path": str(path),
            "worksheets_found": sheet_names,
            "worksheet_parsed": attribution_sheet.title,
            "template_metadata": metadata,
            "header_rows": {
                "band_label_row": band_row,
                "metric_group_row": group_row,
                "leaf_header_row": leaf_row,
                "first_data_row": first_data_row,
            },
            "period_bands_detected": [band.describe() for band in bands],
            "selected_period_band": selected_band.describe(),
            "logical_period_map": period_debug,
            "logical_period_map_summary": [
                f"{key} -> {entry['workbook_band'] or 'NOT FOUND'}"
                for key, entry in period_debug.items()
            ],
            "row_classification_counts": row_counts,
            "sector_rows_found": len(sector_rows),
            "sector_names_found": sorted(x for x in found_sectors if x),
            "sectors_missing": missing_sectors,
            "attributed_bucket_rows": [
                {
                    "label": r.section_label,
                    "row": r.row,
                    "total_effect": round_pct(r.values.get("total_effect")),
                    "portfolio_weight": round_pct(r.values.get("portfolio_weight")),
                }
                for r in bucket_rows
            ],
            "security_rows_scanned_total": len(security_rows_all),
            "security_rows_in_gics_sectors": len(security_rows_in_sectors),
            "security_rows_outside_gics_sectors": len(security_rows_all)
            - len(security_rows_in_sectors),
            "warnings": warnings,
        }
        return parsed, debug
    finally:
        workbook.close()


# =============================================================================
# SECTION 11 - MANAGER LINKAGE AND DOCUMENT ASSEMBLY
# =============================================================================


def link_manager(
    portfolio_name: str, phase1: Phase1Index
) -> Tuple[Optional[ManagerRef], Optional[str]]:
    """
    Strict normalised match of the Template 'Portfolio:' value to a manager.

    No fuzzy matching. Anything other than exactly one match is an error: a
    silent mis-link would attach one manager's attribution to another, which is
    the worst failure this pipeline can produce.
    """
    lookup_key = normalize_lookup_name(portfolio_name)
    if not lookup_key:
        return None, "Template 'Portfolio:' value is empty; cannot link to a manager."

    matches, match_type = phase1.resolve(lookup_key)
    if not matches:
        return None, (
            f"No manager matches portfolio {portfolio_name!r} "
            f"(normalised {lookup_key!r}) in any asset-class JSON."
        )
    if len(matches) > 1:
        where = ", ".join(f"{m.asset_class_file}[{m.index}]" for m in matches)
        return None, (
            f"Portfolio {portfolio_name!r} matched {len(matches)} managers ({where}); "
            f"refusing to guess."
        )

    match = matches[0]
    match.match_type = match_type
    return match, None


def build_attribution_document(
    parsed: Dict[str, Any],
    manager_ref: ManagerRef,
    quarter: Quarter,
    path: Path,
    generated_at: str,
) -> Tuple[Dict[str, Any], Dict[str, Any], List[str]]:
    """Assemble the battle-book-facing attribution JSON for one manager."""
    warnings: List[str] = []
    metadata = parsed["metadata"]
    band: PeriodBand = parsed["selected_band"]
    total_values = parsed["total_values"]

    benchmark_total_return = total_values.get("benchmark_return")

    sectors = build_sector_attribution(parsed["sector_rows"])
    all_securities = build_security_records(
        parsed["security_rows_in_sectors"], benchmark_total_return
    )

    top_contributors = build_top_movers(all_securities, contributors=True)
    top_detractors = build_top_movers(all_securities, contributors=False)
    protected = [
        (record["security"] or "", record["ticker"] or "")
        for record in top_contributors + top_detractors
    ]
    stored_securities = filter_material_securities(all_securities, protected)
    for record in stored_securities:
        record.pop("source_row", None)

    periods = {
        period_key: build_period_record(
            parsed["period_bands"].get(period_key),
            parsed["period_values"].get(period_key),
        )
        for period_key, _ in LOGICAL_PERIODS
    }

    document = {
        "manager": manager_ref.manager,
        "manager_lookup_key": manager_ref.lookup_key,
        "asset_class": manager_ref.asset_class,
        "asset_class_key": manager_ref.asset_class_key,
        "benchmark": metadata.get("benchmark") or manager_ref.benchmark,
        "period": quarter.label,
        "period_start_date": quarter.start_date.isoformat(),
        "period_end_date": quarter.end_date.isoformat(),
        "period_band_label": band.label,
        "model": metadata.get("model") or "three_factor_brinson",
        "classification": metadata.get("classification") or "GICS Sector",
        "lineage": {
            "source_file": path.name,
            "worksheet": "Attribution",
            "exported_at": metadata.get("exported_at", ""),
            "generated_at": generated_at,
        },
        # Market backdrop input: the benchmark's own GICS sector composition.
        # Placed before the manager's attribution because commentary runs
        # market -> attribution -> sector -> security.
        "benchmark_sector_context": build_benchmark_sector_context(
            parsed["sector_rows"], total_values
        ),
        "attribution_summary": build_attribution_summary(total_values),
        "sector_attribution": sectors,
        # Cash sits outside the GICS sector table by design; see §5.3 of the README.
        "cash_attribution": build_cash_attribution(parsed["bucket_rows"]),
        "security_attribution": stored_securities,
        "security_attribution_meta": {
            "securities_in_gics_sectors": len(all_securities),
            "securities_held": sum(1 for s in all_securities if s["held"]),
            "securities_stored": len(stored_securities),
            "materiality_threshold_pp": MATERIALITY_THRESHOLD_PP,
            "materiality_threshold_bps": round(MATERIALITY_THRESHOLD_PP * 100, 1),
            "note": (
                "Securities carry a selection effect only; allocation and "
                "interaction exist at sector level. `held` distinguishes portfolio "
                "holdings from benchmark constituents the manager did not own."
            ),
        },
        "top_contributors": top_contributors,
        "top_detractors": top_detractors,
        "sector_rankings": build_sector_rankings(sectors),
        "attribution_periods": periods,
        "attribution_trends": build_attribution_trends(periods),
        "concentration": build_concentration(all_securities, sectors),
    }

    # Benchmark agreement is a linkage check, not commentary material.
    template_benchmark = (metadata.get("benchmark") or "").strip()
    if (
        template_benchmark
        and manager_ref.benchmark
        and normalize_lookup_name(template_benchmark)
        != normalize_lookup_name(manager_ref.benchmark)
    ):
        warnings.append(
            f"{path.name}: Template benchmark {template_benchmark!r} does not match "
            f"the Phase 1 benchmark {manager_ref.benchmark!r} for "
            f"{manager_ref.manager!r}."
        )

    reconciliation = build_reconciliation(
        sectors,
        all_securities,
        parsed["bucket_rows"],
        total_values,
        parsed["control_values"],
        manager_ref,
    )
    return document, reconciliation, warnings


# =============================================================================
# SECTION 11b - BATTLE BOOK PACK
# -----------------------------------------------------------------------------
# One self-contained file per manager, assembled here rather than at query time,
# so runtime retrieval is a single deterministic file read.
#
# Field names are deliberately kept flat and identical to the names used in the
# asset-class and attribution JSON, so the commentary agent's instructions work
# against a pack without any change of path.
# =============================================================================


def enrich_movers(
    movers: Sequence[Dict[str, Any]], securities: Sequence[Dict[str, Any]]
) -> List[Dict[str, Any]]:
    """Add weights and contribution to each top mover, from its security record."""
    by_identity = {
        (s.get("security"), s.get("ticker")): s for s in securities
    }
    enriched: List[Dict[str, Any]] = []
    for mover in movers:
        record = dict(mover)
        source = by_identity.get((mover.get("security"), mover.get("ticker")))
        if source:
            for field_name in PACK_MOVER_ENRICHMENT:
                if field_name in source:
                    record[field_name] = source[field_name]
        enriched.append(record)
    return enriched


PRIOR_RUN_QUARTERS = ("two_quarters_ago", "three_quarters_ago", "four_quarters_ago")
QUARTER_LABEL_KEYS = PRIOR_RUN_QUARTERS + ("selected_quarter",)


def with_quarter_labels(periods: Dict[str, Any]) -> Dict[str, Any]:
    """
    Name each quarter block by its calendar quarter.

    Phase 1 labels them by distance - "3 Quarters Ago" - and leaves the calendar
    quarter implicit in the dates. Asked which quarters had been strong, the
    agent worked back from the labels and named "Q3 and Q4 2025" when the strong
    ones were Q4 2025 and Q1 2026. The arithmetic is trivial and the dates were
    right there, which is exactly the kind of derivation that keeps going wrong.
    """
    labelled = dict(periods)
    for key in QUARTER_LABEL_KEYS:
        record = labelled.get(key)
        if not isinstance(record, dict):
            continue
        end = record.get("period_end_date")
        if not end:
            continue
        try:
            year, month = int(end[:4]), int(end[5:7])
        except (TypeError, ValueError):
            continue
        labelled[key] = {**record, "quarter": f"{year} Q{(month - 1) // 3 + 1}"}
    return labelled


def with_prior_run(
    trends: Dict[str, Any], periods: Dict[str, Any]
) -> Dict[str, Any]:
    """
    Record the run of quarters that ended with the selected one.

    consecutive_quarters_outperforming and its twin describe the run the
    selected quarter belongs to, so when a quarter breaks a streak both read as
    a small number and the streak it broke is recorded nowhere. Asked to
    describe a reversal, the agent counted the quarters itself and got it
    wrong - claiming T. Rowe Price Mid-Cap Growth had reversed three straight
    quarters of outperformance when the run was two, the fourth quarter back
    having trailed by 105 bps. Computing it here removes the counting.
    """
    if not periods:
        return dict(trends)

    def excess(key: str) -> Optional[float]:
        record = periods.get(key) or {}
        return record.get("excess_return_cumulative") if record.get("available") else None

    current = excess("selected_quarter")
    prior = [excess(key) for key in PRIOR_RUN_QUARTERS]
    if current is None or prior[0] is None:
        return dict(trends)

    # Only a reversal has a run worth naming; an unbroken streak is already
    # covered by consecutive_quarters_*.
    outperformed_before = prior[0] > 0
    if (current > 0) == outperformed_before:
        return dict(trends)

    quarters = 0
    for value in prior:
        if value is None or (value > 0) != outperformed_before:
            break
        quarters += 1

    enriched = dict(trends)
    enriched["prior_run"] = {
        "direction": "outperforming" if outperformed_before else "underperforming",
        "quarters": quarters,
    }
    return enriched


def strip_active_return_total(record: Dict[str, Any]) -> Dict[str, Any]:
    """
    Drop the Brinson total from anything the agent reads.

    The pack already carries excess_return_cumulative, the performance excess.
    total_active_return is the attribution model's own total and differs from it
    - 967 bps against 926.5 for Fidelity Q2 2026 - because attribution covers
    the securities it can classify. Both figures are correct and neither needs
    explaining, but a draft that quotes both leaves a consultant staring at a
    40 bps discrepancy. Instructing the agent to prefer one held on 4.8 and
    failed on 4.7, so the competing figure goes instead. It stays in the full
    attribution JSON and the debug report, which are the audit artifacts.
    """
    return {
        key: value
        for key, value in record.items()
        if key not in ("total_active_return", "total_active_return_bps")
    }


def material_cash_only(cash: Dict[str, Any]) -> Dict[str, Any]:
    """
    Put cash_attribution in the pack only when it is worth writing about.

    Shipping it with material=false invited "the cash position was not
    material" - a sentence about the data rather than the fund, and one that
    echoes a field name. Absent the key entirely there is nothing to report,
    and `material` itself never reaches the agent's vocabulary.
    """
    if not cash.get("material"):
        return {}
    return {
        "cash_attribution": {
            key: value for key, value in cash.items() if key != "material"
        }
    }


def build_pack(
    manager_record: Dict[str, Any],
    asset_class_document: Dict[str, Any],
    attribution: Dict[str, Any],
    quarter: Quarter,
    generated_at: str,
) -> Dict[str, Any]:
    """Assemble the single file a battle book is written from."""
    market_periods = asset_class_document.get("market_data_periods", {})
    extra_market = {
        f"market_summaries_{key}": {
            "label": market_periods.get(key, {}).get("label", ""),
            "period_start_date": market_periods.get(key, {}).get("period_start_date", ""),
            "period_end_date": market_periods.get(key, {}).get("period_end_date", ""),
            "aligned_with_selected_quarter_end": market_periods.get(key, {}).get(
                "aligned_with_selected_quarter_end", False
            ),
            "note": market_periods.get(key, {}).get("note", ""),
            "summaries": market_periods.get(key, {}).get("summaries", {}),
        }
        for key in PACK_EXTRA_MARKET_PERIODS
        # In Q1 the year-to-date window is the quarter itself, so an extra
        # market period would simply duplicate `summaries`.
        if key in market_periods and not (key == "ytd" and quarter.quarter == 1)
    }

    pack = {
        "pack_version": SCRIPT_VERSION,
        "manager": manager_record.get("manager", ""),
        "manager_lookup_key": normalize_lookup_name(manager_record.get("manager")),
        "strategy_name": manager_record.get("strategy_name", ""),
        "ticker": manager_record.get("ticker", ""),
        "portfolio_managers": manager_record.get("portfolio_managers", []),
        "asset_class": attribution.get("asset_class", ""),
        "asset_class_key": attribution.get("asset_class_key", ""),
        "benchmark": manager_record.get("benchmark", ""),
        "period": quarter.label,
        "period_start_date": quarter.start_date.isoformat(),
        "period_end_date": quarter.end_date.isoformat(),
        "lineage": {
            "performance_source": manager_record.get("source_file", ""),
            "attribution_source": attribution.get("lineage", {}).get("source_file", ""),
            "attribution_exported_at": attribution.get("lineage", {}).get(
                "exported_at", ""
            ),
            "generated_at": generated_at,
        },
        # --- headline performance -------------------------------------------
        "return_cumulative": manager_record.get("return_cumulative"),
        "benchmark_return": manager_record.get("benchmark_return"),
        "benchmark_return_calculated": manager_record.get(
            "benchmark_return_calculated", False
        ),
        "peer_percentile": manager_record.get("peer_percentile"),
        "excess_return_cumulative": manager_record.get("excess_return_cumulative"),
        # --- history ---------------------------------------------------------
        "performance_periods": with_quarter_labels(
            manager_record.get("performance_periods", {})
        ),
        "performance_trends": with_prior_run(
            manager_record.get("performance_trends", {}),
            manager_record.get("performance_periods", {}),
        ),
        "ranking_trend": manager_record.get("ranking_trend", {}),
        # --- market backdrop --------------------------------------------------
        "reference_indexes": asset_class_document.get("reference_indexes", []),
        "peer_group_stats": asset_class_document.get("peer_group_stats", {}),
        "summaries": asset_class_document.get("summaries", {}),
        "market_trends": asset_class_document.get("market_trends", {}),
        **extra_market,
        # --- attribution ------------------------------------------------------
        "benchmark_sector_context": attribution.get("benchmark_sector_context", {}),
        "attribution_summary": strip_active_return_total(
            attribution.get("attribution_summary", {})
        ),
        "sector_attribution": attribution.get("sector_attribution", []),
        **material_cash_only(attribution.get("cash_attribution", {})),
        "sector_rankings": attribution.get("sector_rankings", {}),
        "top_contributors": enrich_movers(
            attribution.get("top_contributors", []),
            attribution.get("security_attribution", []),
        ),
        "top_detractors": enrich_movers(
            attribution.get("top_detractors", []),
            attribution.get("security_attribution", []),
        ),
        "attribution_periods": {
            key: strip_active_return_total(record)
            for key, record in attribution.get("attribution_periods", {}).items()
        },
        "attribution_trends": attribution.get("attribution_trends", {}),
        "concentration": attribution.get("concentration", {}),
    }
    return pack


def collect_lookup_keys(
    asset_class_document: Dict[str, Any], manager_index: int
) -> List[str]:
    """
    Every key that resolves to this manager - fund name, ticker, strategy name.

    Taken from the Phase 1 manager_lookup so a pack is reachable by anything the
    asset-class file already answers to, not just the exact fund name.
    """
    keys = [
        key
        for key, index in asset_class_document.get("manager_lookup", {}).items()
        if index == manager_index
    ]
    return sorted(set(keys))


def strip_share_class(name: Any) -> str:
    """Normalised fund name with its share-class suffix removed."""
    current = normalize_lookup_name(name)
    for _ in range(SHARE_CLASS_STRIP_ROUNDS):
        stripped = SHARE_CLASS_SUFFIX_RE.sub("", current).strip()
        if not stripped or stripped == current:
            break
        current = stripped
    return current


PUNCTUATION_VARIANT_RE = re.compile(r"[-/]")


def punctuation_variants(normalized: str) -> List[str]:
    """
    Spellings of a normalized name a user is likely to type instead.

    The lookup key strips periods and collapses spaces but leaves hyphens and
    slashes intact, so Morningstar's "T. Rowe Price Mid-Cap Growth" resolves
    only for someone who types the hyphen. Asking for "mid cap growth" returned
    not_found on a manager whose pack was sitting in the folder. Both the
    spaced and the joined form are generated, since "Mid Cap" and "MidCap" are
    equally plausible typings.
    """
    if not PUNCTUATION_VARIANT_RE.search(normalized):
        return []
    spaced = re.sub(
        r"\s+", " ", PUNCTUATION_VARIANT_RE.sub(" ", normalized)
    ).strip()
    joined = PUNCTUATION_VARIANT_RE.sub("", normalized)
    return sorted({spaced, joined} - {normalized, ""})


def build_share_class_aliases(
    phase1: "Phase1Index",
) -> Tuple[Dict[str, str], Dict[str, List[str]]]:
    """
    Map alternative spellings of a fund name onto the one manager they can mean.

    Covers two kinds: share-class suffixes stripped off ("Fidelity Blue Chip
    Growth" for the K shares) and punctuation the user is unlikely to reproduce
    ("mid cap" for "Mid-Cap"). Both are generated for the suffix-stripped form
    too, so "t rowe price mid cap growth" resolves even when the fund carries a
    share class.

    The check runs across every manager in every asset class, not just those
    with a pack, so an alias is only created when it is unambiguous in the whole
    universe. Returns (safe aliases, rejected ambiguous names).
    """
    candidates: Dict[str, List[str]] = {}
    exact: set = set()

    for document in phase1.documents.values():
        for manager in document.get("managers", []):
            name = manager.get("manager")
            if not name:
                continue
            full = normalize_lookup_name(name)
            exact.add(full)
            forms = {full}
            stripped = strip_share_class(name)
            if stripped:
                forms.add(stripped)
            for form in list(forms):
                forms.update(punctuation_variants(form))
            for form in forms:
                if form and form != full:
                    candidates.setdefault(form, []).append(full)

    aliases: Dict[str, str] = {}
    ambiguous: Dict[str, List[str]] = {}
    for stripped, owners in candidates.items():
        unique = sorted(set(owners))
        if len(unique) > 1:
            ambiguous[stripped] = unique
        elif stripped in exact:
            # The stripped form is itself another fund's full name; leave it
            # pointing at that fund.
            continue
        else:
            aliases[stripped] = unique[0]
    return aliases, ambiguous


def build_pack_manifest(
    entries: Sequence[Dict[str, Any]],
    quarter: Quarter,
    generated_at: str,
    share_class_aliases: Optional[Dict[str, str]] = None,
    ambiguous_names: Optional[Dict[str, List[str]]] = None,
) -> Dict[str, Any]:
    """
    Resolve any normalised manager key to a pack path.

    Only managers with a pack appear, so absence from the manifest carries the
    same meaning as a missing pack: no attribution was supplied, so no battle
    book. `packs` is a flat dictionary for a single-step lookup by callers that
    cannot easily search.
    """
    packs: Dict[str, str] = {}
    collisions: List[str] = []
    for entry in entries:
        for key in entry["keys"]:
            if key in packs and packs[key] != entry["path"]:
                collisions.append(key)
                continue
            packs[key] = entry["path"]

    # Suffix-stripped names, added only where they resolve to one manager and do
    # not overwrite a key that already resolves exactly.
    derived: Dict[str, str] = {}
    by_full_name = {
        normalize_lookup_name(entry["manager"]): entry["path"] for entry in entries
    }
    for alias, full_name in (share_class_aliases or {}).items():
        path = by_full_name.get(full_name)
        if path and alias not in packs:
            packs[alias] = path
            derived[alias] = full_name

    return {
        "period": quarter.label,
        "generated_at": generated_at,
        "pack_count": len(entries),
        "packs": dict(sorted(packs.items())),
        "share_class_aliases": dict(sorted(derived.items())),
        "ambiguous_names": dict(sorted((ambiguous_names or {}).items())),
        "managers": [
            {
                "manager": entry["manager"],
                "asset_class": entry["asset_class"],
                "path": entry["path"],
                "keys": entry["keys"],
            }
            for entry in entries
        ],
        "key_collisions": sorted(set(collisions)),
        "note": (
            "Normalise the requested manager name (lowercase, remove periods, "
            "collapse repeated spaces, trim) and look it up in `packs`. Do not "
            "rebuild the filename from the manager name; slugs strip characters "
            "that normalisation keeps. `packs` also accepts fund names without "
            "their share-class suffix where that is unambiguous; names listed in "
            "`ambiguous_names` are deliberately absent and should be asked about."
        ),
    }


# =============================================================================
# SECTION 12 - PATCHING PHASE 1 ASSET-CLASS FILES
# =============================================================================


def patch_asset_class_files(
    phase1: Phase1Index,
    linked: Dict[str, Dict[str, Any]],
    output_folder: Path,
) -> Dict[str, Any]:
    """
    Add attribution pointers to the Phase 1 asset-class JSONs.

    Additive and idempotent: every manager gains `attribution_available` and
    `attribution_path`, and each document gains an `attribution_index`. Nothing
    else is touched, so Phase 1 output is not degraded.
    """
    summary: Dict[str, Any] = {}

    for asset_class_key, filename in ASSET_CLASS_FILES:
        document = phase1.documents.get(asset_class_key)
        path = phase1.paths.get(asset_class_key)
        if document is None or path is None:
            continue

        available_keys: List[str] = []
        slugs: Dict[str, str] = {}
        patched = 0

        for manager in document.get("managers", []):
            key = normalize_lookup_name(manager.get("manager"))
            entry = linked.get(key)
            if entry is not None and entry["asset_class_key"] == asset_class_key:
                manager["attribution_available"] = True
                manager["attribution_path"] = entry["relative_path"]
                available_keys.append(key)
                slugs[key] = entry["slug"]
                patched += 1
            else:
                manager["attribution_available"] = False
                manager["attribution_path"] = None

        document["attribution_index"] = {
            "available": available_keys,
            "count": len(available_keys),
            "folder": ATTRIBUTION_OUTPUT_SUBFOLDER,
            "path_pattern": f"{ATTRIBUTION_OUTPUT_SUBFOLDER}/{{manager_slug}}.json",
            "manager_slugs": slugs,
        }

        write_json(path, document)
        summary[asset_class_key] = {
            "file": filename,
            "managers_with_attribution": patched,
            "managers_total": len(document.get("managers", [])),
        }
    return summary


# =============================================================================
# SECTION 13 - ORCHESTRATION
# =============================================================================


def build_attribution(
    input_root: Path,
    output_root: Path,
    requested_quarter: str,
    patch_phase1: bool = True,
) -> Dict[str, Any]:
    """Run the attribution build for one quarter. Returns the debug report."""
    quarter, quarter_folder = resolve_quarter(input_root, requested_quarter)
    LOG.info("Building attribution for %s", quarter.label)

    generated_at = now_iso()
    output_folder = output_root / quarter.label
    attribution_folder = output_folder / ATTRIBUTION_OUTPUT_SUBFOLDER
    pack_folder = output_folder / PACK_OUTPUT_SUBFOLDER
    input_folder = quarter_folder / ATTRIBUTION_INPUT_SUBFOLDER

    phase1 = load_phase1_index(output_folder)
    LOG.info(
        "Loaded Phase 1 index: %d managers across %d asset classes",
        sum(len(refs) for refs in phase1.by_manager_name.values()),
        len(phase1.documents),
    )

    global_warnings: List[str] = []
    errors: List[Dict[str, Any]] = []
    workbook_reports: List[Dict[str, Any]] = []
    linked: Dict[str, Dict[str, Any]] = {}
    files_written: List[str] = []
    packs_written: List[str] = []
    pack_entries: List[Dict[str, Any]] = []

    workbooks: List[Path] = []
    if not input_folder.is_dir():
        message = (
            f"No attribution folder at {input_folder}. Nothing to build; the "
            f"asset-class files will still be patched to record that no "
            f"attribution is available."
        )
        LOG.warning(message)
        global_warnings.append(message)
    else:
        workbooks = sorted(
            path
            for path in input_folder.glob("*.xls*")
            if not path.name.startswith("~$")
        )
        if not workbooks:
            message = f"Attribution folder {input_folder} contains no workbooks."
            LOG.warning(message)
            global_warnings.append(message)

    for path in workbooks:
        try:
            parsed, debug = parse_attribution_workbook(path, quarter)
        except Exception as error:  # noqa: BLE001 - report and continue
            message = f"Failed to parse {path.name}: {error}"
            LOG.error(message)
            errors.append({"file": path.name, "stage": "parse", "error": str(error)})
            workbook_reports.append({"file": path.name, "error": str(error)})
            continue

        portfolio = parsed["metadata"].get("portfolio", "")
        manager_ref, link_error = link_manager(portfolio, phase1)
        if manager_ref is None:
            message = f"{path.name}: {link_error}"
            LOG.error(message)
            errors.append(
                {
                    "file": path.name,
                    "stage": "manager_linkage",
                    "portfolio": portfolio,
                    "error": link_error,
                }
            )
            debug["manager_linkage"] = {
                "portfolio": portfolio,
                "normalized_key": normalize_lookup_name(portfolio),
                "linked": False,
                "error": link_error,
            }
            workbook_reports.append(debug)
            continue

        document, reconciliation, doc_warnings = build_attribution_document(
            parsed, manager_ref, quarter, path, generated_at
        )
        slug = manager_slug(manager_ref.manager)
        relative_path = f"{ATTRIBUTION_OUTPUT_SUBFOLDER}/{slug}.json"
        write_json(attribution_folder / f"{slug}.json", document)
        files_written.append(relative_path)

        # The battle book pack: everything for this manager in one file, so
        # runtime retrieval is a single deterministic read.
        asset_class_document = phase1.documents.get(manager_ref.asset_class_key, {})
        manager_record = asset_class_document.get("managers", [])[manager_ref.index]
        pack = build_pack(
            manager_record, asset_class_document, document, quarter, generated_at
        )
        pack_path = pack_folder / f"{slug}.json"
        write_json(pack_path, pack, compact=True)
        pack_relative = f"{PACK_OUTPUT_SUBFOLDER}/{slug}.json"
        packs_written.append(pack_relative)

        pack_entries.append({
            "manager": manager_ref.manager,
            "asset_class": manager_ref.asset_class,
            "path": pack_relative,
            "keys": collect_lookup_keys(asset_class_document, manager_ref.index),
        })

        pack_kb = pack_path.stat().st_size / 1024
        if pack_kb > PACK_SIZE_WARN_KB:
            message = (
                f"{pack_relative} is {pack_kb:.1f} KB, above the {PACK_SIZE_WARN_KB} KB "
                f"guideline. The whole pack is returned in one response, so check it "
                f"still fits the retrieval path."
            )
            LOG.warning(message)
            global_warnings.append(message)

        linked[manager_ref.lookup_key] = {
            "slug": slug,
            "relative_path": relative_path,
            "asset_class_key": manager_ref.asset_class_key,
        }

        debug["manager_linkage"] = {
            "portfolio": portfolio,
            "normalized_key": manager_ref.lookup_key,
            "linked": True,
            "match_type": manager_ref.match_type,
            "manager": manager_ref.manager,
            "asset_class": manager_ref.asset_class,
            "asset_class_file": manager_ref.asset_class_file,
            "manager_index": manager_ref.index,
            "benchmark_template": parsed["metadata"].get("benchmark", ""),
            "benchmark_phase1": manager_ref.benchmark,
            "benchmark_matches": normalize_lookup_name(
                parsed["metadata"].get("benchmark", "")
            )
            == normalize_lookup_name(manager_ref.benchmark),
        }
        debug["output_file"] = relative_path
        debug["pack_file"] = pack_relative
        debug["pack_size_kb"] = round(pack_kb, 1)
        debug["securities_stored"] = document["security_attribution_meta"][
            "securities_stored"
        ]
        debug["securities_filtered_out"] = (
            document["security_attribution_meta"]["securities_in_gics_sectors"]
            - document["security_attribution_meta"]["securities_stored"]
        )
        debug["reconciliation"] = reconciliation
        debug["warnings"] = list(debug.get("warnings", [])) + doc_warnings
        workbook_reports.append(debug)
        global_warnings.extend(parsed["warnings"])
        global_warnings.extend(doc_warnings)

    # Attribution workbooks arrive one battle book at a time, so the output
    # folder accumulates. Flag any attribution JSON with no matching workbook in
    # this run: the pointers will say it is unavailable while the file is still
    # on disk and retrievable. Reported, never deleted.
    orphans: List[str] = []
    if attribution_folder.is_dir():
        expected = {f"{entry['slug']}.json" for entry in linked.values()}
        for path in sorted(attribution_folder.glob("*.json")):
            if path.name not in expected:
                orphans.append(path.name)
        if orphans:
            message = (
                f"Attribution JSON files present with no matching workbook in "
                f"{input_folder}: {orphans}. They remain on disk but are marked "
                f"unavailable in the asset-class files. Keep the source workbooks "
                f"in the quarter folder, or remove these files."
            )
            LOG.warning(message)
            global_warnings.append(message)

    # Manifest of every key that resolves to a pack. Written whenever packs
    # were produced, so it can never describe a stale set.
    manifest_relative = ""
    if pack_entries:
        share_class_aliases, ambiguous_names = build_share_class_aliases(phase1)
        manifest = build_pack_manifest(
            pack_entries, quarter, generated_at, share_class_aliases, ambiguous_names
        )
        if ambiguous_names:
            message = (
                f"Share-class aliases not created for ambiguous names "
                f"{sorted(ambiguous_names)}; those funds must be asked for by "
                f"their full name."
            )
            LOG.warning(message)
            global_warnings.append(message)
        write_json(pack_folder / PACK_MANIFEST_FILENAME, manifest)
        manifest_relative = f"{PACK_OUTPUT_SUBFOLDER}/{PACK_MANIFEST_FILENAME}"
        if manifest["key_collisions"]:
            message = (
                f"Pack manifest: keys resolving to more than one manager "
                f"{manifest['key_collisions']}; the first pack wins. Ask for those "
                f"managers by their full fund name."
            )
            LOG.warning(message)
            global_warnings.append(message)

    patch_summary: Dict[str, Any] = {}
    if patch_phase1:
        patch_summary = patch_asset_class_files(phase1, linked, output_folder)
    else:
        LOG.info("--no-patch supplied; asset-class files left untouched.")

    report = {
        "script_version": SCRIPT_VERSION,
        "period": quarter.label,
        "period_start_date": quarter.start_date.isoformat(),
        "period_end_date": quarter.end_date.isoformat(),
        "generated_at": generated_at,
        "input_folder": str(input_folder),
        "output_folder": str(attribution_folder),
        "materiality_threshold_pp": MATERIALITY_THRESHOLD_PP,
        "workbooks_found": [p.name for p in workbooks],
        "workbooks_processed": len(files_written),
        "files_written": files_written,
        "packs_written": packs_written,
        "pack_manifest": manifest_relative,
        "pack_folder": str(pack_folder),
        "unmatched_attribution_files": [
            e["file"] for e in errors if e["stage"] == "manager_linkage"
        ],
        "orphaned_attribution_json": orphans,
        "errors": errors,
        "asset_class_patch_summary": patch_summary,
        "phase1_patched": patch_phase1,
        "workbooks": workbook_reports,
        "warnings": global_warnings,
    }
    write_json(output_folder / DEBUG_REPORT_FILENAME, report)
    return report


# =============================================================================
# SECTION 14 - CLI
# =============================================================================


def configure_logging(verbose: bool) -> None:
    logging.basicConfig(
        level=logging.DEBUG if verbose else logging.INFO,
        format="%(levelname)-8s %(message)s",
        stream=sys.stdout,
    )


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        prog="build_attribution_json.py",
        description=(
            "Convert Morningstar Direct attribution exports into per-manager JSON "
            "for the Battle Book commentary agent (Phase 2)."
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Examples:\n"
            '  python build_attribution_json.py --quarter "2026 Q2"\n'
            "  python build_attribution_json.py --quarter latest\n"
            '  python build_attribution_json.py --quarter "2026 Q2" --no-patch\n'
        ),
    )
    parser.add_argument(
        "--quarter",
        default="latest",
        help='Quarter to build, e.g. "2026 Q2", or "latest" (default: latest).',
    )
    parser.add_argument("--input-root", default=DEFAULT_INPUT_ROOT)
    parser.add_argument("--output-root", default=DEFAULT_OUTPUT_ROOT)
    parser.add_argument(
        "--no-patch",
        action="store_true",
        help="Do not write attribution pointers into the asset-class JSON files.",
    )
    parser.add_argument("--verbose", action="store_true", help="Enable debug logging.")
    return parser.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = parse_args(argv)
    configure_logging(args.verbose)

    try:
        report = build_attribution(
            Path(args.input_root),
            Path(args.output_root),
            args.quarter,
            patch_phase1=not args.no_patch,
        )
    except (FileNotFoundError, ValueError) as error:
        LOG.error("%s", error)
        return 2
    except PermissionError as error:
        LOG.error(
            "Permission denied: %s\n"
            "       Close the workbook in Excel (or pause OneDrive sync) and retry.",
            error,
        )
        return 2

    found = len(report["workbooks_found"])
    processed = report["workbooks_processed"]
    LOG.info(
        "Done. %s: %d of %d attribution workbooks written (%d warnings, %d errors).",
        report["period"],
        processed,
        found,
        len(report["warnings"]),
        len(report["errors"]),
    )
    if report["errors"]:
        LOG.error(
            "%d attribution workbook(s) failed. See %s for detail.",
            len(report["errors"]),
            DEBUG_REPORT_FILENAME,
        )
        return 1
    if found and not processed:
        LOG.error("No attribution files were produced despite finding %d workbook(s).", found)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
