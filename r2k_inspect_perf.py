"""Quick structural dump of the Morningstar performance workbook so the parser can
be fixed to its actual layout. Prints the first sheet's top rows + which cells the
date parser recognizes. RUN: python r2k_inspect_perf.py"""
from pathlib import Path
import os
import openpyxl
from r2k_perf_io import find_performance_file, parse_month, ID_COLS

BASE = Path(os.environ.get("R2KG_BASE", "."))


def main():
    path = find_performance_file()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print(f"FILE: {path.name}")
    print(f"SHEETS ({len(wb.sheetnames)}): {wb.sheetnames[:10]}")
    ws = wb[wb.sheetnames[0]]
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        rows.append(r)
        if i >= 14: break
    ncol = max((len(r) for r in rows), default=0)
    print(f"first sheet '{wb.sheetnames[0]}': showing {len(rows)} rows x up to 18 cols\n")
    for i, r in enumerate(rows):
        cells = [("" if c is None else str(c))[:16] for c in r[:18]]
        # mark id-col hits and date hits on this row
        idhits = [str(c).strip().lower() for c in r if c is not None and str(c).strip().lower() in ID_COLS]
        ndates = sum(1 for c in r if parse_month(c))
        flag = f"  <== idcols={idhits} dateCols={ndates}" if (idhits or ndates >= 6) else ""
        print(f"row {i:>2}: " + " | ".join(cells) + flag)
    # also show the LAST few rows (where the index/benchmark rows live)
    allrows = list(ws.iter_rows(values_only=True))
    print(f"\ntotal rows: {len(allrows)}")
    print("last 6 rows (first 6 cols):")
    for r in allrows[-6:]:
        print("   " + " | ".join(("" if c is None else str(c))[:24] for c in r[:6]))
    wb.close()


if __name__ == "__main__":
    main()
